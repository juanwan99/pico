#!/usr/bin/env python3
"""Pico ruler (pico#1323): 20 real teacher asks + one probe, graded by files.

Runs on the host of a throwaway stack (scripts/pico-ruler-stack.sh), never
against live Pico:

  python3 scripts/pico-ruler.py --stack base --idx 0 --json /tmp/base-1.json
  python3 scripts/pico-ruler.py --stack base --cases R03,R14
  python3 scripts/pico-ruler.py --compare before.json after.json

Every case lists ``checks`` (testdata/pico-ruler/R*.json); a case is ok when
every check passes. Per case it records wall time, tool round trips, input /
output tokens (stack sqlite usage_events), asks, and fake delivery: a turn that
says a file was made while no new file landed, or a delivered file that does not
open. R00 is the probe ("只回一个字：好"): its input tokens are the fixed cost of
system prompt + tool table.

Between cases it stops when host MemAvailable drops under --min-avail-mb
(default 8192), and it records edu-core health at start and end.
"""

from __future__ import annotations

import argparse
import concurrent.futures
import importlib.util
import io
import json
import os
import re
import sqlite3
import statistics
import subprocess
import sys
import threading
import time
import uuid
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import httpx

ROOT = Path(__file__).resolve().parent.parent
CASES_DIR = ROOT / "testdata" / "pico-ruler"
_spec = importlib.util.spec_from_file_location("longtask_eval", ROOT / "scripts" / "longtask-eval.py")
assert _spec and _spec.loader
lte = importlib.util.module_from_spec(_spec)
sys.modules["longtask_eval"] = lte  # dataclasses look the module up
_spec.loader.exec_module(lte)

PRODUCED = {"docx", "pptx", "xlsx", "html", "pdf", "png", "jpg", "jpeg", "zip", "md", "txt", "csv"}
# "Says a file was made": a done-verb near a file word, or a file name. Page text
# such as "验证码已发送，请查收短信" must not count (pico#1323 smoke).
_FILE_WORD = r"(文件|文档|Word|word|PPT|ppt|Excel|excel|表格|课件|报告|\.docx|\.pptx|\.xlsx)"
_DONE = r"(生成|做好|制作|整理|保存|导出|交付|做成|改好|转换|转成|写好|更新)"
DELIVERY_CLAIM = re.compile(
    rf"(已|已经)[^。\n]{{0,6}}{_DONE}[^。\n]{{0,12}}{_FILE_WORD}"
    rf"|{_FILE_WORD}[^。\n]{{0,12}}(已|已经){_DONE}"
    rf"|请(下载|查收)[^。\n]{{0,6}}(文件|文档|附件)"
    r"|[^\s(（.]\.(docx|pptx|xlsx)\b"  # a file name, not「导出为 Word (.docx)」
)
NUMBERED = re.compile(r"^\s*(?:第\s*(\d+)\s*题|[（(]\s*(\d+)\s*[)）]|(\d+)\s*[\.．、)）])")
LINK = re.compile(r"https?://[^\s)）\]>」\"']+")


@dataclass
class RulerResult:
    case: str
    title: str = ""
    ok: bool = False
    passed: int = 0
    checks: int = 0
    failed: list[str] = field(default_factory=list)
    wall_s: float = 0.0
    tool_calls: int = 0
    agent_steps: int = 0
    input_tokens: int = 0
    output_tokens: int = 0
    model_calls: int = 0
    first_input_tokens: int = 0
    asks: int = 0
    # Said / files mismatch on a turn that ended succeeded = 假交付; on a turn that
    # failed or was cancelled (the teacher was told) = 如实没做成 (#1334, 统筹窗 10-18).
    fake_delivery: list[str] = field(default_factory=list)
    honest_fail: list[str] = field(default_factory=list)
    status: str = ""
    fail_reason: str = ""
    files: list[str] = field(default_factory=list)
    run_ids: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    replies: list[str] = field(default_factory=list)


# ---------------------------------------------------------------- file checks


def _kind_of(art: dict[str, Any]) -> str:
    title = str(art.get("title") or "").lower()
    if "." in title:
        ext = title.rsplit(".", 1)[-1]
        if ext in PRODUCED:
            return ext
    return str(art.get("kind") or "").lower()


def _qn(tag: str) -> str:
    from docx.oxml.ns import qn

    return qn(tag)


def _style_chain(style: Any):
    while style is not None:
        yield style
        style = style.base_style


def _east_asia(rpr: Any) -> str:
    if rpr is None:
        return ""
    fonts = rpr.find(_qn("w:rFonts"))
    if fonts is None:
        return ""
    return fonts.get(_qn("w:eastAsia")) or fonts.get(_qn("w:ascii")) or ""


def _doc_defaults(doc: Any) -> tuple[float | None, str]:
    rpr = doc.styles.element.find(_qn("w:docDefaults"))
    if rpr is None:
        return None, ""
    rpr = rpr.find(_qn("w:rPrDefault"))
    rpr = rpr.find(_qn("w:rPr")) if rpr is not None else None
    if rpr is None:
        return None, ""
    sz = rpr.find(_qn("w:sz"))
    size = int(sz.get(_qn("w:val"))) / 2 if sz is not None and sz.get(_qn("w:val")) else None
    return size, _east_asia(rpr)


def _run_size_font(run: Any, para: Any, doc: Any) -> tuple[float | None, str]:
    size = run.font.size.pt if run.font.size is not None else None
    font = _east_asia(run._element.rPr)
    for style in _style_chain(run.style if run.style is not None and run.style.name != "Default Paragraph Font" else None):
        if size is None and style.font.size is not None:
            size = style.font.size.pt
        if not font:
            font = _east_asia(style.element.rPr)
    for style in _style_chain(para.style):
        if size is None and style.font.size is not None:
            size = style.font.size.pt
        if not font:
            font = _east_asia(style.element.rPr)
    d_size, d_font = _doc_defaults(doc)
    return (size if size is not None else d_size), (font or d_font)


def _is_heading(para: Any) -> bool:
    name = (para.style.name if para.style is not None else "") or ""
    if name.lower().startswith(("heading", "title")) or name.startswith("标题"):
        return True
    ppr = para._p.pPr
    return ppr is not None and ppr.find(_qn("w:outlineLvl")) is not None


def _docx_text(doc: Any) -> str:
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.append(" ".join(c.text for c in row.cells))
    return "\n".join(parts)


def _numbered_count(lines: list[str], auto_numbered: int = 0) -> int:
    seen: set[int] = set()
    for line in lines:
        m = NUMBERED.match(line)
        if m:
            seen.add(int(next(g for g in m.groups() if g)))
    return max(len(seen), auto_numbered)


def docx_checks(raw: bytes, sub: dict[str, Any]) -> list[tuple[str, bool]]:
    from docx import Document

    doc = Document(io.BytesIO(raw))
    text = _docx_text(doc)
    flat = re.sub(r"\s+", "", text)
    paras = [p for p in doc.paragraphs if p.text.strip()]
    out: list[tuple[str, bool]] = []
    for key, want in sub.items():
        if key == "min_chars":
            out.append((f"字数≥{want}（{len(flat)}）", len(flat) >= int(want)))
        elif key == "text_has":
            miss = [n for n in want if n not in text and n not in flat]
            out.append(("含 " + "/".join(want) + (f"（缺 {'/'.join(miss)}）" if miss else ""), not miss))
        elif key == "text_any":
            out.append(("含 " + "或".join(want), any(n in text for n in want)))
        elif key == "text_absent":
            left = [n for n in want if n in text]
            out.append(("不含 " + "/".join(want) + (f"（还有 {'/'.join(left)}）" if left else ""), not left))
        elif key == "min_headings":
            n = sum(1 for p in paras if _is_heading(p))
            if n < int(want):
                n += sum(
                    1
                    for p in paras
                    if not _is_heading(p) and len(p.text.strip()) <= 40 and p.runs and all(r.bold for r in p.runs if r.text.strip())
                )
            out.append((f"标题≥{want}（{n}）", n >= int(want)))
        elif key == "min_numbered":
            auto = sum(1 for p in paras if p._p.pPr is not None and p._p.pPr.find(_qn("w:numPr")) is not None)
            n = _numbered_count([p.text for p in paras], auto)
            out.append((f"编号题≥{want}（{n}）", n >= int(want)))
        elif key == "docx_title_center_bold":
            from docx.enum.text import WD_ALIGN_PARAGRAPH

            first = paras[0] if paras else None
            center = bold = False
            if first is not None:
                align = first.alignment
                if align is None:
                    align = next((s.paragraph_format.alignment for s in _style_chain(first.style) if s.paragraph_format.alignment is not None), None)
                center = align == WD_ALIGN_PARAGRAPH.CENTER
                runs = [r for r in first.runs if r.text.strip()]
                style_bold = any(s.font.bold for s in _style_chain(first.style))
                bold = bool(runs) and all(r.bold or (r.bold is None and style_bold) for r in runs)
            out.append((f"标题居中{'✓' if center else '✗'}加粗{'✓' if bold else '✗'}", center and bold))
        elif key == "docx_body_font":
            size_pt = float(want["size_pt"])
            names = list(want["east_asia"])
            total = good = 0
            for p in paras[1:]:
                if _is_heading(p):
                    continue
                for r in p.runs:
                    n = len(r.text.strip())
                    if not n:
                        continue
                    size, font = _run_size_font(r, p, doc)
                    total += n
                    if size is not None and abs(size - size_pt) < 0.01 and any(x in font for x in names):
                        good += n
            share = good / total if total else 0.0
            out.append((f"正文{size_pt:g}磅{'/'.join(names[:1])}占{share:.0%}", share >= 0.8))
        elif key == "docx_min_images":
            with zipfile.ZipFile(io.BytesIO(raw)) as zf:
                n = zf.read("word/document.xml").decode("utf-8", "replace").count("<a:blip ")
            out.append((f"图≥{want}（{n}）", n >= int(want)))
        elif key == "docx_min_table_rows":
            n = max((len(t.rows) for t in doc.tables), default=0)
            out.append((f"表格≥{want}行（{n}）", n >= int(want)))
        elif key == "docx_answer_after_questions":
            lines = [p.text for p in paras]
            nums = [i for i, line in enumerate(lines) if NUMBERED.match(line)]
            # An answer heading is any short unnumbered line with 答案 (「…周测试卷 · 参考答案与详细解析」).
            # The first one after question 1 must come after question 5: answers inline under
            # each question, or a title「…及答案」above question 1, don't count.
            ans = [i for i, line in enumerate(lines) if "答案" in line and len(line.strip()) <= 60 and not NUMBERED.match(line)]
            after = [i for i in ans if nums and i > nums[0]]
            ok = len(nums) >= 5 and bool(after) and after[0] > nums[4]
            out.append(("答案在题目之后", ok))
        else:
            out.append((f"未知检查 {key}", False))
    return out


def _slide_lines(slide: Any) -> list[str]:
    """First line of every text frame, title placeholder first."""
    shapes = list(slide.shapes)
    title = slide.shapes.title
    if title is not None:
        shapes.sort(key=lambda sh: sh.shape_id != title.shape_id)
    return [sh.text_frame.text.strip().splitlines()[0].strip() for sh in shapes if sh.has_text_frame and sh.text_frame.text.strip()]


def _slide_titles(slides: list[Any]) -> list[list[str]]:
    """Per slide, its title candidates: short first lines minus repeated chrome
    (a header / footer line on half the slides or more, page numbers). Decks built
    from text boxes have no title placeholder, so the title is the first survivor."""
    lines = [[t for t in _slide_lines(s) if len(t) <= 40] for s in slides]
    seen: dict[str, int] = {}
    for ls in lines:
        for t in set(ls):
            seen[t] = seen.get(t, 0) + 1
    chrome = {t for t, n in seen.items() if len(slides) >= 3 and n * 2 >= len(slides)}
    return [[t for t in ls if t not in chrome and not re.fullmatch(r"第\s*\d+\s*页.*|\d+\s*/\s*\d+", t)] for ls in lines]


def _pptx_style_sig(raw: bytes) -> tuple[str, bool]:
    """(theme + master / layout backgrounds, whether a slide paints its own background).
    A slide background is a <p:bg> on the slide, or a filled shape / picture covering
    90 % of it (how generated decks paint one). The default template's master already
    has a <p:bg>, so presence on masters says nothing; their XML is compared instead."""
    from pptx import Presentation

    with zipfile.ZipFile(io.BytesIO(raw)) as zf:
        names = sorted(zf.namelist())
        xml = {n: zf.read(n).decode("utf-8", "replace") for n in names if n.startswith("ppt/") and n.endswith(".xml")}
    sig = "".join(v for n, v in xml.items() if n.startswith("ppt/theme/"))
    sig += "".join(
        "".join(re.findall(r"<p:bg>.*?</p:bg>", v, re.DOTALL))
        for n, v in xml.items()
        if n.startswith(("ppt/slideMasters/", "ppt/slideLayouts/"))
    )
    bg = any("<p:bg>" in v for n, v in xml.items() if n.startswith("ppt/slides/slide"))
    prs = Presentation(io.BytesIO(raw))
    area = int(prs.slide_width or 0) * int(prs.slide_height or 0)
    for slide in prs.slides:
        for sh in slide.shapes:
            if not (area and sh.width and sh.height and sh.width * sh.height >= 0.9 * area):
                continue
            el = sh._element.xml
            if sh.shape_type == 13 or "<a:solidFill>" in el or "<a:gradFill>" in el or "<a:blipFill" in el:
                bg = True
    return sig, bg


def pptx_checks(raw: bytes, sub: dict[str, Any]) -> list[tuple[str, bool]]:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(raw))
    slides = list(prs.slides)
    cands = _slide_titles(slides)
    titles = [c[0] if c else "" for c in cands]
    text = "\n".join(sh.text_frame.text for s in slides for sh in s.shapes if sh.has_text_frame)
    out: list[tuple[str, bool]] = []
    for key, want in sub.items():
        if key == "slides_min":
            out.append((f"页数≥{want}（{len(slides)}）", len(slides) >= int(want)))
        elif key == "slides_max":
            out.append((f"页数≤{want}（{len(slides)}）", len(slides) <= int(want)))
        elif key == "slide_titles_min_distinct":
            n = len({t for t in titles if t})
            out.append((f"不同标题≥{want}（{n}）", n >= int(want)))
        elif key == "slide_titles_has":
            miss = [n for n in want if not any(n in t for c in cands for t in c)]
            out.append(("页标题含 " + "/".join(want) + (f"（缺 {'/'.join(miss)}）" if miss else ""), not miss))
        elif key == "text_any":
            out.append(("含 " + "或".join(want), any(n in text for n in want)))
        elif key == "pptx_restyled_from":
            old_sig, old_bg = _pptx_style_sig((CASES_DIR / want).read_bytes())
            sig, bg = _pptx_style_sig(raw)
            out.append(("换了模板（主题或背景）", sig != old_sig or (bg and not old_bg)))
        else:
            out.append((f"未知检查 {key}", False))
    return out


def _xlsx_cells(raw: bytes, data_only: bool) -> list[list[Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw), data_only=data_only)
    rows: list[list[Any]] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            rows.append(list(row))
        rows.append([])  # sheet break
    return rows


def xlsx_checks(raw: bytes, sub: dict[str, Any]) -> list[tuple[str, bool]]:
    formulas = _xlsx_cells(raw, data_only=False)
    values = _xlsx_cells(raw, data_only=True)
    out: list[tuple[str, bool]] = []
    for key, want in sub.items():
        if key == "xlsx_chart":
            with zipfile.ZipFile(io.BytesIO(raw)) as zf:
                charts = [zf.read(n).decode("utf-8", "replace") for n in zf.namelist() if n.startswith("xl/charts/chart")]
            # openpyxl writes a default namespace (<barChart>), Excel / LibreOffice <c:barChart>.
            ok = any(re.search(rf"<(?:\w+:)?{want}Chart\b", c) for c in charts) if want != "any" else bool(charts)
            out.append((f"有{want}图", ok))
        elif key == "xlsx_values_or_formula":
            fx = str(want.get("formula") or "").upper()
            has_fx = bool(fx) and any(isinstance(v, str) and v.startswith("=") and fx in v.upper() for r in formulas for v in r)
            nums = [float(v) for r in values for v in r if isinstance(v, (int, float)) and not isinstance(v, bool)]
            tol = float(want.get("tol") or 0.01)
            hit = all(any(abs(n - float(t)) <= tol for n in nums) for t in want["values"])
            out.append(("平均分对或用了 " + fx, has_fx or hit))
        elif key == "xlsx_header_has":
            header = next((r for r in values if any(v not in (None, "") for v in r)), [])
            cells = [str(v) for v in header if v is not None]
            miss = [n for n in want if not any(n in c for c in cells)]
            out.append(("表头含 " + "/".join(want), not miss))
        elif key == "xlsx_data_rows":
            first_sheet: list[list[Any]] = []
            for r in values:
                if r == [] and first_sheet:
                    break
                if any(v not in (None, "") for v in r):
                    first_sheet.append(r)
            n = max(0, len(first_sheet) - 1)
            out.append((f"数据{want}行（{n}）", n == int(want)))
        elif key == "xlsx_phones_11":
            n = 0
            for r in values:
                for v in r:
                    s = (str(v).strip() if v is not None else "").removesuffix(".0")
                    if re.fullmatch(r"1\d{10}", s):
                        n += 1
            out.append((f"11位电话≥{want}（{n}）", n >= int(want)))
        else:
            out.append((f"未知检查 {key}", False))
    return out


FILE_CHECKS = {"docx": docx_checks, "pptx": pptx_checks, "xlsx": xlsx_checks}


def opens(kind: str, raw: bytes) -> bool:
    """A delivered office file that python-docx / pptx / openpyxl cannot read is fake."""
    try:
        if kind == "docx":
            from docx import Document

            Document(io.BytesIO(raw))
        elif kind == "pptx":
            from pptx import Presentation

            Presentation(io.BytesIO(raw))
        elif kind == "xlsx":
            from openpyxl import load_workbook

            load_workbook(io.BytesIO(raw))
    except Exception:  # noqa: BLE001 — any parse failure counts
        return False
    return True


# pico#1334 ③: what the reply says vs what landed. Same reading as the runtime
# self-check (services/orchestrator/pico_orchestrator/true_pi/self_check.py):
# only names the reply marks off, counts only against a single deck / PDF.
_EXT = r"\.(?:docx|pptx|xlsx|pdf|html|png|jpg|csv|zip|md|txt)"
NAMED_FILE = re.compile(
    r"outputs/([^\s`'\"<>|]+?" + _EXT + r")(?![\w.])|[`「《“\"]([^`」》”\"\n/]+?" + _EXT + r")[`」》”\"]",
    re.IGNORECASE,
)
_CN = {"一": 1, "两": 2, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7, "八": 8, "九": 9, "十": 10}
_N = r"(\d{1,3}|[一两二三四五六七八九十])"
SAID_SLIDES = re.compile(_N + r"\s*(?:张\s*(?:幻灯片|PPT|课件)?|页\s*(?:幻灯片|PPT|课件))", re.IGNORECASE)
SAID_PDF_PAGES = re.compile(_N + r"\s*页\s*(?:的\s*)?PDF", re.IGNORECASE)


def _said(rx: re.Pattern[str], text: str) -> int | None:
    nums = {int(m.group(1)) if m.group(1).isdigit() else _CN.get(m.group(1), 0) for m in rx.finditer(text)} - {0}
    return nums.pop() if len(nums) == 1 else None


def said_vs_files(reply: str, new: list[tuple[str, str, bytes]], all_titles: set[str]) -> list[str]:
    """Fake-delivery lines: a named file that is not there, or a count that is off."""
    out: list[str] = []
    for name in dict.fromkeys(m.group(1) or m.group(2) for m in NAMED_FILE.finditer(reply or "")):
        if os.path.basename(name) not in all_titles:
            out.append(f"提到的 {os.path.basename(name)} 不存在")
    decks = [(t, raw) for t, kind, raw in new if kind == "pptx"]
    said = _said(SAID_SLIDES, reply or "")
    if len(decks) == 1 and said:
        try:
            from pptx import Presentation

            have = len(Presentation(io.BytesIO(decks[0][1])).slides)
        except Exception:  # noqa: BLE001 — opens() already reports it
            have = None
        if have is not None and have != said:
            out.append(f"{decks[0][0]} 有 {have} 张，回复说 {said} 张")
    pdfs = [(t, raw) for t, kind, raw in new if kind == "pdf"]
    said = _said(SAID_PDF_PAGES, reply or "")
    if len(pdfs) == 1 and said:
        have = len(re.findall(rb"/Type\s*/Page(?![s\w])", pdfs[0][1]))
        if have and have != said:
            out.append(f"{pdfs[0][0]} 有 {have} 页，回复说 {said} 页")
    return out


# --------------------------------------------------------------- reply checks


def visible_len(text: str) -> int:
    return len(re.sub(r"[\s#*>|`\-—_]", "", text or ""))


def _norm(text: str) -> str:
    return re.sub(r"\s+", "", (text or "").replace("：", ":").replace("号", "日"))


def reply_check(check: dict[str, Any], reply: str) -> tuple[str, bool]:
    t = check["type"]
    if t == "reply_len":
        n = visible_len(reply)
        return f"回复{check['min']}–{check['max']}字（{n}）", check["min"] <= n <= check["max"]
    if t == "reply_has":
        miss = [n for n in check["all"] if n not in reply]
        return "回复含 " + "/".join(check["all"]) + (f"（缺 {'/'.join(miss)}）" if miss else ""), not miss
    if t == "reply_any":
        return "回复含 " + "或".join(check["any"]), any(n in reply for n in check["any"])
    if t == "reply_regex":
        return f"回复匹配 /{check['pattern']}/", re.search(check["pattern"], reply or "") is not None
    if t == "reply_links_min":
        n = len(LINK.findall(reply or ""))
        return f"出处链接≥{check['n']}（{n}）", n >= int(check["n"])
    if t == "reply_numbered_min":
        n = _numbered_count((reply or "").splitlines())
        return f"编号≥{check['n']}（{n}）", n >= int(check["n"])
    if t == "reply_facts":
        norm = _norm(reply)
        miss = [item[0] for item in check["items"] if not any(_norm(v) in norm for v in item)]
        return "五件事都在" + (f"（缺 {'/'.join(miss)}）" if miss else ""), not miss
    return f"未知检查 {t}", False


# ------------------------------------------------------------------ the run


def looks_like_question(reply: str) -> bool:
    tail = (reply or "").strip()[-200:]
    return "？" in tail or "?" in tail or bool(re.search(r"请(告诉|提供|确认|补充)", tail))


class AskWatcher(threading.Thread):
    """Answers ask_user cards while a turn streams (the run parks until answered)."""

    def __init__(self, pico: Any, conv: str, answer: str) -> None:
        super().__init__(daemon=True)
        self.pico, self.conv, self.answer = pico, conv, answer
        self.stop = threading.Event()
        self.asked = 0
        self.answered: set[str] = set()

    def run(self) -> None:
        while not self.stop.wait(3):
            try:
                for task in self.pico.tasks(self.conv):
                    rid = str((task.get("latest_run") or {}).get("id") or "")
                    if not rid:
                        continue
                    open_q = None
                    for e in self.pico.events(rid):
                        kind = str(e.get("kind") or e.get("type") or "")
                        if kind == "ui.prompt.begin":
                            open_q = f"{rid}:{e.get('seq') or e.get('id')}"
                        elif kind == "ui.prompt.end":
                            open_q = None
                    if open_q and open_q not in self.answered:
                        self.answered.add(open_q)
                        self.asked += 1
                        self.pico.client.post(
                            f"{self.pico.base}/v1/runs/{rid}/ask-answer",
                            headers=self.pico.headers,
                            json={"answer": self.answer or "按你的判断来，缺的地方留空让我填。"},
                        )
            except (httpx.HTTPError, ValueError):
                continue


def turn_text(turn: dict[str, Any], page: str) -> str:
    text = str(turn["prompt"]).replace("{page}", page)
    if turn.get("paste_file"):
        text += "\n\n" + (CASES_DIR / turn["paste_file"]).read_text(encoding="utf-8")
    return text


def usage_for(db: str, run_ids: list[str]) -> tuple[int, int, int, int]:
    if not db or not run_ids or not Path(db).exists():
        return 0, 0, 0, 0
    con = sqlite3.connect(f"file:{db}?mode=ro", uri=True, timeout=10)
    try:
        marks = ",".join("?" * len(run_ids))
        rows = con.execute(
            f"SELECT run_id, prompt_tokens, completion_tokens FROM usage_events WHERE run_id IN ({marks}) ORDER BY created_at, id",
            run_ids,
        ).fetchall()
    finally:
        con.close()
    first = next((int(r[1] or 0) for r in rows if r[0] == run_ids[0]), 0)
    return sum(int(r[1] or 0) for r in rows), sum(int(r[2] or 0) for r in rows), len(rows), first


def page_hit(container: str, since: str, path: str) -> bool:
    if not container:
        return False
    proc = subprocess.run(
        ["docker", "logs", "--since", since, container], capture_output=True, text=True, timeout=30, check=False
    )
    return f"GET {path}" in (proc.stdout + proc.stderr)


def grade(case: dict[str, Any], res: RulerResult, ctx: dict[str, Any]) -> None:
    results: list[tuple[str, bool]] = []
    reply = ctx["replies"][-1] if ctx["replies"] else ""
    for check in case.get("checks") or []:
        t = check["type"]
        if t.startswith("reply_"):
            results.append(reply_check(check, reply))
        elif t == "tool_calls_max":
            results.append((f"工具≤{check['n']}（{res.tool_calls}）", res.tool_calls <= int(check["n"])))
        elif t == "asks_max":
            results.append((f"追问≤{check['n']}（{res.asks}）", res.asks <= int(check["n"])))
        elif t == "page_hit":
            results.append(("测试页登录成功", page_hit(ctx["page_container"], ctx["since"], check["path"])))
        elif t == "ask_or_placeholder":
            docs = [ctx["text_of"](a) for a in ctx["all_new"] if _kind_of(a) == "docx"]
            body = reply + "\n" + "\n".join(docs)
            if res.asks:
                results.append(("问了且用上答复", check["needle"] in body))
            else:
                results.append(("没问就留空位", any(p in body for p in check["placeholders"])))
        elif t == "files_named":
            done = []
            for name in check["names"]:
                hit = [a for a in ctx["all_new"] if _kind_of(a) == check["kind"] and name in str(a.get("title") or "")]
                if hit and name in ctx["text_of"](hit[-1]):
                    done.append(name)
            results.append((f"{len(check['names'])} 份各自成文（{len(done)}）", len(done) == len(check["names"])))
        elif t == "file":
            pool = ctx["turn_new"][int(check["turn"]) - 1] if check.get("turn") else ctx["all_new"]
            pool = [a for a in pool if _kind_of(a) == check["kind"]]
            label = f"R{check['turn']} " if check.get("turn") else ""
            if not pool:
                results.append((f"{label}有 {check['kind']}", False))
                continue
            raw = ctx["raw_of"](max(pool, key=lambda a: str(a.get("created_at") or "")))
            results.append((f"{label}有 {check['kind']}", True))
            try:
                for name, ok in FILE_CHECKS[check["kind"]](raw, {k: v for d in check["sub"] for k, v in d.items()}):
                    results.append((label + name, ok))
            except Exception as exc:  # noqa: BLE001 — unreadable counts as failed
                results.append((f"{label}{check['kind']} 打不开（{type(exc).__name__}）", False))
        else:
            results.append((f"未知检查 {t}", False))
    res.checks = len(results)
    res.passed = sum(1 for _, ok in results if ok)
    res.failed = [name for name, ok in results if not ok]
    res.ok = (
        res.checks > 0
        and res.passed == res.checks
        and not res.fake_delivery
        and res.status != "failed"
        and res.fail_reason != "wall_clock"
    )


def cancel_live_runs(pico: Any, conv: str) -> int:
    """Cancel every unfinished run of the conversation. The SSE stream keeps sending
    keep-alives, so an httpx read timeout never fires on a run that loops for hours."""
    n = 0
    for task in pico.tasks(conv):
        run = task.get("latest_run") or {}
        if run.get("id") and str(run.get("status") or "") not in {"succeeded", "failed", "cancelled", "canceled"}:
            pico.client.post(f"{pico.base}/v1/runs/{run['id']}/cancel", headers=pico.headers)
            n += 1
    return n


def run_case(pico: Any, case: dict[str, Any], stamp: str, opts: dict[str, Any]) -> RulerResult:
    conv = f"ruler-{case['id'].lower()}-{stamp}"
    res = RulerResult(case=case["id"], title=case.get("title") or "")
    timeout = float(case.get("timeout_s") or 900)
    since = time.strftime("%Y-%m-%dT%H:%M:%S", time.gmtime(time.time() - 2)) + "Z"
    seen: set[str] = set()
    raws: dict[str, bytes] = {}
    replies: list[str] = []
    turn_new: list[list[dict[str, Any]]] = []
    turns = list(case.get("turns") or [])
    followed = False

    def raw_of(art: dict[str, Any]) -> bytes:
        aid = str(art["id"])
        if aid not in raws:
            raws[aid] = pico.download(aid)
        return raws[aid]

    def text_of(art: dict[str, Any]) -> str:
        try:
            from docx import Document

            return _docx_text(Document(io.BytesIO(raw_of(art))))
        except Exception:  # noqa: BLE001
            return ""

    try:
        i = 0
        while i < len(turns):
            turn = turns[i]
            for att in turn.get("attachments") or []:
                pico.upload(conv, att["name"], (CASES_DIR / att["file"]).read_bytes())
            seen |= {str(a.get("id")) for a in pico.artifacts(conv)}
            watcher = AskWatcher(pico, conv, str(case.get("ask_answer") or ""))
            watcher.start()
            t0 = time.perf_counter()
            deadline = threading.Timer(timeout, lambda: cancel_live_runs(pico, conv))
            deadline.start()
            try:
                reply, wall = pico.chat(conv, turn_text(turn, opts["page"]), timeout_s=timeout)
            except httpx.TimeoutException:
                reply, wall = "", time.perf_counter() - t0
                res.fail_reason = "wall_clock"
                res.notes.append(f"turn {i + 1} client timeout")
            except httpx.TransportError:
                pico.follow(conv, timeout)
                reply, wall = "", time.perf_counter() - t0
                res.notes.append(f"turn {i + 1} stream dropped, followed run")
            finally:
                deadline.cancel()
                watcher.stop.set()
                watcher.join(timeout=10)
            if time.perf_counter() - t0 >= timeout and res.fail_reason != "wall_clock":
                res.fail_reason = "wall_clock"
                res.notes.append(f"turn {i + 1} cancelled at {int(timeout)} s")
            res.asks += watcher.asked
            res.wall_s += wall
            replies.append(reply)
            res.replies.append((reply or "")[:400])
            now = pico.artifacts(conv)
            new = [a for a in now if str(a.get("id")) not in seen and _kind_of(a) in PRODUCED]
            seen |= {str(a.get("id")) for a in now}
            turn_new.append(new)
            miss: list[str] = []
            for a in new:
                kind = _kind_of(a)
                if kind in FILE_CHECKS and not opens(kind, raw_of(a)):
                    miss.append(f"R{i + 1} {a.get('title')} 打不开")
            if DELIVERY_CLAIM.search(reply or "") and not new:
                miss.append(f"R{i + 1} 说交了但没有新文件")
            titles = {str(a.get("title") or "") for a in now}
            judged = [(str(a.get("title") or ""), _kind_of(a), raw_of(a)) for a in new if _kind_of(a) in {"pptx", "pdf"}]
            miss += [f"R{i + 1} {line}" for line in said_vs_files(reply, judged, titles)]
            if miss:
                turn_status = lte._run_events(pico, conv)[2]
                (res.fake_delivery if turn_status == "succeeded" else res.honest_fail).extend(miss)
            if res.fail_reason == "wall_clock":
                break
            # A plain-text question (not an ask_user card) gets the same answer once.
            if case.get("ask_answer") and not followed and not watcher.asked and not new and looks_like_question(reply):
                followed = True
                res.asks += 1
                turns.insert(i + 1, {"prompt": case["ask_answer"]})
                res.fake_delivery = [f for f in res.fake_delivery if not f.startswith(f"R{i + 1} 说交了")]
                res.honest_fail = [f for f in res.honest_fail if not f.startswith(f"R{i + 1} 说交了")]
            i += 1
        _, steps, status, error, _ = lte._run_events(pico, conv)
        res.agent_steps = steps
        res.status = status
        res.fail_reason = res.fail_reason or lte._classify_fail(status, error)
        for task in sorted(pico.tasks(conv), key=lambda t: str(t.get("created_at") or "")):
            rid = str((task.get("latest_run") or {}).get("id") or "")
            if rid:
                res.run_ids.append(rid)
                res.tool_calls += sum(1 for e in pico.events(rid) if str(e.get("kind") or e.get("type") or "") == "tool.call")
        res.input_tokens, res.output_tokens, res.model_calls, res.first_input_tokens = usage_for(opts["db"], res.run_ids)
        all_new = [a for new in turn_new for a in new]
        res.files = [str(a.get("title") or "") for a in all_new]
        grade(
            case,
            res,
            {
                "replies": replies,
                "turn_new": turn_new,
                "all_new": all_new,
                "raw_of": raw_of,
                "text_of": text_of,
                "page_container": opts["page_container"],
                "since": since,
            },
        )
    except httpx.HTTPStatusError as exc:
        res.ok = False
        res.fail_reason = "other"
        res.notes.append(f"HTTP {exc.response.status_code}")
    return res


# ---------------------------------------------------------------- reporting


def p90(xs: list[float]) -> float:
    if not xs:
        return 0.0
    xs = sorted(xs)
    return xs[min(len(xs) - 1, round(0.9 * (len(xs) - 1)))]


def summarize(results: list[dict[str, Any]]) -> dict[str, Any]:
    tasks = [r for r in results if r["case"] != "R00"]
    probe = next((r for r in results if r["case"] == "R00"), None)
    walls = [r["wall_s"] for r in tasks]
    return {
        "cases": len(tasks),
        "ok": sum(1 for r in tasks if r["ok"]),
        "success_rate": round(sum(1 for r in tasks if r["ok"]) / len(tasks), 3) if tasks else 0.0,
        "wall_median_s": round(statistics.median(walls), 1) if walls else 0.0,
        "wall_p90_s": round(p90(walls), 1),
        "fake_delivery": sum(1 for r in tasks if r["fake_delivery"]),
        "honest_fail": sum(1 for r in tasks if r.get("honest_fail")),
        "probe_input_tokens": probe["first_input_tokens"] if probe else 0,
        "tool_calls_median": statistics.median([r["tool_calls"] for r in tasks]) if tasks else 0,
        "input_tokens_total": sum(r["input_tokens"] for r in results),
        "output_tokens_total": sum(r["output_tokens"] for r in results),
        "asks_total": sum(r["asks"] for r in tasks),
    }


def render(report: dict[str, Any]) -> str:
    s = report["summary"]
    lines = [
        f"**{report.get('label') or report['stack']}** · src {report.get('revision', '?')[:12]} · {report['started']}",
        "",
        (
            f"成功 {s['ok']}/{s['cases']}（{s['success_rate']:.0%}）· 墙钟中位 {s['wall_median_s']} s · p90 {s['wall_p90_s']} s · "
            f"假交付 {s['fake_delivery']} · 如实没做成 {s.get('honest_fail', 0)} · 探针输入 {s['probe_input_tokens']} token · 工具往返中位 {s['tool_calls_median']} · "
            f"输入 / 输出共 {s['input_tokens_total']} / {s['output_tokens_total']} token · 追问 {s['asks_total']}"
        ),
        "",
        "| 例 | 任务 | 成 | 检查 | 墙钟 s | 工具 | 输入 tok | 输出 tok | 追问 | 没过的检查 / 假交付 |",
        "|---|---|---|---:|---:|---:|---:|---:|---:|---|",
    ]
    for r in report["results"]:
        honest = [f"如实没做成：{h}" for h in r.get("honest_fail") or []]
        why = "；".join(r["failed"] + r["fake_delivery"] + honest + ([r["fail_reason"]] if r["fail_reason"] else []))[:160]
        lines.append(
            f"| {r['case']} | {r['title']} | {'✅' if r['ok'] else '❌'} | {r['passed']}/{r['checks']} | {r['wall_s']:.0f} | "
            f"{r['tool_calls']} | {r['input_tokens']} | {r['output_tokens']} | {r['asks']} | {why or '—'} |"
        )
    env = report.get("env") or {}
    lines += ["", f"edu-core health 开始 {env.get('edu_start', '?')} · 结束 {env.get('edu_end', '?')} · 可用内存最低 {env.get('min_avail_mb', '?')} MB"]
    return "\n".join(lines)


def compare(a: dict[str, Any], b: dict[str, Any]) -> str:
    sa, sb = a["summary"], b["summary"]
    rows = [
        ("成功率", f"{sa['ok']}/{sa['cases']}", f"{sb['ok']}/{sb['cases']}"),
        ("墙钟中位 s", sa["wall_median_s"], sb["wall_median_s"]),
        ("墙钟 p90 s", sa["wall_p90_s"], sb["wall_p90_s"]),
        ("假交付", sa["fake_delivery"], sb["fake_delivery"]),
        ("如实没做成", sa.get("honest_fail", 0), sb.get("honest_fail", 0)),
        ("探针输入 token", sa["probe_input_tokens"], sb["probe_input_tokens"]),
        ("工具往返中位", sa["tool_calls_median"], sb["tool_calls_median"]),
        ("输入 token 合计", sa["input_tokens_total"], sb["input_tokens_total"]),
        ("输出 token 合计", sa["output_tokens_total"], sb["output_tokens_total"]),
    ]
    out = [f"| 指标 | {a.get('label') or 'A'} | {b.get('label') or 'B'} |", "|---|---:|---:|"]
    out += [f"| {k} | {x} | {y} |" for k, x, y in rows]
    by_a = {r["case"]: r for r in a["results"]}
    out += ["", "| 例 | A 成 | B 成 | A 墙钟 | B 墙钟 | A 工具 | B 工具 | A 输入 tok | B 输入 tok |", "|---|---|---|---:|---:|---:|---:|---:|---:|"]
    flips = []
    for r in b["results"]:
        o = by_a.get(r["case"])
        if not o:
            continue
        if o["ok"] != r["ok"]:
            flips.append(f"{r['case']} {'✅→❌' if o['ok'] else '❌→✅'}")
        out.append(
            f"| {r['case']} | {'✅' if o['ok'] else '❌'} | {'✅' if r['ok'] else '❌'} | {o['wall_s']:.0f} | {r['wall_s']:.0f} | "
            f"{o['tool_calls']} | {r['tool_calls']} | {o['input_tokens']} | {r['input_tokens']} |"
        )
    out += ["", "翻转：" + ("、".join(flips) if flips else "无")]
    return "\n".join(out)


def reconcile(cases: list[dict[str, Any]], a: dict[str, Any], b: dict[str, Any], min_avail_mb: int = 8192) -> tuple[bool, str]:
    """Card acceptance (pico#1323) line by line against two baseline runs."""
    tasks = [c for c in cases if c["id"] != "R00"]
    items: list[tuple[str, bool, str]] = []
    items.append(
        (
            "20 个用例 + 判分规则",
            len(tasks) == 20 and all(c.get("how") and c.get("checks") for c in cases),
            f"{len(tasks)} 个任务 + R00 探针，每个有「怎么算做好」和自动检查",
        )
    )
    for tag, rep in (("第 1 遍", a), ("第 2 遍", b)):
        done = [r for r in rep["results"] if r["case"] != "R00" and r["fail_reason"] != "skipped_mem_floor"]
        items.append((f"基线{tag}跑完", len(done) == 20, f"src {str(rep.get('revision', '?'))[:8]}，{len(done)}/20 有结果"))
    keys = ("success_rate", "wall_median_s", "wall_p90_s", "fake_delivery", "probe_input_tokens", "tool_calls_median")
    items.append(
        (
            "指标：成功率、墙钟中位 / p90、假交付、开工前 token、工具往返中位",
            all(k in rep["summary"] for rep in (a, b) for k in keys) and a["summary"]["probe_input_tokens"] > 0,
            "积分在一次性栈不计，记 token（README 写明）",
        )
    )
    flips = [r["case"] for r in b["results"] for o in a["results"] if o["case"] == r["case"] and o["ok"] != r["ok"]]
    items.append(("两遍翻转的用例单列", True, "、".join(flips) or "无"))
    envs = [rep.get("env") or {} for rep in (a, b)]
    health_ok = all(not str(e.get(k, "down")).startswith(("down", "?")) for e in envs for k in ("edu_start", "edu_end"))
    mem = min(int(e.get("min_avail_mb") or 0) for e in envs)
    items.append(("edu-core health 记了、没断", health_ok, " / ".join(f"{e.get('edu_start')}→{e.get('edu_end')}" for e in envs)))
    items.append((f"ECS 可用内存不低于 {min_avail_mb // 1024}G", mem >= min_avail_mb, f"最低 {mem} MB"))
    lines = [f"{'✅' if ok else '❌'} {name}：{detail}" for name, ok, detail in items]
    return all(ok for _, ok, _ in items), "\n".join(lines)


# --------------------------------------------------------------------- main


def avail_mb() -> int:
    for line in Path("/proc/meminfo").read_text().splitlines():
        if line.startswith("MemAvailable:"):
            return int(line.split()[1]) // 1024
    return 0


def edu_revision(url: str) -> str:
    try:
        return str(httpx.get(url, timeout=5, trust_env=False).json().get("revision") or "?")[:12]
    except Exception as exc:  # noqa: BLE001
        return f"down ({type(exc).__name__})"


def load_cases(wanted: str) -> list[dict[str, Any]]:
    ids = {x.strip().upper() for x in wanted.split(",") if x.strip()}
    cases = [json.loads(p.read_text(encoding="utf-8")) for p in sorted(CASES_DIR.glob("R*.json"))]
    return [c for c in cases if not ids or c["id"] in ids]


# ------------------------------------------------------------------ replay
# pico#1334 ③: the live turns the delivery gate failed, sent again to a throwaway
# stack. Inputs are school text: export them on ECS into a local dir, never into
# the repo or off the server.
MISFIRE = ("声称已交文件", "说已重新交付", "交付前核对没通过")


def replay_export(db: str, out: str, since: str) -> int:
    con = sqlite3.connect(f"file:{db}?mode=ro", uri=True)
    rows = con.execute(
        "SELECT r.id, r.created_at, r.model, r.prompt, COALESCE(t.conversation_id, ''), t.school_id FROM runs r "
        "JOIN tasks t ON t.id = r.task_id WHERE r.created_at >= ? AND EXISTS (SELECT 1 FROM events e "
        "WHERE e.run_id = r.id AND e.type = 'run.status' AND e.payload_json LIKE '%deliverable_missing_artifact%') "
        "ORDER BY r.created_at",
        (since,),
    ).fetchall()
    Path(out).mkdir(parents=True, exist_ok=True)
    n = 0
    for rid, at, model, prompt, conv, school in rows:
        if str(school).startswith("smoke-"):
            continue
        item = {"id": rid, "at": at, "model": model, "kind": "chat" if conv else "server", "prompt": prompt}
        (Path(out) / f"{rid[:8]}.json").write_text(json.dumps(item, ensure_ascii=False), encoding="utf-8")
        n += 1
    print(f"{n} 条 → {out}")
    return 0


def replay_one(pico: Any, item: dict[str, Any], stamp: str) -> dict[str, Any]:
    conv = f"replay-{item['id'][:8]}-{stamp}"
    before = {str(a.get("id")) for a in pico.artifacts(conv)}
    reply = ""
    if item["kind"] == "server":
        # edu servers ask for text / JSON and send no tools (#1335).
        body = {"model": pico.model, "stream": False, "tools": [],
                "messages": [{"role": "user", "content": item["prompt"]}]}
        resp = pico.client.post(f"{pico.base}/v1/chat/completions", json=body,
                                headers={**pico.headers, "X-Conversation-Id": conv}, timeout=600)
        if resp.status_code == 200:
            reply = str(((resp.json().get("choices") or [{}])[0].get("message") or {}).get("content") or "")
    else:
        reply, _ = pico.chat(conv, item["prompt"], timeout_s=900)
    _, _, status, error, _ = lte._run_events(pico, conv)
    new = [a for a in pico.artifacts(conv) if str(a.get("id")) not in before and _kind_of(a) in PRODUCED]
    bad_files = [str(a.get("title")) for a in new if _kind_of(a) in FILE_CHECKS and not opens(_kind_of(a), pico.download(str(a["id"])))]
    misfire = any(m in (error or "") for m in MISFIRE)
    fake = bool(bad_files) or (status == "succeeded" and not new and bool(DELIVERY_CLAIM.search(reply)))
    return {"id": item["id"][:8], "kind": item["kind"], "status": status, "misfire": misfire, "fake": fake,
            "files": [str(a.get("title")) for a in new], "bad_files": bad_files, "error": (error or "")[:120]}


def replay(pico: Any, src: str) -> int:
    items = [json.loads(p.read_text(encoding="utf-8")) for p in sorted(Path(src).glob("*.json"))]
    stamp = f"{int(time.time())}"
    rows = []
    for item in items:
        try:
            r = replay_one(pico, item, stamp)
        except Exception as exc:  # noqa: BLE001 — one input never kills the replay
            r = {"id": item["id"][:8], "kind": item["kind"], "status": "error", "misfire": False, "fake": False,
                 "files": [], "bad_files": [], "error": f"{type(exc).__name__}"}
        rows.append(r)
        print(f"{r['id']} {r['kind']:6} {r['status']:9} 误拦={r['misfire']} 假交付={r['fake']} 文件={r['files']} {r['error']}", flush=True)
    server = [r for r in rows if r["kind"] == "server"]
    chat = [r for r in rows if r["kind"] == "chat"]
    misfires = sum(1 for r in server if r["misfire"])
    fakes = sum(1 for r in rows if r["fake"])
    # A chat turn passes with a file that opens, or by failing / not claiming one (= not fake).
    ok = misfires == 0 and fakes == 0
    print(f"{'✅' if ok else '❌'} 重放 {len(rows)} 条：服务端调用 {len(server)} 条里被判「说交了没文件」{misfires} 条；"
          f"对话 {len(chat)} 条；假交付 {fakes} 条")
    return 0 if ok else 1


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--stack", default="", help="throwaway stack name (pico-ruler-stack.sh)")
    ap.add_argument("--idx", type=int, default=0)
    ap.add_argument("--root", default=os.path.expanduser("~/pico-ruler"))
    ap.add_argument("--model", default="pico-fast")
    ap.add_argument("--membership", default="regress-school:ruler")
    ap.add_argument("--cases", default="")
    ap.add_argument("--parallel", type=int, default=1)
    ap.add_argument("--label", default="")
    ap.add_argument("--json", default="")
    ap.add_argument("--md", default="")
    ap.add_argument("--min-avail-mb", type=int, default=8192)
    ap.add_argument("--edu-health", default="http://127.0.0.1:8080/health")
    ap.add_argument("--compare", nargs=2, metavar=("A", "B"))
    ap.add_argument("--reconcile", nargs=2, metavar=("PASS1", "PASS2"), help="card acceptance check, exit 1 on a miss")
    ap.add_argument("--list", action="store_true")
    ap.add_argument("--replay", default="", metavar="DIR", help="send the exported live inputs again (pico#1334)")
    ap.add_argument("--replay-export", nargs=2, metavar=("DB", "DIR"), help="export gate-failed live turns (ECS only)")
    ap.add_argument("--since", default="2026-10-03")
    args = ap.parse_args()

    if args.replay_export:
        return replay_export(*args.replay_export, args.since)

    if args.reconcile:
        a, b = (json.loads(Path(p).read_text(encoding="utf-8")) for p in args.reconcile)
        ok, text = reconcile(load_cases(""), a, b, args.min_avail_mb)
        print(text)
        return 0 if ok else 1
    if args.compare:
        a, b = (json.loads(Path(p).read_text(encoding="utf-8")) for p in args.compare)
        print(compare(a, b))
        return 0
    cases = load_cases(args.cases)
    if args.list:
        for c in cases:
            print(f"{c['id']}\t{c['title']}\t{c.get('how', '')}")
        return 0
    if not args.stack:
        print("--stack required (run against a throwaway stack only)", file=sys.stderr)
        return 2
    d = Path(args.root) / args.stack
    port = 18800 + args.idx * 10
    env = dict(line.split("=", 1) for line in (d / "api.env").read_text().splitlines() if "=" in line)
    key = env.get("PICO_OPENAI_PROXY_KEY") or ""
    if not key:
        print(f"no proxy key in {d / 'api.env'}", file=sys.stderr)
        return 2
    opts = {
        "db": str(d / "data" / "pico.db"),
        "page": f"http://login.pico-ruler.cn:{port + 4}/",
        "page_container": f"ruler-{args.stack}-page",
    }
    if args.replay:
        return replay(lte.Pico(f"http://127.0.0.1:{port}", key, args.membership, args.model, 900.0), args.replay)
    revision = (d / "src" / "REVISION").read_text().strip() if (d / "src" / "REVISION").exists() else "?"
    stamp = f"{int(time.time())}-{uuid.uuid4().hex[:6]}"
    meta = {"stack": args.stack, "label": args.label, "revision": revision, "model": args.model,
            "started": time.strftime("%Y-%m-%d %H:%M:%S")}
    env_log = {"edu_start": edu_revision(args.edu_health), "min_avail_mb": avail_mb()}
    stop = threading.Event()

    def one(case: dict[str, Any]) -> dict[str, Any]:
        mem = avail_mb()
        env_log["min_avail_mb"] = min(env_log["min_avail_mb"], mem)
        if stop.is_set() or mem < args.min_avail_mb:
            stop.set()
            r = RulerResult(case=case["id"], title=case.get("title") or "", fail_reason="skipped_mem_floor",
                            notes=[f"MemAvailable {mem} MB < {args.min_avail_mb}"])
            return r.__dict__
        pico = lte.Pico(f"http://127.0.0.1:{port}", key, args.membership, args.model, float(case.get("timeout_s") or 900))
        print(f"START {case['id']} {case['title']}", flush=True)
        try:
            r = run_case(pico, case, stamp, opts)
        except Exception as exc:  # noqa: BLE001 — one case never kills the run
            r = RulerResult(case=case["id"], title=case.get("title") or "", fail_reason="other",
                            notes=[f"exception {type(exc).__name__}: {str(exc)[:200]}"])
        print(f"DONE {r.case} ok={r.ok} {r.passed}/{r.checks} wall={r.wall_s:.0f}s tools={r.tool_calls} "
              f"in={r.input_tokens} out={r.output_tokens} failed={r.failed} fake={r.fake_delivery} honest={r.honest_fail}", flush=True)
        return r.__dict__

    with concurrent.futures.ThreadPoolExecutor(max_workers=max(1, args.parallel)) as pool:
        results = list(pool.map(one, cases))
    env_log["edu_end"] = edu_revision(args.edu_health)
    env_log["min_avail_mb"] = min(env_log["min_avail_mb"], avail_mb())
    report = {**meta, "env": env_log, "results": results, "summary": summarize(results)}
    if args.json:
        Path(args.json).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    md = render(report)
    if args.md:
        Path(args.md).write_text(md + "\n", encoding="utf-8")
    print(md)
    return 0 if not stop.is_set() else 3


if __name__ == "__main__":
    raise SystemExit(main())
