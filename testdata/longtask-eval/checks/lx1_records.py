"""LX1: 80 records past one context window → 期末评语.xlsx, every label against truth.

Truth is rebuilt from the same seed (lx_corpus.py). Points:
  structure (essential) · 进步方向 accuracy (essential) · 家校配合 accuracy
  (essential) · 评语 rules (essential) · 统计 sheet matches its own columns.
Accuracy bar: ≥90% overall and ≥85% on the first 20 files — those are read
first, so they are the ones a compaction would push out of the window.
"""

import os
import sys

sys.path.insert(0, "/w/check")
from _common import Score, find
from lx_corpus import HOME, PROGRESS, students

HEADER = ["学号", "姓名", "进步方向", "家校配合", "期末评语"]
FORBIDDEN = ("但是", "排名", "分数")
BAR_ALL, BAR_EARLY = 0.90, 0.85


def cell(v) -> str:
    return "" if v is None else str(v).strip()


def main() -> None:
    from openpyxl import load_workbook

    s = Score()
    path = find("期末评语.xlsx")
    if not path:
        s.fail("期末评语.xlsx not delivered")
        s.emit()
    wb = load_workbook(path, data_only=True)
    if "评语" not in wb.sheetnames:
        s.fail(f"no 评语 sheet: {wb.sheetnames}")
        s.emit()
    rows = [[cell(v) for v in r] for r in wb["评语"].iter_rows(values_only=True)]
    rows = [r for r in rows if any(r)]
    head, body = (rows[0][:5], rows[1:]) if rows else ([], [])
    truth = students()
    by_id = {t["id"]: t for t in truth}
    ids = [r[0] for r in body]
    s.check(
        head == HEADER
        and len(body) == len(truth)
        and ids == [t["id"] for t in truth]
        and all(by_id[r[0]]["name"] == r[1] for r in body if r[0] in by_id),
        f"structure: header={head} rows={len(body)} ids_sorted_match={ids == [t['id'] for t in truth]}",
    )
    got = {r[0]: r for r in body if r[0] in by_id}
    early = {t["id"] for t in truth if t["file"] < "21_"}

    def accuracy(col: int, key: str, label: str) -> None:
        hit = {i for i, r in got.items() if r[col] == by_id[i][key]}
        all_rate = len(hit) / len(truth)
        early_rate = len(hit & early) / len(early)
        s.notes.append(f"{label} {len(hit)}/{len(truth)} · first20 {len(hit & early)}/{len(early)}")
        s.check(all_rate >= BAR_ALL and early_rate >= BAR_EARLY, f"{label} below bar")

    accuracy(2, "progress", "进步方向")
    accuracy(3, "home", "家校配合")

    bad = []
    for i, r in got.items():
        text = r[4] if len(r) > 4 else ""
        if not (60 <= len(text) <= 120) or not text.startswith(by_id[i]["name"]) or any(
            w in text for w in FORBIDDEN
        ):
            bad.append(i)
    s.notes.append(f"评语 rules broken {len(bad)}/{len(truth)}")
    s.check(len(got) == len(truth) and len(bad) <= len(truth) * 0.05, "评语 rules")

    counts: dict[str, int] = {}
    for r in got.values():
        for v in (r[2], r[3]):
            counts[v] = counts.get(v, 0) + 1
    stat_ok = False
    if "统计" in wb.sheetnames:
        seen: dict[str, int] = {}
        for r in wb["统计"].iter_rows(values_only=True):
            vals = [cell(v) for v in r]
            for k, v in enumerate(vals[:-1]):
                if v in PROGRESS or v in HOME:
                    try:
                        seen[v] = int(float(vals[k + 1]))
                    except ValueError:
                        pass
        stat_ok = all(seen.get(lbl) == counts.get(lbl, 0) for lbl in [*PROGRESS, *HOME])
    s.check(stat_ok, "统计 sheet does not match 评语 columns", essential=False)
    s.emit()


if __name__ == "__main__":
    os.chdir(os.environ.get("TMPDIR", "/tmp"))
    main()
