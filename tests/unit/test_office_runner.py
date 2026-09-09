"""pico-office runner (#959): a computer, not a jail. Collect rule, receipts, cleanup."""

from __future__ import annotations

import os
import sys
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "services"))
sys.path.insert(0, str(ROOT / "services" / "orchestrator"))

from sandbox_worker.office_runner import (
    MAX_TIMEOUT_S,
    OFFICE_KINDS,
    RunBody,
    app,
    run_office_job,
)


def test_kinds_and_bad_args() -> None:
    assert set(OFFICE_KINDS) == {"docx", "xlsx", "pptx"}
    assert run_office_job(kind="pdf", source="x=1")["code"] == "tool.invalid_arguments"
    assert run_office_job(kind="docx", source="   ")["code"] == "tool.invalid_arguments"
    assert run_office_job(kind="docx", source="x" * 300_000)["code"] == "tool.invalid_arguments"


def test_output_path_receipt_and_job_dir_removed(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("PICO_OFFICE_TMP", str(tmp_path))
    source = """
import os, json
from openpyxl import Workbook
wb = Workbook()
ws = wb.active
ws["A1"] = "ok"
ws["B1"] = len(os.listdir("."))
print("stdout-marker")
wb.save(OUTPUT_PATH)
"""
    receipt = run_office_job(kind="xlsx", source=source)
    assert receipt["ok"] is True, receipt
    assert receipt["exit_code"] == 0
    assert "stdout-marker" in receipt["stdout_tail"]
    assert receipt["output_name"] == "out.xlsx"
    assert receipt["wall_ms"] >= 0
    import base64

    ws = load_workbook(BytesIO(base64.b64decode(receipt["output_b64"]))).active
    assert ws["A1"].value == "ok"
    # job dir is gone
    assert not any(p.name.startswith("job-") for p in tmp_path.iterdir())


def test_newest_matching_file_collected_and_input_excluded(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("PICO_OFFICE_TMP", str(tmp_path))
    from openpyxl import Workbook

    buf = BytesIO()
    wb = Workbook()
    wb.active["A1"] = "原件"
    wb.save(buf)
    source = """
from openpyxl import load_workbook
wb = load_workbook(INPUT_PATH)
wb.active["A2"] = "改了"
wb.save("renamed.xlsx")
"""
    receipt = run_office_job(
        kind="xlsx", source=source, input_bytes=buf.getvalue(), input_name="grades.xlsx"
    )
    assert receipt["ok"] is True, receipt
    assert receipt["output_name"] == "renamed.xlsx"
    import base64

    ws = load_workbook(BytesIO(base64.b64decode(receipt["output_b64"]))).active
    assert ws["A1"].value == "原件" and ws["A2"].value == "改了"


def test_failure_receipt_carries_stderr() -> None:
    receipt = run_office_job(kind="docx", source="raise SystemExit('nope-marker')")
    assert receipt["ok"] is False
    assert receipt["code"] == "sandbox.docx_failed"
    assert "nope-marker" in receipt["stderr_tail"]


def test_no_output_receipt() -> None:
    receipt = run_office_job(kind="pptx", source="x = 1")
    assert receipt["ok"] is False
    assert receipt["code"] == "sandbox.no_output"
    assert "save_deck" in receipt["message"]


def test_timeout_kills_and_reports() -> None:
    receipt = run_office_job(kind="docx", source="import time\ntime.sleep(30)", timeout_s=1.5)
    assert receipt["ok"] is False
    assert receipt["timed_out"] is True
    assert receipt["code"] == "sandbox.exec_timeout"
    assert receipt["wall_ms"] < 15_000


def test_timeout_is_capped() -> None:
    assert MAX_TIMEOUT_S == 180.0


def test_env_does_not_leak_host_secrets(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("PICO_SANDBOX_TOKEN", "sekrit-token-value")
    monkeypatch.setenv("DEEPSEEK_API_KEY", "sk-should-not-leak")
    source = """
import os
from docx import Document
doc = Document()
doc.add_paragraph("|".join(sorted(k for k in os.environ)))
doc.add_paragraph("LEAK" if any("sekrit" in v or "sk-should" in v for v in os.environ.values()) else "CLEAN")
save_doc(doc)
"""
    receipt = run_office_job(kind="docx", source=source)
    assert receipt["ok"] is True, receipt
    import base64

    from docx import Document

    paras = [p.text for p in Document(BytesIO(base64.b64decode(receipt["output_b64"]))).paragraphs]
    assert paras[1] == "CLEAN"
    assert "PICO_SANDBOX_TOKEN" not in paras[0]
    assert "DEEPSEEK_API_KEY" not in paras[0]


def test_http_surface_shape() -> None:
    paths = {r.path for r in app.routes}
    assert "/v1/internal/office/run" in paths
    assert "/health" in paths
    body = RunBody(kind="docx", source="x=1")
    assert body.images == {}
    assert body.timeout_s is None


@pytest.mark.skipif(os.name != "posix", reason="rlimit is posix-only")
def test_child_has_rlimits(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("PICO_OFFICE_RLIMIT_NPROC", "64")
    source = """
import resource
from docx import Document
soft, _ = resource.getrlimit(resource.RLIMIT_NPROC)
fsize, _ = resource.getrlimit(resource.RLIMIT_FSIZE)
doc = Document()
doc.add_paragraph(f"nproc={soft}")
doc.add_paragraph(f"fsize={fsize}")
save_doc(doc)
"""
    receipt = run_office_job(kind="docx", source=source)
    assert receipt["ok"] is True, receipt
    import base64

    from docx import Document

    paras = [p.text for p in Document(BytesIO(base64.b64decode(receipt["output_b64"]))).paragraphs]
    assert paras[0] == "nproc=64"
    assert paras[1] == f"fsize={200 * 1024 * 1024}"
