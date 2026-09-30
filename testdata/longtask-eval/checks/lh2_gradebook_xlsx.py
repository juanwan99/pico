"""LH2: grade-level analysis from 270 students x 12 marks — every number against truth.

Truth is recomputed from the attachment (/w/in/年级成绩.xlsx). Points (all
essential): 班级均分 36 rows · 进步榜 top 10 in order · 年级概况 per subject ·
Word report long enough and naming the right best/worst class-subject and No.1.
Formula-only cells are recalculated with LibreOffice before reading.
"""

import os
import re
import subprocess
import sys

sys.path.insert(0, "/w/check")
from _common import IN, OUT, Score

SUBJECTS = ["语文", "数学", "英语", "物理", "化学", "生物"]
TOL = 0.011


def num(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    text = str(v).strip().replace("％", "%")
    try:
        return float(text.rstrip("%")) / (100 if text.endswith("%") else 1)
    except ValueError:
        return None


def truth():
    from openpyxl import load_workbook

    wb = load_workbook(os.path.join(IN, "年级成绩.xlsx"), data_only=True)
    sheets = {}
    for title in ("期中", "期末"):
        rows = list(wb[title].iter_rows(values_only=True))
        head = [str(h).strip() for h in rows[0]]
        recs = {}
        for r in rows[1:]:
            d = dict(zip(head, r))
            recs[str(d["学号"]).strip()] = {
                "cls": str(d["班级"]).strip(),
                "name": str(d["姓名"]).strip(),
                "m": {s: num(d[s]) for s in SUBJECTS},
            }
        sheets[title] = recs
    mid, fin = sheets["期中"], sheets["期末"]
    classes = sorted({r["cls"] for r in mid.values()})
    means = {}
    for c in classes:
        for s in SUBJECTS:
            a = [r["m"][s] for r in mid.values() if r["cls"] == c and r["m"][s] is not None]
            b = [r["m"][s] for r in fin.values() if r["cls"] == c and r["m"][s] is not None]
            ma, mb = sum(a) / len(a), sum(b) / len(b)
            means[(c, s)] = (ma, mb, mb - ma)
    board = []
    for sid, r in mid.items():
        f = fin.get(sid)
        if not f or None in r["m"].values() or None in f["m"].values():
            continue
        t0, t1 = sum(r["m"].values()), sum(f["m"].values())
        board.append((-(t1 - t0), sid, r["cls"], r["name"], t0, t1))
    board.sort()
    overview = {}
    for s in SUBJECTS:
        v = [r["m"][s] for r in fin.values() if r["m"][s] is not None]
        overview[s] = (sum(v) / len(v), max(v), sum(1 for x in v if x >= 60) / len(v))
    return classes, means, board[:10], overview


def delivered(ext):
    found = []
    for dirpath, _dirs, files in os.walk(OUT):
        found += [os.path.join(dirpath, f) for f in files if f.lower().endswith(ext)]
    return found


def open_book(path):
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    formulas = load_workbook(path, data_only=False)
    stale = any(
        isinstance(c.value, str) and c.value.startswith("=") and wb[ws.title][c.coordinate].value is None
        for ws in formulas.worksheets
        for row in ws.iter_rows()
        for c in row
    )
    if stale:
        os.makedirs("recalc", exist_ok=True)
        subprocess.run(
            ["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", "recalc", path],
            capture_output=True, timeout=180, check=False, env={**os.environ, "HOME": os.path.abspath(".")},
        )
        again = os.path.join("recalc", os.path.basename(path))
        if os.path.exists(again):
            wb = load_workbook(again, data_only=True)
    return wb


def table(wb, key, must):
    """Rows under the first header row that holds every name in ``must``."""
    ws = next((w for w in wb.worksheets if key in w.title), None)
    if ws is None:
        return None, f"no sheet named like {key}"
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    for i, r in enumerate(rows):
        cells = [str(c).strip() if c is not None else "" for c in r]
        cols = {}
        for want in must:
            hit = next((j for j, c in enumerate(cells) if want in c), None)
            if hit is None:
                break
            cols[want] = hit
        else:
            body = [x for x in rows[i + 1:] if any(v not in (None, "") for v in x)]
            return [{k: x[j] if j < len(x) else None for k, j in cols.items()} for x in body], ""
    return None, f"{key}: header with {must} not found"


def main():
    sc = Score()
    books = [p for p in delivered(".xlsx") if "年级成绩" not in os.path.basename(p)]
    docs = delivered(".docx")
    if not books:
        sc.fail("no .xlsx delivered")
        sc.emit()
    _classes, means, board, overview = truth()
    wb = open_book(min(books, key=len))

    rows, why = table(wb, "班级均分", ["班级", "科目", "期中", "期末", "进步"])
    bad = [why] if why else []
    if rows is not None:
        got = {(str(r["班级"]).strip(), str(r["科目"]).strip()): r for r in rows}
        if len(rows) != 36 or set(got) != set(means):
            bad.append(f"班级均分 rows={len(rows)} keys off: {sorted(set(got) ^ set(means))[:3]}")
        for k, (ma, mb, d) in means.items():
            r = got.get(k)
            if r is None:
                continue
            vals = (num(r["期中"]), num(r["期末"]), num(r["进步"]))
            if any(v is None or abs(v - t) > TOL for v, t in zip(vals, (ma, mb, d))):
                bad.append(f"{k[0]}{k[1]} got {vals} want ({ma:.2f}, {mb:.2f}, {d:.2f})")
                break
    sc.check(not bad, "班级均分: " + " | ".join(bad[:3]))

    rows, why = table(wb, "进步榜", ["班级", "姓名", "期中", "期末", "进步"])
    bad = [why] if why else []
    if rows is not None:
        got = [(str(r["班级"]).strip(), str(r["姓名"]).strip(), num(r["进步"])) for r in rows[:10]]
        want = [(b[2], b[3], b[5] - b[4]) for b in board]
        if [g[:2] for g in got] != [w[:2] for w in want]:
            bad.append(f"top10 got {[g[1] for g in got]} want {[w[1] for w in want]}")
        elif any(g[2] is None or abs(g[2] - w[2]) > TOL for g, w in zip(got, want)):
            bad.append("进步 values off")
    sc.check(not bad, "进步榜: " + " | ".join(bad[:2]))

    rows, why = table(wb, "年级概况", ["科目", "均分", "最高", "及格率"])
    bad = [why] if why else []
    if rows is not None:
        got = {str(r["科目"]).strip(): r for r in rows}
        for s, (m, hi, rate) in overview.items():
            r = got.get(s)
            if r is None:
                bad.append(f"no row {s}")
                continue
            pr = num(r["及格率"])
            if pr is not None and pr > 1.0001:
                pr /= 100
            if (
                num(r["均分"]) is None or abs(num(r["均分"]) - m) > TOL
                or num(r["最高"]) != hi or pr is None or abs(pr - rate) > 0.0006
            ):
                bad.append(f"{s} got ({r['均分']}, {r['最高']}, {r['及格率']}) want ({m:.2f}, {hi}, {rate:.1%})")
    sc.check(not bad, "年级概况: " + " | ".join(bad[:2]))

    bad = []
    if not docs:
        bad.append("no .docx")
    else:
        from docx import Document

        d = Document(min(docs, key=len))
        text = "\n".join(p.text for p in d.paragraphs)
        text += "\n".join(c.text for t in d.tables for row in t.rows for c in row.cells)
        n = len(re.sub(r"\s", "", text))
        if n < 1500:
            bad.append(f"docx {n} chars < 1500")
        best = max(means, key=lambda k: means[k][2])
        worst = min(means, key=lambda k: means[k][2])
        flat = re.sub(r"\s", "", text)
        for tag, (c, s) in (("best", best), ("worst", worst)):
            if not re.search(re.escape(c) + r".{0,30}" + s + "|" + s + r".{0,30}" + re.escape(c), flat):
                bad.append(f"{tag} {c}{s} not named")
        if board[0][3] not in flat:
            bad.append(f"No.1 {board[0][3]} not named")
    sc.check(not bad, "报告: " + " | ".join(bad))
    sc.emit()


main()
