"""pico-office: run teacher-facing office Python inside the isolated container.

This is the office computer, not a jail. Isolation is the container contract
(``network_mode: none``, read-only root, tmpfs workdir, rlimit, no secrets).
The script gets a full Python 3 with python-docx / openpyxl / python-pptx and
the whole standard library. Pico does not maintain an import allowlist here.

Started as ``python -m sandbox_worker.office_runner --uds /run/pico-office/office.sock``
in the ``pico-office`` service, or called in-process (``run_office_job``) when
``PICO_OFFICE_URL=embedded`` (tests / dev).
"""

from __future__ import annotations

import asyncio
import base64
import contextlib
import json
import os
import shutil
import signal
import subprocess
import sys
import tempfile
import time
import uuid
from pathlib import Path
from typing import Any

from fastapi import FastAPI, Header, HTTPException
from pydantic import BaseModel, Field

OFFICE_KINDS = ("docx", "xlsx", "pptx")
DEFAULT_TIMEOUT_S = 60.0
MAX_TIMEOUT_S = 180.0
MAX_SOURCE_CHARS = 200_000
TAIL_CHARS = 4000
# Set by compose for the container (uid 65532 owns only this service). Left at 0
# elsewhere: on a shared host the uid already has more processes than any cap.
RLIMIT_NPROC = int(os.environ.get("PICO_OFFICE_RLIMIT_NPROC") or "0")
RLIMIT_FSIZE_BYTES = 200 * 1024 * 1024
_CONCURRENCY = int(os.environ.get("PICO_OFFICE_CONCURRENCY") or "2")

_PRELUDE = '''"""Convenience names for office scripts. Not a sandbox; full Python is available."""
import json as _json
import os as _os

KIND = _os.environ.get("PICO_KIND", "")
INPUT_PATH = _os.environ.get("PICO_INPUT_PATH", "")
OUTPUT_PATH = _os.environ.get("PICO_OUTPUT_PATH", "")
WORKDIR = _os.environ.get("PICO_JOB", _os.getcwd())

import pptx as _pptx_pkg
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx_helpers import (
    ImagePathMap,
    add_content_slide,
    add_table,
    add_title_slide,
    install_blank_slide_compat,
)

_pptx_pkg.Inches = Inches
_pptx_pkg.Pt = Pt
_pptx_pkg.Emu = Emu
_pptx_pkg.RGBColor = RGBColor
install_blank_slide_compat()

IMAGE_PATHS = ImagePathMap(_json.loads(_os.environ.get("PICO_IMAGE_PATHS") or "{}"))

# Receipt, not a redirect: remember where the script saved so the runner can
# collect it even when the path is not OUTPUT_PATH. `.save()` with no argument
# lands on OUTPUT_PATH as a convenience.
_SAVED = []


def _track_save(cls):
    orig = cls.save

    def save(self, path=None, *args, **kwargs):
        target = OUTPUT_PATH if path is None else path
        result = orig(self, target, *args, **kwargs)
        try:
            _SAVED.append(_os.path.abspath(_os.fspath(target)))
        except TypeError:
            pass
        return result

    cls.save = save


from pptx.presentation import Presentation as _PrsCls
from docx.document import Document as _DocCls
from openpyxl.workbook.workbook import Workbook as _WbCls

for _cls in (_PrsCls, _DocCls, _WbCls):
    _track_save(_cls)


def _flush_saved():
    try:
        with open(_os.path.join(WORKDIR, ".pico_saved.json"), "w", encoding="utf-8") as fh:
            _json.dump(_SAVED, fh)
    except OSError:
        pass


import atexit as _atexit

_atexit.register(_flush_saved)


def _need_input(label):
    if not INPUT_PATH:
        raise ValueError(
            f"没有原件。改已有{label}请传 artifact_id；新建不要调用 load_*。"
        )
    return INPUT_PATH


def load_doc():
    return Document(_need_input(" Word"))


def load_book():
    return load_workbook(_need_input(" Excel"))


def load_deck():
    return Presentation(_need_input(" PPT"))


def save_doc(doc):
    doc.save(OUTPUT_PATH)


def save_book(wb):
    wb.save(OUTPUT_PATH)


def save_deck(prs):
    prs.save(OUTPUT_PATH)


__all__ = [
    "KIND", "INPUT_PATH", "OUTPUT_PATH", "WORKDIR", "IMAGE_PATHS",
    "Presentation", "Document", "Workbook", "load_workbook",
    "Inches", "Pt", "Emu", "RGBColor",
    "add_title_slide", "add_content_slide", "add_table",
    "load_doc", "load_book", "load_deck", "save_doc", "save_book", "save_deck",
]
'''

_RUNNER = """import runpy
import pico_prelude

_names = {k: getattr(pico_prelude, k) for k in pico_prelude.__all__}
runpy.run_path("user_script.py", init_globals=_names, run_name="__main__")
"""


def _helpers_source() -> str:
    from pico_orchestrator.office import pptx_helpers

    return Path(pptx_helpers.__file__).read_text(encoding="utf-8")


def _tail(raw: bytes | None) -> str:
    text = (raw or b"").decode("utf-8", errors="replace")
    return text[-TAIL_CHARS:]


def _limits() -> None:
    """Child-only rlimits. Memory is the container cgroup's job."""
    with contextlib.suppress(Exception):  # best-effort on non-Linux hosts
        import resource

        nproc = int(os.environ.get("PICO_OFFICE_RLIMIT_NPROC") or RLIMIT_NPROC or 0)
        if nproc > 0:
            resource.setrlimit(resource.RLIMIT_NPROC, (nproc, nproc))
        resource.setrlimit(resource.RLIMIT_FSIZE, (RLIMIT_FSIZE_BYTES, RLIMIT_FSIZE_BYTES))


def _sniff_image_ext(blob: bytes) -> str:
    if blob[:8] == b"\x89PNG\r\n\x1a\n":
        return ".png"
    if blob[:3] == b"\xff\xd8\xff":
        return ".jpg"
    if blob[:4] == b"GIF8":
        return ".gif"
    return ".png"


def _safe_ext(name: str, default: str) -> str:
    suffix = Path(str(name or "")).suffix.lower()
    if suffix and len(suffix) <= 8 and suffix[1:].isalnum():
        return suffix
    return default


def _collect(job: Path, *, kind: str, out_path: Path, in_path: Path | None) -> Path | None:
    if out_path.is_file() and out_path.stat().st_size > 0:
        return out_path
    manifest = job / ".pico_saved.json"
    if manifest.is_file():
        try:
            saved = json.loads(manifest.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            saved = []
        for raw in reversed([s for s in saved if isinstance(s, str)]):
            p = Path(raw)
            if (
                p.suffix.lower() == f".{kind}"
                and p.is_file()
                and p.stat().st_size > 0
                and (in_path is None or p.resolve() != in_path.resolve())
            ):
                return p
    candidates = [
        p
        for p in job.rglob(f"*.{kind}")
        if p.is_file()
        and p.stat().st_size > 0
        and (in_path is None or p.resolve() != in_path.resolve())
    ]
    if not candidates:
        return None
    return max(candidates, key=lambda p: p.stat().st_mtime_ns)


def run_office_job(
    *,
    kind: str,
    source: str,
    input_bytes: bytes | None = None,
    input_name: str | None = None,
    images: dict[str, bytes] | None = None,
    timeout_s: float | None = None,
) -> dict[str, Any]:
    """Run one office script in a throwaway job dir. Returns a JSON-safe receipt."""
    kind = str(kind or "").strip().lower()
    if kind not in OFFICE_KINDS:
        return {
            "ok": False,
            "code": "tool.invalid_arguments",
            "message": "kind 必须是 docx、xlsx 或 pptx。",
        }
    text = str(source or "")
    if not text.strip():
        return {
            "ok": False,
            "code": "tool.invalid_arguments",
            "message": "source 必须是非空的办公脚本。",
        }
    if len(text) > MAX_SOURCE_CHARS:
        return {
            "ok": False,
            "code": "tool.invalid_arguments",
            "message": f"source 不能超过 {MAX_SOURCE_CHARS} 字。",
        }
    timeout = float(timeout_s or DEFAULT_TIMEOUT_S)
    timeout = max(1.0, min(timeout, MAX_TIMEOUT_S))

    root = Path(os.environ.get("PICO_OFFICE_TMP") or tempfile.gettempdir())
    root.mkdir(parents=True, exist_ok=True)
    job = Path(tempfile.mkdtemp(prefix=f"job-{uuid.uuid4().hex[:8]}-", dir=str(root)))
    started = time.perf_counter()
    try:
        (job / "pico_prelude.py").write_text(_PRELUDE, encoding="utf-8")
        (job / "pptx_helpers.py").write_text(_helpers_source(), encoding="utf-8")
        (job / "user_script.py").write_text(text, encoding="utf-8")
        (job / "run.py").write_text(_RUNNER, encoding="utf-8")

        out_path = job / f"out.{kind}"
        in_path: Path | None = None
        if input_bytes:
            in_path = job / f"in{_safe_ext(input_name or '', f'.{kind}')}"
            in_path.write_bytes(bytes(input_bytes))

        image_paths: dict[str, str] = {}
        for idx, (key, blob) in enumerate((images or {}).items()):
            if not blob:
                continue
            dest = job / f"img_{idx}{_sniff_image_ext(blob)}"
            dest.write_bytes(blob)
            image_paths[str(key)] = str(dest)

        env = {
            "PATH": os.environ.get("PATH", "/usr/local/bin:/usr/bin:/bin"),
            "HOME": str(job),
            "LANG": os.environ.get("LANG", "C.UTF-8"),
            "LC_ALL": os.environ.get("LC_ALL", "C.UTF-8"),
            "PYTHONDONTWRITEBYTECODE": "1",
            "PYTHONUNBUFFERED": "1",
            "PYTHONIOENCODING": "utf-8",
            "PICO_JOB": str(job),
            "PICO_KIND": kind,
            "PICO_INPUT_PATH": str(in_path) if in_path else "",
            "PICO_OUTPUT_PATH": str(out_path),
            "PICO_IMAGE_PATHS": json.dumps(image_paths, ensure_ascii=False),
        }
        for key in ("VIRTUAL_ENV", "PYTHONPATH", "SYSTEMROOT", "TEMP", "TMP"):
            if key in os.environ and key not in env:
                env[key] = os.environ[key]

        popen_kwargs: dict[str, Any] = {
            "cwd": str(job),
            "env": env,
            "stdout": subprocess.PIPE,
            "stderr": subprocess.PIPE,
            "stdin": subprocess.DEVNULL,
        }
        if os.name == "posix":
            popen_kwargs["start_new_session"] = True
            popen_kwargs["preexec_fn"] = _limits
        proc = subprocess.Popen([sys.executable, "-B", "run.py"], **popen_kwargs)
        timed_out = False
        try:
            stdout, stderr = proc.communicate(timeout=timeout)
        except subprocess.TimeoutExpired:
            timed_out = True
            if os.name == "posix":
                try:
                    os.killpg(proc.pid, signal.SIGKILL)
                except ProcessLookupError:
                    pass
            else:
                proc.kill()
            stdout, stderr = proc.communicate()
        wall_ms = int((time.perf_counter() - started) * 1000)

        files = [
            {"name": str(p.relative_to(job)), "bytes": p.stat().st_size}
            for p in sorted(job.rglob("*"))
            if p.is_file()
            and p.name
            not in {
                "pico_prelude.py",
                "pptx_helpers.py",
                "run.py",
                "user_script.py",
                ".pico_saved.json",
            }
        ]
        receipt: dict[str, Any] = {
            "ok": False,
            "kind": kind,
            "exit_code": proc.returncode,
            "timed_out": timed_out,
            "stdout_tail": _tail(stdout),
            "stderr_tail": _tail(stderr),
            "wall_ms": wall_ms,
            "files": files,
            "output_b64": None,
            "output_name": None,
        }
        if timed_out:
            receipt.update(
                code="sandbox.exec_timeout", message=f"办公脚本超过 {int(timeout)} 秒，已终止。"
            )
            return receipt
        if proc.returncode != 0:
            receipt.update(
                code=f"sandbox.{kind}_failed",
                message="办公脚本运行失败，见 stderr。",
            )
            return receipt
        produced = _collect(job, kind=kind, out_path=out_path, in_path=in_path)
        if produced is None:
            receipt.update(
                code="sandbox.no_output",
                message=f"脚本跑完了，但没有写出 .{kind} 文件。请调用 save_{'doc' if kind == 'docx' else 'book' if kind == 'xlsx' else 'deck'}(...) 或 .save(OUTPUT_PATH)。",
            )
            return receipt
        receipt.update(
            ok=True,
            output_b64=base64.b64encode(produced.read_bytes()).decode("ascii"),
            output_name=produced.name,
        )
        return receipt
    finally:
        shutil.rmtree(job, ignore_errors=True)


# --- HTTP surface (pico-office service) ---------------------------------------

app = FastAPI(title="pico-office", version="0.1.0")
_slots = asyncio.Semaphore(max(1, _CONCURRENCY))


def _token_ok(header: str | None) -> bool:
    expected = (os.environ.get("PICO_SANDBOX_TOKEN") or "").strip()
    if not expected:
        return True
    return (header or "").strip() == expected


class RunBody(BaseModel):
    kind: str
    source: str
    input_name: str | None = None
    input_b64: str | None = None
    images: dict[str, str] = Field(default_factory=dict)
    timeout_s: float | None = None


@app.get("/health")
async def health() -> dict[str, Any]:
    return {
        "ok": True,
        "service": "pico-office",
        "python": sys.version.split()[0],
        "kinds": list(OFFICE_KINDS),
        "network": "none",
        "jail": False,
        "claim_wb": "NO",
    }


@app.post("/v1/internal/office/run")
async def run_office(
    body: RunBody,
    x_pico_sandbox_token: str | None = Header(default=None),
) -> dict[str, Any]:
    if not _token_ok(x_pico_sandbox_token):
        raise HTTPException(
            status_code=401, detail={"code": "sandbox.denied", "message": "sidecar token mismatch"}
        )
    try:
        input_bytes = base64.b64decode(body.input_b64, validate=False) if body.input_b64 else None
        images = {
            k: base64.b64decode(v, validate=False) for k, v in (body.images or {}).items() if v
        }
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail={
                "code": "tool.invalid_arguments",
                "message": "input_b64 / images 不是有效 base64",
            },
        ) from exc
    async with _slots:
        return await asyncio.to_thread(
            run_office_job,
            kind=body.kind,
            source=body.source,
            input_bytes=input_bytes,
            input_name=body.input_name,
            images=images,
            timeout_s=body.timeout_s,
        )


def main() -> None:
    import argparse

    import uvicorn

    parser = argparse.ArgumentParser(description="pico-office runner")
    parser.add_argument("--uds", default=os.environ.get("PICO_OFFICE_UDS") or "")
    parser.add_argument("--host", default="127.0.0.1")
    parser.add_argument(
        "--port", type=int, default=int(os.environ.get("PICO_OFFICE_PORT") or "18769")
    )
    args = parser.parse_args()
    if args.uds:
        sock = Path(args.uds)
        parent = sock.parent
        parent.mkdir(parents=True, exist_ok=True)
        if not os.access(parent, os.W_OK):
            print(
                f"FATAL: {parent} is not writable by uid {os.getuid()}. "
                "Office sock dir must be a host bind with mode 1777, not a named volume.",
                file=sys.stderr,
            )
            raise SystemExit(2)
        try:
            sock.unlink()
        except FileNotFoundError:
            pass
        uvicorn.run("sandbox_worker.office_runner:app", uds=args.uds, factory=False)
    else:
        uvicorn.run(
            "sandbox_worker.office_runner:app", host=args.host, port=args.port, factory=False
        )


if __name__ == "__main__":
    main()
