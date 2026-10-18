"""Grading rules of the Pico ruler (pico#1323): scripts/pico-ruler.py."""

from __future__ import annotations

import importlib.util
import io
import json
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
_spec = importlib.util.spec_from_file_location("pico_ruler", ROOT / "scripts" / "pico-ruler.py")
assert _spec and _spec.loader
ruler = importlib.util.module_from_spec(_spec)
sys.modules["pico_ruler"] = ruler
_spec.loader.exec_module(ruler)

docx = pytest.importorskip("docx")
pptx = pytest.importorskip("pptx")
openpyxl = pytest.importorskip("openpyxl")


def _letter(size_pt: float, east_asia: str = "宋体") -> bytes:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = docx.Document()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("致高三家长的一封信").bold = True
    for text in ("尊敬的各位家长：寒假将至，请关注孩子的安全。", "寒假期间请督促孩子按计划复习。"):
        run = doc.add_paragraph().add_run(text)
        run.font.size = Pt(size_pt)
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), east_asia)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _checks(results):
    return {name.split("（")[0]: ok for name, ok in results}


def test_docx_title_and_body_font_pass():
    out = ruler.docx_checks(
        _letter(12), {"docx_title_center_bold": True, "docx_body_font": {"size_pt": 12, "east_asia": ["宋体"]}}
    )
    assert all(ok for _, ok in out), out


def test_docx_body_font_wrong_size_or_face_fails():
    assert not ruler.docx_checks(_letter(14), {"docx_body_font": {"size_pt": 12, "east_asia": ["宋体"]}})[0][1]
    assert not ruler.docx_checks(_letter(12, "黑体"), {"docx_body_font": {"size_pt": 12, "east_asia": ["宋体"]}})[0][1]


def test_docx_body_font_inherited_from_normal_style():
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = docx.Document()
    normal = doc.styles["Normal"]
    normal.font.size = Pt(12)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    doc.add_paragraph("标题")
    doc.add_paragraph("正文第一段，字号和字体都从正文样式继承。")
    buf = io.BytesIO()
    doc.save(buf)
    assert ruler.docx_checks(buf.getvalue(), {"docx_body_font": {"size_pt": 12, "east_asia": ["宋体"]}})[0][1]


def test_docx_numbered_questions_and_answers_after():
    doc = docx.Document()
    doc.add_paragraph("初二数学周测")
    for i in range(1, 13):
        doc.add_paragraph(f"{i}. 一次函数 y=kx+b 的第 {i} 题")
    doc.add_paragraph("参考答案")
    doc.add_paragraph("1. A")
    buf = io.BytesIO()
    doc.save(buf)
    got = _checks(ruler.docx_checks(buf.getvalue(), {"min_numbered": 10, "docx_answer_after_questions": True, "text_has": ["答案"]}))
    assert got == {"编号题≥10": True, "答案在题目之后": True, "含 答案": True}


def test_docx_answer_heading_inside_line_and_inline_answers():
    def paper(lines):
        doc = docx.Document()
        for line in lines:
            doc.add_paragraph(line)
        buf = io.BytesIO()
        doc.save(buf)
        return ruler.docx_checks(buf.getvalue(), {"docx_answer_after_questions": True})[0][1]

    qs = [f"{i}．一次函数第 {i} 题" for i in range(1, 13)]
    assert paper(["一次函数周测试卷及答案解析", *qs, "八年级《一次函数》周测试卷 · 参考答案与详细解析", "1.C"])
    inline = [x for i, q in enumerate(qs) for x in (q, f"答案：{'ABCD'[i % 4]}")]
    assert not paper(["一次函数周测卷", *inline])


def test_docx_text_absent_and_tables():
    doc = docx.Document()
    doc.add_paragraph("地点：学校报告厅。李老师")
    table = doc.add_table(rows=16, cols=3)
    table.cell(0, 0).text = "姓名"
    buf = io.BytesIO()
    doc.save(buf)
    got = _checks(ruler.docx_checks(buf.getvalue(), {"text_absent": ["三楼会议室"], "docx_min_table_rows": 16, "docx_min_images": 1}))
    assert got == {"不含 三楼会议室": True, "表格≥16行": True, "图≥1": False}


def test_pptx_titles_and_restyle_against_fixture():
    raw = (ruler.CASES_DIR / "fixtures" / "r07_deck.pptx").read_bytes()
    got = _checks(ruler.pptx_checks(raw, {"slides_min": 7, "slide_titles_has": ["本周回顾"], "pptx_restyled_from": "fixtures/r07_deck.pptx"}))
    assert got == {"页数≥7": False, "页标题含 本周回顾": True, "换了模板": False}


def test_pptx_text_box_deck_with_header_chrome_and_painted_background():
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    prs = pptx.Presentation()
    blank = prs.slide_layouts[6]
    titles = ["封面", "本周回顾", "作业情况", "下周安排", "期中考试安排", "家长配合事项", "谢谢各位家长"]
    for i, title in enumerate(titles, 1):
        slide = prs.slides.add_slide(blank)
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = RGBColor(0xF5, 0xF7, 0xFA)
        for text in ("高三（2）班家长会", title, f"第 {i} 页 / 共 7 页"):
            slide.shapes.add_textbox(Inches(1), Inches(i % 3 + 1), Inches(6), Inches(1)).text_frame.text = text
    buf = io.BytesIO()
    prs.save(buf)
    got = _checks(
        ruler.pptx_checks(
            buf.getvalue(),
            {"slide_titles_has": titles[1:6], "slide_titles_min_distinct": 7, "pptx_restyled_from": "fixtures/r07_deck.pptx"},
        )
    )
    assert got == {"页标题含 本周回顾/作业情况/下周安排/期中考试安排/家长配合事项": True, "不同标题≥7": True, "换了模板": True}, got


def test_xlsx_chart_values_header_and_phones():
    from openpyxl.chart import BarChart, Reference

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["姓名", "班级", "电话"])
    for i in range(12):
        ws.append([f"学生{i}", "高一（3）班", f"139123456{i:02d}"])
    ws.append([])
    ws["E1"] = "=AVERAGE(B2:B13)"
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=13))
    ws.add_chart(chart, "G2")
    buf = io.BytesIO()
    wb.save(buf)
    got = _checks(
        ruler.xlsx_checks(
            buf.getvalue(),
            {
                "xlsx_chart": "bar",
                "xlsx_values_or_formula": {"values": [1.0], "tol": 0.01, "formula": "AVERAGE"},
                "xlsx_header_has": ["姓名", "班级", "电话"],
                "xlsx_data_rows": 12,
                "xlsx_phones_11": 12,
            },
        )
    )
    assert all(got.values()), got


def test_reply_checks():
    assert ruler.reply_check({"type": "reply_facts", "items": [["6:30"], ["11月6日"]]}, "早餐 6：30 开始；11月6号考试")[1]
    assert not ruler.reply_check({"type": "reply_facts", "items": [["21:40"]]}, "晚自习调整")[1]
    assert ruler.reply_check({"type": "reply_len", "min": 5, "max": 10}, "## 家长 会\n- 感谢")[1]
    assert ruler.reply_check({"type": "reply_numbered_min", "n": 5}, "\n".join(f"{i}. x²-{i}=0" for i in range(1, 6)))[1]
    assert ruler.reply_check({"type": "reply_links_min", "n": 1}, "出处：https://www.moe.gov.cn/x.html")[1]
    assert ruler.reply_check({"type": "reply_regex", "pattern": "每周[^。\\n]{0,12}(1|一)\\s*课时"}, "劳动课每周不少于1课时")[1]


def test_delivery_claim_and_question():
    assert ruler.DELIVERY_CLAIM.search("已为你生成 Word 文档，请下载")
    assert not ruler.DELIVERY_CLAIM.search("下面是发言稿正文")
    assert not ruler.DELIVERY_CLAIM.search("页面提示：验证码已发送，请查收短信")
    assert not ruler.DELIVERY_CLAIM.search("已成功登录！当前已进入控制台首页")
    assert ruler.DELIVERY_CLAIM.search("家长信.docx 在这里")
    assert ruler.DELIVERY_CLAIM.search("Word 文档已经做好了")
    assert not ruler.DELIVERY_CLAIM.search("需要直接导出为 **Word (.docx)** 打印排版文件，随时告诉我")
    assert ruler.looks_like_question("请告诉我开会的时间和地点？")
    assert not ruler.looks_like_question("通知已写好。")


def test_opens_rejects_broken_office_file():
    assert ruler.opens("docx", _letter(12))
    assert not ruler.opens("docx", b"PK\x03\x04 not really")


def _report(oks):
    results = [
        {"case": c, "title": c, "ok": ok, "wall_s": 10.0, "tool_calls": 2, "input_tokens": 100, "output_tokens": 10,
         "first_input_tokens": 100, "fake_delivery": [], "asks": 0, "passed": 1, "checks": 1, "failed": [], "fail_reason": ""}
        for c, ok in oks
    ]
    return {"stack": "t", "label": "", "results": results, "summary": ruler.summarize(results)}


def test_summary_and_compare_flags_flips():
    a = _report([("R00", True), ("R01", True), ("R02", False)])
    b = _report([("R00", True), ("R01", False), ("R02", True)])
    assert a["summary"]["cases"] == 2 and a["summary"]["ok"] == 1 and a["summary"]["probe_input_tokens"] == 100
    text = ruler.compare(a, b)
    assert "R01 ✅→❌" in text and "R02 ❌→✅" in text


KNOWN_TOP = {
    "reply_len", "reply_has", "reply_any", "reply_regex", "reply_links_min", "reply_numbered_min", "reply_facts",
    "tool_calls_max", "asks_max", "page_hit", "ask_or_placeholder", "files_named", "file",
}
KNOWN_SUB = {
    "docx": {"min_chars", "text_has", "text_any", "text_absent", "min_headings", "min_numbered", "docx_title_center_bold",
             "docx_body_font", "docx_min_images", "docx_min_table_rows", "docx_answer_after_questions"},
    "pptx": {"slides_min", "slides_max", "slide_titles_min_distinct", "slide_titles_has", "text_any", "pptx_restyled_from"},
    "xlsx": {"xlsx_chart", "xlsx_values_or_formula", "xlsx_header_has", "xlsx_data_rows", "xlsx_phones_11"},
}


def test_case_files_are_complete():
    cases = ruler.load_cases("")
    assert [c["id"] for c in cases] == [f"R{i:02d}" for i in range(21)]
    for case in cases:
        assert case["title"] and case["how"] and case["turns"] and case["checks"], case["id"]
        for turn in case["turns"]:
            for att in turn.get("attachments") or []:
                assert (ruler.CASES_DIR / att["file"]).is_file(), att
            if turn.get("paste_file"):
                assert (ruler.CASES_DIR / turn["paste_file"]).is_file()
        for check in case["checks"]:
            assert check["type"] in KNOWN_TOP, (case["id"], check["type"])
            if check["type"] == "file":
                keys = {k for d in check["sub"] for k in d}
                assert keys <= KNOWN_SUB[check["kind"]], (case["id"], keys - KNOWN_SUB[check["kind"]])
    assert (ruler.CASES_DIR / "www" / "login.html").read_text(encoding="utf-8").count("ok.txt") == 2
    json.dumps(cases, ensure_ascii=False)


def test_reconcile_flags_missing_pass_and_low_memory():
    cases = ruler.load_cases("")
    full = [("R00", True)] + [(f"R{i:02d}", i != 7) for i in range(1, 21)]
    a, b = _report(full), _report(full)
    for rep in (a, b):
        rep["revision"] = "638dd663"
        rep["env"] = {"edu_start": "cf43f706117c", "edu_end": "b6b610b7040a", "min_avail_mb": 23000}
    ok, text = ruler.reconcile(cases, a, b)
    assert ok, text
    b["env"]["min_avail_mb"] = 6000
    b["results"] = b["results"][:10]
    ok, text = ruler.reconcile(cases, a, b)
    assert not ok and "❌ 基线第 2 遍跑完" in text and "❌ ECS 可用内存" in text


def _deck(n: int) -> bytes:
    p = pptx.Presentation()
    for _ in range(n):
        p.slides.add_slide(p.slide_layouts[1]).shapes.title.text = "页"
    buf = io.BytesIO()
    p.save(buf)
    return buf.getvalue()


def test_said_vs_files_names_and_counts():
    # pico#1334 ③: ruler s13 R16 turn 2 — renamed in the reply, never written.
    reply = "已修改，文件已保存到 `/workspace/outputs/高三（2）班班级公约.docx`。"
    assert ruler.said_vs_files(reply, [], {"班级公约.docx"}) == ["提到的 高三（2）班班级公约.docx 不存在"]
    assert ruler.said_vs_files("文件已保存到 `班级公约.docx`。", [], {"班级公约.docx"}) == []
    deck = [("课件.pptx", "pptx", _deck(6))]
    assert ruler.said_vs_files("已生成 10 张幻灯片。", deck, {"课件.pptx"}) == ["课件.pptx 有 6 张，回复说 10 张"]
    assert ruler.said_vs_files("已生成六页 PPT。", deck, {"课件.pptx"}) == []
    pdf = b"%PDF-1.4\n1 0 obj<</Type /Pages /Kids[2 0 R 3 0 R]>>\n2 0 obj<</Type /Page>>\n3 0 obj<</Type/Page>>"
    assert ruler.said_vs_files("已导出 3 页 PDF。", [("卷.pdf", "pdf", pdf)], {"卷.pdf"}) == ["卷.pdf 有 2 页，回复说 3 页"]


def test_replay_export_keeps_only_gate_failed_turns(tmp_path):
    import sqlite3

    db = tmp_path / "p.db"
    con = sqlite3.connect(db)
    con.executescript(
        "CREATE TABLE tasks(id, school_id, conversation_id);"
        "CREATE TABLE runs(id, task_id, created_at, model, prompt);"
        "CREATE TABLE events(run_id, type, payload_json);"
        "INSERT INTO tasks VALUES('t1','s-1',NULL),('t2','s-1','c-2'),('t3','smoke-school','c-3');"
        "INSERT INTO runs VALUES('r1aaaaaaaa','t1','2026-10-07 15:32','pico-fast','细则'),"
        "('r2bbbbbbbb','t2','2026-10-08 07:06','pico-fast','做通知'),('r3cccccccc','t3','2026-10-17','pico-fast','x'),"
        "('r4dddddddd','t1','2026-10-07','pico-fast','ok');"
        "INSERT INTO events VALUES('r1aaaaaaaa','run.status','{\"reason\":\"deliverable_missing_artifact\"}'),"
        "('r2bbbbbbbb','run.status','{\"reason\":\"deliverable_missing_artifact\"}'),"
        "('r3cccccccc','run.status','{\"reason\":\"deliverable_missing_artifact\"}'),"
        "('r4dddddddd','run.status','{\"status\":\"succeeded\"}');"
    )
    con.commit()
    out = tmp_path / "replay"
    assert ruler.replay_export(str(db), str(out), "2026-10-03") == 0
    items = {p.stem: json.loads(p.read_text(encoding="utf-8")) for p in out.glob("*.json")}
    assert set(items) == {"r1aaaaaa", "r2bbbbbb"}
    assert items["r1aaaaaa"]["kind"] == "server" and items["r2bbbbbb"]["kind"] == "chat"
