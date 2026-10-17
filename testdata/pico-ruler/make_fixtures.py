#!/usr/bin/env python3
"""Build the pico-ruler attachments (pico#1323). Deterministic: same bytes every run
except docx/pptx/xlsx zip timestamps, which the grader never reads.

  python3 testdata/pico-ruler/make_fixtures.py          # writes fixtures/
  python3 testdata/pico-ruler/make_fixtures.py --check  # prints expected numbers

The PDF (r11_notice.pdf) is converted from r11_notice.docx with LibreOffice on a
host that has it (the pico-sandbox image), see README.md.
"""

from __future__ import annotations

import json
import random
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
OUT = HERE / "fixtures"
FONT = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"

SURNAMES = "王李张刘陈杨赵黄周吴徐孙胡朱高林何郭马罗梁宋郑谢韩唐冯于董萧程曹袁邓许傅沈曾彭吕苏卢蒋蔡贾丁魏"
GIVEN = "子涵一诺欣怡梓轩浩然雨桐宇轩思琪俊杰诗涵博文佳怡天佑紫萱明轩雅静晨曦嘉怡若曦宇航"


def students(n: int = 45) -> list[tuple[str, int, int, int]]:
    rnd = random.Random(1323)
    rows = []
    seen: set[str] = set()
    while len(rows) < n:
        name = rnd.choice(SURNAMES) + rnd.choice(GIVEN) + rnd.choice(GIVEN)
        if name in seen:
            continue
        seen.add(name)
        rows.append((name, rnd.randint(68, 138), rnd.randint(52, 148), rnd.randint(60, 142)))
    return rows


def averages() -> dict[str, float]:
    rows = students()
    return {
        subject: round(sum(r[i] for r in rows) / len(rows), 2)
        for i, subject in ((1, "语文"), (2, "数学"), (3, "英语"))
    }


def scores_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "期中成绩"
    ws.append(["姓名", "语文", "数学", "英语"])
    for row in students():
        ws.append(list(row))
    wb.save(path)


ROSTER = [
    ("陈思琪", "高一（3）班", "13912345601", "物理、化学、生物"),
    ("李浩然", "高一（3）班", "13912345602", "物理、化学、地理"),
    ("王雨桐", "高一（3）班", "13912345603", "历史、政治、地理"),
    ("张宇航", "高一（3）班", "13912345604", "物理、化学、生物"),
    ("刘佳怡", "高一（3）班", "13912345605", "历史、政治、生物"),
    ("赵俊杰", "高一（3）班", "13912345606", "物理、生物、地理"),
    ("周诗涵", "高一（3）班", "13912345607", "历史、化学、生物"),
    ("吴博文", "高一（3）班", "13912345608", "物理、化学、政治"),
    ("孙紫萱", "高一（3）班", "13912345609", "历史、政治、地理"),
    ("黄明轩", "高一（3）班", "13912345610", "物理、化学、生物"),
    ("林雅静", "高一（3）班", "13912345611", "历史、地理、生物"),
    ("何晨曦", "高一（3）班", "13912345612", "物理、化学、地理"),
    ("郭天佑", "高一（3）班", "13912345613", "物理、政治、地理"),
    ("马若曦", "高一（3）班", "13912345614", "历史、政治、化学"),
    ("罗嘉怡", "高一（3）班", "13912345615", "物理、化学、生物"),
]


def roster_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "选科名单"
    ws.append(["姓名", "班级", "家长电话", "选考科目"])
    for row in ROSTER:
        ws.append(list(row))
    wb.save(path)


def letter_docx(path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph("家长会通知")
    doc.add_paragraph("尊敬的各位家长：")
    doc.add_paragraph("您好！为加强家校沟通，经学校研究，定于2025年9月3日（星期三）下午3:00召开高三家长会。")
    doc.add_paragraph("地点：学校三楼会议室。")
    doc.add_paragraph("会议内容：介绍本学期教学安排、通报月考情况、交流备考经验。")
    doc.add_paragraph("请您准时参加，如有特殊情况请提前与班主任联系。")
    doc.add_paragraph("高三（2）班 王老师")
    doc.save(path)


def deck_pptx(path: Path) -> None:
    from pptx import Presentation

    prs = Presentation()
    for title, body in (
        ("高三（2）班家长会", "2026 年 10 月"),
        ("本周回顾", "月考结束，整体平稳"),
        ("作业情况", "按时提交率 92%"),
        ("下周安排", "专题复习：函数与导数"),
        ("谢谢各位家长", "欢迎随时联系"),
    ):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    prs.save(path)


NOTICE_LINES = [
    "关于2026年秋季学期期中考试安排的通知",
    "各年级、各班级：",
    "根据学校教学计划，本学期期中考试定于11月6日至11月7日举行。",
    "考试科目为语文、数学、英语、物理、化学、生物、政治、历史、地理。",
    "考场安排将于11月4日在教学楼一楼公告栏张贴，请各班主任提前组织学生核对考号。",
    "考试期间严禁携带手机等电子设备进入考场，违者按考试违纪处理。",
    "阅卷工作于11月8日在实验楼三楼集中进行，成绩预计11月12日发布。",
    "教务处",
    "2026年10月20日",
]


def notice_docx(path: Path) -> None:
    from docx import Document

    doc = Document()
    for line in NOTICE_LINES:
        doc.add_paragraph(line)
    doc.save(path)


def questions_png(path: Path) -> None:
    from PIL import Image, ImageDraw, ImageFont

    img = Image.new("RGB", (1200, 520), "white")
    draw = ImageDraw.Draw(img)
    font = ImageFont.truetype(FONT, 34)
    lines = [
        "期中复习 · 一元二次方程",
        "1. 解方程：x² - 5x + 6 = 0",
        "2. 解方程：2x² + 3x - 2 = 0",
        "3. 已知方程 x² + kx + 9 = 0 有两个相等的实数根，求 k 的值。",
    ]
    for i, line in enumerate(lines):
        draw.text((60, 60 + i * 100), line, fill="black", font=font)
    img.save(path, optimize=True)


FACTS = [
    "期中考试定于11月6日至7日进行，9个科目全部参加。",
    "家长开放日安排在11月15日上午，家长可以进班听课。",
    "食堂从下周一起增设早餐窗口，营业时间6:30到7:20。",
    "高三晚自习结束时间由22:00调整为21:40。",
    "学校将在11月20日组织消防疏散演练，各班提前熟悉路线。",
]
FILLER = [
    "会议首先回顾了上个月的教学工作，各年级组长依次发言，介绍了本组的进度和遇到的问题。",
    "有老师提出，部分学生作业完成质量不高，需要在作业设计上多花心思，减少机械重复。",
    "德育处通报了近期学生行为规范检查的情况，总体良好，个别班级午休纪律需要加强。",
    "后勤处介绍了教室设备的维护情况，几间教室的投影仪已经更换，空调检修将在周末完成。",
    "教研组长建议各学科加强集体备课，统一进度，统一练习，避免各班差距过大。",
    "大家还讨论了家校沟通的方式，认为班级群消息要简明，重要事项要单独提醒。",
    "心理老师提醒，考试前后学生压力较大，班主任要多关注情绪变化，必要时及时转介。",
    "图书馆本学期新到图书一千余册，欢迎老师推荐给学生，借阅时间延长到下午五点半。",
    "信息中心提醒，教师账号密码要定期修改，不要在公共电脑上保存登录状态。",
    "体育组汇报了秋季运动会的筹备进展，项目设置和往年基本一致，增加了趣味接力。",
]


def long_paste(path: Path) -> None:
    rnd = random.Random(2026)
    paras: list[str] = ["【10月16日教职工例会记录（整理稿）】"]
    facts = list(FACTS)
    while sum(len(p) for p in paras) < 8000:
        paras.append("".join(rnd.choice(FILLER) for _ in range(3)))
        if facts and rnd.random() < 0.18:
            paras.append(facts.pop(0))
    paras.extend(facts)
    path.write_text("\n\n".join(paras) + "\n", encoding="utf-8")


def main() -> int:
    if "--check" in sys.argv:
        print(json.dumps({"averages": averages(), "facts": FACTS}, ensure_ascii=False))
        return 0
    OUT.mkdir(exist_ok=True)
    scores_xlsx(OUT / "r08_scores.xlsx")
    letter_docx(OUT / "r05_letter.docx")
    deck_pptx(OUT / "r07_deck.pptx")
    roster_xlsx(OUT / "r10_roster.xlsx")
    notice_docx(OUT / "r11_notice.docx")
    questions_png(OUT / "r12_questions.png")
    long_paste(OUT / "r17_meeting.txt")
    print(json.dumps({"averages": averages(), "files": sorted(p.name for p in OUT.iterdir())}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
