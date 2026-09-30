"""LC5: quiz module after three turns — API, exclude + history, CLI to Word."""

import json
import os
import shutil
import sys
from pathlib import Path

sys.path.insert(0, "/w/check")
from _common import IN, Score, find, run, tail, workdir

API = r"""
import json, sys
from quiz import make_paper
from history import record, seen_ids
bank = json.load(open("bank.json", encoding="utf-8"))
ratio = {"易": 0.5, "中": 0.3, "难": 0.2}
ids = lambda p: [q["id"] for q in p]
lv = {q["id"]: q["level"] for q in bank}
out = {}
p1 = make_paper(bank, 10, ratio, seed=1)
out["det"] = ids(p1) == ids(make_paper(bank, 10, ratio, seed=1))
out["n"] = len(p1) == 10 and len(set(ids(p1))) == 10
out["mix"] = sorted(lv[i] for i in ids(p1)) == sorted(["易"] * 5 + ["中"] * 3 + ["难"] * 2)
p2 = make_paper(bank, 20, {"难": 1.0}, seed=2)
out["topup"] = len(p2) == 20 and len(set(ids(p2))) == 20
p3 = make_paper(bank, 10, ratio, seed=1, exclude_ids=set(ids(p1)))
out["exclude"] = len(p3) == 10 and not set(ids(p3)) & set(ids(p1))
record("张三", p1, "h.json")
out["history"] = set(ids(p1)) <= set(seen_ids("张三", "h.json")) and not set(seen_ids("李四", "h.json"))
print(json.dumps(out))
"""


def stems_in(docx_path, bank):
    from docx import Document

    doc = Document(docx_path)
    text = "\n".join(p.text for p in doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            text += "\n" + " ".join(c.text for c in row.cells)
    return {q["id"] for q in bank if q["stem"] in text}, text


def main():
    sc = Score()
    quiz = find("quiz.py")
    if not quiz or not os.path.exists(os.path.join(os.path.dirname(quiz), "history.py")):
        sc.fail("quiz.py + history.py not delivered side by side")
        sc.emit()
    work = workdir(os.path.dirname(quiz))
    shutil.copy(os.path.join(IN, "bank.json"), os.path.join(work, "bank.json"))
    bank = json.loads(Path(work, "bank.json").read_text(encoding="utf-8"))
    with open(os.path.join(work, "_api_check.py"), "w", encoding="utf-8") as fh:
        fh.write(API)
    proc = run([sys.executable, "_api_check.py"], cwd=work)
    try:
        res = json.loads(proc.stdout.strip().splitlines()[-1])
    except (ValueError, IndexError):
        res = {}
        sc.notes.append(f"api run: {tail(proc.stderr)}")
    sc.check(all(res.get(k) for k in ("det", "n", "mix", "topup")), f"make_paper: {res}")
    sc.check(bool(res.get("exclude") and res.get("history")), f"exclude/history: {res}")
    cli = [
        sys.executable,
        "quiz.py",
        "--bank",
        "bank.json",
        "--n",
        "10",
        "--student",
        "王五",
        "--history",
        "cli.json",
    ]
    first = run(cli + ["--out", "paper1.docx"], cwd=work)
    second = run(cli + ["--out", "paper2.docx"], cwd=work)
    ok = False
    note = f"cli rc={first.returncode}/{second.returncode}: {tail(first.stderr + second.stderr)}"
    if all(os.path.exists(os.path.join(work, f)) for f in ("paper1.docx", "paper2.docx")):
        s1, text1 = stems_in(os.path.join(work, "paper1.docx"), bank)
        s2, _t = stems_in(os.path.join(work, "paper2.docx"), bank)
        ok = len(s1) == 10 and len(s2) == 10 and not s1 & s2 and "答案" in text1
        note = f"cli papers: {len(s1)}/{len(s2)} stems, overlap {len(s1 & s2)}, 答案页 {'答案' in text1}"
    sc.check(ok, note)
    tests = os.path.join(work, "test_quiz.py")
    if sc.check(os.path.exists(tests), "test_quiz.py not delivered", essential=False):
        t = run([sys.executable, "-m", "unittest", "test_quiz"], cwd=work, timeout=120)
        sc.check(t.returncode == 0, f"own tests fail: {tail(t.stderr)}", essential=False)
    sc.emit()


main()
