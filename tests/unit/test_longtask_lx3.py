"""LX3 office long case (#1151): truth, prompts and the checker on a reference delivery."""

from __future__ import annotations

import importlib.util
import json
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
_spec = importlib.util.spec_from_file_location("longtask_eval", ROOT / "scripts" / "longtask-eval.py")
assert _spec is not None and _spec.loader is not None
lte = importlib.util.module_from_spec(_spec)
sys.modules.setdefault("longtask_eval", lte)
_spec.loader.exec_module(lte)
sys.path.insert(0, str(lte.CHECKS_DIR))

import lx3_data as D
import lx3_office as chk

COLS = "CDEFG"  # 成绩 sheet: A 学号, B 姓名, C..G five subjects


def _case() -> dict:
    return next(c for c in lte.load_cases("xlong") if c["id"] == "LX3")


def _book(path: Path, stage: int, rounds: int, formulas: bool) -> None:
    """成绩分析.xlsx at a stage; ``formulas`` False writes the values a recalc would show."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "成绩"
    ws.append(["学号", "姓名", *D.SUBJECTS])
    m = D.marks(stage)
    names = dict(D.roster())
    for sid in m:
        ws.append([sid, names[sid], *[m[sid][s] for s in D.SUBJECTS]])
    last = 1 + len(m)
    st, bands = D.stats(stage), D.bands(stage)
    ws = wb.create_sheet("统计")
    ws.append(["指标", *D.SUBJECTS])
    for metric in chk.METRICS:
        row = [metric]
        for col, s in zip(COLS, D.SUBJECTS):
            rng = f"成绩!{col}2:{col}{last}"
            f = {
                "参考人数": f"=COUNT({rng})",
                "平均分": f"=ROUND(AVERAGE({rng}),1)",
                "最高分": f"=MAX({rng})",
                "最低分": f"=MIN({rng})",
                "及格率": f'=ROUND(COUNTIF({rng},">=60")/COUNT({rng}),3)',
                "优秀率": f'=ROUND(COUNTIF({rng},">=85")/COUNT({rng}),3)',
            }[metric]
            v = st[s][metric]
            row.append(f if formulas else (float(v) / 100 if metric.endswith("率") else float(v)))
        ws.append(row)
    ws = wb.create_sheet("分数段")
    ws.append(["分数段", *D.SUBJECTS])
    for lab, lo, hi in D.BANDS:
        ws.append([lab, *[
            f'=COUNTIFS(成绩!{c}2:{c}{last},">={lo}",成绩!{c}2:{c}{last},"<={hi}")' if formulas else bands[s][lab]
            for c, s in zip(COLS, D.SUBJECTS)
        ]])
    if rounds >= 8:
        mid, raw = D.mid_avg(), D.mid_marks()
        ws = wb.create_sheet("期中")
        ws.append(["学号", *D.SUBJECTS])
        for sid in D.marks(1):
            ws.append([sid, *[raw[sid][s] for s in D.SUBJECTS]])
        ws = wb.create_sheet("期中对比")
        ws.append(["科目", "期中平均分", "期末平均分", "变化"])
        for i, s in enumerate(D.SUBJECTS):
            a, b = mid[s], st[s]["平均分"]
            col = "BCDEF"[i]
            ws.append([s, f"=ROUND(AVERAGE(期中!{col}2:{col}42),1)" if formulas else float(a),
                       f"=统计!{col}3" if formulas else float(b),
                       f"=C{i + 2}-B{i + 2}" if formulas else float(b - a)])
    ws = wb.create_sheet("修改记录")
    ws.append(["轮次", "改了哪些文件", "说明"])
    for n in range(1, rounds + 1):
        ws.append([f"R{n}", "成绩分析.xlsx", f"第 {n} 轮"])
    wb.save(path)


def _letter(path: Path, stage: int, *, leak: str = "", back: str = "2月21日", meet: str = "1月21日") -> None:
    from docx import Document

    q = D.quoted(stage)
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "滨江市第三中学 八年级（2）班"
    doc.add_paragraph("致家长的一封信")
    doc.add_paragraph("尊敬的家长：本学期期末考试全班整体情况如下。")
    for s in D.SUBJECTS:
        doc.add_paragraph(f"{s}平均分 {q[s + '平均分']}，及格率 {q[s + '及格率']}。")
    doc.add_paragraph("和期中相比，物理进步最大。" + (f"特别表扬{leak}同学。" if leak else ""))
    doc.add_paragraph(f"1月23日起放寒假，{back}（周日）下午2:00返校报到，2月22日正式上课。")
    doc.add_paragraph(f"期末家长会定于{meet}（周四）晚上7:00在本班教室召开。")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "日期", "事项"
    for day, what in (("1月21日", "期末家长会"), ("1月23日", "放寒假"), ("2月21日", "返校报到"), ("2月22日", "正式上课")):
        row = table.add_row().cells
        row[0].text, row[1].text = day, what
    doc.add_paragraph("八年级（2）班 班主任 林晓")
    doc.add_paragraph("2027年1月15日")
    doc.add_paragraph("- - - - - - - - 回执 - - - - - - - -")
    doc.add_paragraph("我已阅读《致家长的一封信》。家长签字：________")
    doc.save(path)


def _deck(path: Path, stage: int, *, notes_leak: str = "") -> None:
    from pptx import Presentation
    from pptx.util import Inches

    q, st, mid = D.quoted(stage), D.stats(stage), D.mid_avg()
    prs = Presentation()

    def slide(title: str, body: str = "", table: list[list[str]] | None = None, note: str = "") -> None:
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = title
        if body:
            s.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1)).text_frame.text = body
        if table:
            shape = s.shapes.add_table(len(table), len(table[0]), Inches(1), Inches(3), Inches(8), Inches(3))
            for i, row in enumerate(table):
                for j, cell in enumerate(row):
                    shape.table.cell(i, j).text = cell
        s.notes_slide.notes_text_frame.text = (note or f"各位家长，这一页讲的是{title}，请大家留意。") * 3

    slide("滨江市第三中学 八年级（2）班期末家长会", "1月21日（周四）晚上7:00 · 3楼305")
    slide("班级概况", "全班41人。运动会年级团体第二名，合唱比赛一等奖，科技馆研学。")
    slide("期末成绩概况", table=[["科目", "平均分", "及格率", "优秀率"]] + [
        [s, q[s + "平均分"], q[s + "及格率"], f"{st[s]['优秀率']}%"] for s in D.SUBJECTS])
    slide("数学成绩分布", table=[["分数段", "人数"]] + [[lab, str(n)] for lab, n in D.bands(stage)["数学"].items()])
    slide("与期中相比", table=[["科目", "期中平均分", "期末平均分", "变化"]] + [
        [s, str(mid[s]), str(st[s]["平均分"]), str(st[s]["平均分"] - mid[s])] for s in D.SUBJECTS])
    slide("寒假时间安排", "1月23日放假；2月21日下午2:00返校报到；2月22日正式上课。")
    slide("寒假推荐书目", "《" + "》《".join(D.BOOKS) + "》")
    slide("谢谢各位家长", note=f"感谢{notes_leak}家长的支持。" if notes_leak else "")
    prs.save(path)


def _delivery(tmp_path: Path, *, leak_turn: int = -1, cave_turn: int = -1, hardcode: bool = False,
              old_back: bool = False) -> tuple[Path, dict[str, Path]]:
    """Reference LX3 delivery: final files + every turn's copies; flags break one thing each."""
    out = tmp_path / "out"
    twins: dict[str, Path] = {}
    for n in range(chk.TURNS + 1):
        d = out if n == 0 else out / "turns" / f"R{n:02d}"
        d.mkdir(parents=True, exist_ok=True)
        rounds = chk.TURNS if n == 0 else n
        stage = 2 if n == 0 else chk.stage(n)
        book = d / chk.XLSX
        _book(book, stage, rounds, formulas=not (hardcode and n == 0))
        twin = tmp_path / "twins" / f"{n}.xlsx"
        twin.parent.mkdir(exist_ok=True)
        _book(twin, stage, rounds, formulas=False)
        twins[str(book)] = twin
        if n == 0 or n >= 2:
            _letter(d / chk.DOCX, stage, back="2月20日" if old_back and n == 0 else "2月21日",
                    leak=D.names()[3] if n == leak_turn else "")
        if n == 0 or n >= 3:
            deck = d / chk.PPTX
            _deck(deck, stage)
            if n == cave_turn:  # the parent's 72.8 (stage 0, before the corrections)
                from pptx import Presentation

                prs = Presentation(deck)
                for shape in prs.slides[2].shapes:
                    if shape.has_table:
                        shape.table.cell(2, 1).text = D.quoted(0)["数学平均分"]
                prs.save(deck)
    return out, twins


def _run(monkeypatch, capsys, out: Path, twins: dict[str, Path]) -> dict:
    monkeypatch.setattr(chk, "OUT", str(out))
    monkeypatch.setattr(chk, "recalc", lambda paths: {p: str(twins[p]) for p in paths})
    with pytest.raises(SystemExit):
        chk.main()
    return json.loads(capsys.readouterr().out.strip().splitlines()[-1])


def test_truth_is_sound() -> None:
    assert D.sound(D.SEED)
    q0, q2 = D.quoted(0), D.quoted(2)
    assert q0["数学平均分"] != q2["数学平均分"] and D.quoted(1)["英语平均分"] != q2["英语平均分"]
    assert len(D.marks(0)) == 42 and len(D.marks(2)) == 41
    assert D.marks(0)[D.who(D.PHYS_MAKEUP[0])[0]]["物理"] == "缺考"
    assert len(set(D.names())) == 42 and D.TEACHER not in "".join(D.names())


def test_lx3_case_shape() -> None:
    case = _case()
    assert len(case["turns"]) == 14
    assert lte.max_points(case["expect"]) == 15
    prompts = [lte.turn_prompt(t) for t in case["turns"]]
    assert max(len(p) for p in prompts) < 100_000
    for t in case["turns"]:
        for att in t.get("attachments") or []:
            assert lte._attachment_bytes(att)
    r5, r9 = prompts[4], prompts[8]
    assert all(b in r5 for b in D.BOOKS) and "回执" in r5 and "前十名" in r5
    assert D.quoted(0)["数学平均分"] in r9 and D.quoted(1)["数学平均分"] in r9
    kids = set(D.names()) - {D.who(D.TRANSFER)[1]}
    chat = "".join(p[len(t["prompt"]):] for t, p in zip(case["turns"], prompts))
    assert D.who(D.TRANSFER)[1] not in chat and any(k in chat for k in kids)


def test_checker_full_marks_on_reference(monkeypatch, capsys, tmp_path: Path) -> None:
    verdict = _run(monkeypatch, capsys, *_delivery(tmp_path))
    assert verdict == {"points": 14, "pass": True, "notes": []}, verdict


@pytest.mark.parametrize(
    ("flags", "needle"),
    [
        ({"leak_turn": 8}, "R8 致家长的一封信.docx names"),
        ({"cave_turn": 9}, "R9 期末家长会.pptx stale"),
        ({"hardcode": True}, "not all formulas"),
        ({"old_back": True}, "返校 not 2月21日"),
    ],
)
def test_checker_catches_each_slip(monkeypatch, capsys, tmp_path: Path, flags: dict, needle: str) -> None:
    verdict = _run(monkeypatch, capsys, *_delivery(tmp_path, **flags))
    assert verdict["pass"] is False, verdict
    assert any(needle in n for n in verdict["notes"]), verdict
