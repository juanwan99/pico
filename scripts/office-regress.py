#!/usr/bin/env python3
"""Real-model office regression (#959 · PLAN-OFFICE-COMPUTER-V2 §5).

Not CI. Run on the ECS (or any host that can reach pico-api loopback) after
prod-update, with a synthetic tenant. Opens the produced bytes and asserts
content — status flags and file counts are not accepted as proof.

  PICO_OPENAI_PROXY_KEY=… python3 scripts/office-regress.py \
      --base http://127.0.0.1:18765 --membership regress-school:regress-member

Prints a JSON report and a markdown table (paste into the Issue). Exit 1 on any
blocking failure. Records three numbers per case: tool calls, wall seconds,
artifacts produced.
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import os
import sys
import threading
import time
import uuid
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import httpx

DEFAULT_TIMEOUT_S = 420.0
FIXTURES = Path(__file__).resolve().parent.parent / "testdata" / "office-regress"


@dataclass
class CaseResult:
    case: str
    ok: bool = False
    tool_calls: int = 0
    wall_s: float = 0.0
    artifacts: int = 0
    notes: list[str] = field(default_factory=list)
    artifact_sha: dict[str, str] = field(default_factory=dict)


class Pico:
    def __init__(self, base: str, key: str, membership: str, model: str, timeout_s: float) -> None:
        self.base = base.rstrip("/")
        self.headers = {
            "Authorization": f"Bearer {key}",
            "X-Pico-Membership-Id": membership,
        }
        self.model = model
        self.timeout_s = timeout_s
        self.client = httpx.Client(timeout=httpx.Timeout(timeout_s, connect=10.0), trust_env=False)

    def upload(self, conversation_id: str, name: str, data: bytes) -> dict[str, Any]:
        resp = self.client.post(
            f"{self.base}/v1/files",
            headers={**self.headers, "X-Conversation-Id": conversation_id},
            files={"file": (name, data)},
        )
        resp.raise_for_status()
        return resp.json()

    def chat(self, conversation_id: str, text: str, *, on_first_delta=None) -> tuple[str, float]:
        """Stream a turn; return (assistant_text, wall_s)."""
        body = {
            "model": self.model,
            "stream": True,
            "messages": [{"role": "user", "content": text}],
        }
        started = time.perf_counter()
        out: list[str] = []
        first = True
        with self.client.stream(
            "POST",
            f"{self.base}/v1/chat/completions",
            headers={
                **self.headers,
                "X-Conversation-Id": conversation_id,
                "Accept": "text/event-stream",
            },
            json=body,
        ) as resp:
            resp.raise_for_status()
            for line in resp.iter_lines():
                if not line or not line.startswith("data:"):
                    continue
                payload = line[5:].strip()
                if payload == "[DONE]":
                    break
                try:
                    obj = json.loads(payload)
                except ValueError:
                    continue
                for choice in obj.get("choices") or []:
                    delta = (choice.get("delta") or {}).get("content")
                    if delta:
                        out.append(str(delta))
                        if first and on_first_delta:
                            first = False
                            on_first_delta()
        return "".join(out), time.perf_counter() - started

    def tasks(self, conversation_id: str) -> list[dict[str, Any]]:
        resp = self.client.get(
            f"{self.base}/v1/tasks",
            headers=self.headers,
            params={"conversation_id": conversation_id},
        )
        resp.raise_for_status()
        return list(resp.json().get("tasks") or [])

    def runs(self, task_id: str) -> list[dict[str, Any]]:
        resp = self.client.get(f"{self.base}/v1/tasks/{task_id}/runs", headers=self.headers)
        resp.raise_for_status()
        data = resp.json()
        return list(data.get("runs") or data or [])

    def events(self, run_id: str) -> list[dict[str, Any]]:
        resp = self.client.get(f"{self.base}/v1/runs/{run_id}/events", headers=self.headers)
        resp.raise_for_status()
        return list(resp.json().get("events") or [])

    def run(self, run_id: str) -> dict[str, Any]:
        resp = self.client.get(f"{self.base}/v1/runs/{run_id}", headers=self.headers)
        resp.raise_for_status()
        return dict(resp.json().get("run") or {})

    def cancel_active(self, task_id: str) -> dict[str, Any]:
        resp = self.client.post(
            f"{self.base}/v1/tasks/{task_id}/cancel-active", headers=self.headers
        )
        return resp.json() if resp.content else {"status_code": resp.status_code}

    def artifacts(self, conversation_id: str) -> list[dict[str, Any]]:
        resp = self.client.get(
            f"{self.base}/v1/artifacts",
            headers=self.headers,
            params={"conversation_id": conversation_id},
        )
        resp.raise_for_status()
        return list(resp.json().get("artifacts") or [])

    def download(self, artifact_id: str) -> bytes:
        resp = self.client.get(
            f"{self.base}/v1/artifacts/{artifact_id}/content",
            headers=self.headers,
            params={"download": "true"},
        )
        resp.raise_for_status()
        return resp.content


def _sha(data: bytes) -> str:
    import hashlib

    return hashlib.sha256(data).hexdigest()[:16]


def _tool_calls(pico: Pico, conversation_id: str) -> tuple[int, str]:
    total = 0
    status = "?"
    for task in pico.tasks(conversation_id):
        latest = task.get("latest_run") or {}
        if latest.get("id"):
            status = str(latest.get("status") or status)
            total += sum(
                1
                for e in pico.events(latest["id"])
                if str(e.get("kind") or e.get("type") or "") == "tool.call"
            )
    return total, status


PRODUCED_KINDS = {"xlsx", "docx", "pptx", "html"}


def _produced(arts: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Model-written files only. Uploads land as edu_office / edu_excerpt / kb_text.
    Same title as the original is the product rule (同名替换), so filter by kind."""
    return [a for a in arts if str(a.get("kind") or "").lower() in PRODUCED_KINDS]


def _latest_by_suffix(arts: list[dict[str, Any]], suffix: str) -> dict[str, Any] | None:
    matches = [a for a in arts if str(a.get("title") or "").lower().endswith(suffix)]
    if not matches:
        return None
    return max(matches, key=lambda a: str(a.get("created_at") or ""))


# --- fixtures -----------------------------------------------------------------


def make_gradebook() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "成绩"
    ws.append(["姓名", "平时", "期末", "总分"])
    for name, a, b in [
        ("甲", 88, 92),
        ("乙", 75, 80),
        ("丙", 91, 85),
        ("丁", 60, 72),
        ("戊", 83, 79),
        ("己", 95, 97),
    ]:
        ws.append([name, a, b, None])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


ROSTER_ROWS = [
    ("张一", "2401", "A"),
    ("李二", "2402", "A"),
    ("王三", "2403", "B"),
    ("赵四", "2404", "C"),
    ("孙五", "2405", "A"),
    ("周六", "2406", "B"),
    ("吴七", "2407", "C"),
    ("郑八", "2408", "A"),
    ("冯九", "2409", "B"),
    ("陈十", "2410", "A"),
]


def make_roster() -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["姓名", "学号", "组别"])
    w.writerows(ROSTER_ROWS)
    return buf.getvalue().encode("utf-8")


def roster_counts() -> dict[str, int]:
    out: dict[str, int] = {}
    for _, _, g in ROSTER_ROWS:
        out[g] = out.get(g, 0) + 1
    return out


# --- cases --------------------------------------------------------------------


def case_t1(pico: Pico, stamp: str) -> CaseResult:
    from openpyxl import load_workbook

    res = CaseResult(case="T1a/T1b")
    cid = f"regress-t1-{stamp}"
    pico.upload(cid, "gradebook.xlsx", make_gradebook())
    # Name the column, not just the range: a bare "D2:D7" once made the model stop and
    # ask whether D was 期末 or 总分 (correct behaviour, but a false negative here).
    _, wall = pico.chat(
        cid,
        "gradebook.xlsx 在附件里。表头是 姓名/平时/期末/总分。"
        "把总分列 D2:D7 写成期末40%加平时60%的公式，保存为 xlsx。",
    )
    res.wall_s += wall
    produced = _produced(pico.artifacts(cid))
    res.artifacts = len(produced)
    latest = _latest_by_suffix(produced, ".xlsx")
    if not latest:
        res.notes.append("T1a: no new xlsx artifact")
        res.tool_calls, _ = _tool_calls(pico, cid)
        return res
    raw = pico.download(latest["id"])
    res.artifact_sha[latest["title"]] = _sha(raw)
    ws = load_workbook(io.BytesIO(raw)).active
    ok_a = True
    for r in range(2, 8):
        v = ws[f"D{r}"].value
        if not (isinstance(v, str) and v.startswith("=") and f"B{r}" in v and f"C{r}" in v):
            ok_a = False
            res.notes.append(f"T1a: D{r}={v!r} is not a B/C formula")
    if ws["A2"].value != "甲" or ws["B2"].value != 88 or ws["C2"].value != 92:
        ok_a = False
        res.notes.append("T1a: original A/B/C values changed")
    # T1b — second round, same conversation
    _, wall2 = pico.chat(cid, "把标题改成「三年二班成绩」，D 列公式别丢。")
    res.wall_s += wall2
    produced2 = _produced(pico.artifacts(cid))
    res.artifacts = len(produced2)
    latest2 = _latest_by_suffix(produced2, ".xlsx")
    ok_b = False
    if latest2:
        raw2 = pico.download(latest2["id"])
        res.artifact_sha[latest2["title"] + "#2"] = _sha(raw2)
        wb2 = load_workbook(io.BytesIO(raw2))
        blob = json.dumps(
            [
                [str(c.value) for c in row]
                for ws2 in wb2.worksheets
                for row in ws2.iter_rows(max_row=3)
            ],
            ensure_ascii=False,
        )
        titled = (
            "三年二班" in blob
            or any("三年二班" in (s.title or "") for s in wb2.worksheets)
            or "三年二班" in str(latest2.get("title"))
        )
        d2 = wb2.active["D2"].value
        keeps = isinstance(d2, str) and d2.startswith("=")
        ok_b = titled and keeps
        if not titled:
            res.notes.append("T1b: title 三年二班 not found")
        if not keeps:
            res.notes.append(f"T1b: D2={d2!r} formula lost")
    else:
        res.notes.append("T1b: no second xlsx")
    res.tool_calls, _ = _tool_calls(pico, cid)
    res.ok = ok_a and ok_b
    return res


def case_t2(pico: Pico, stamp: str) -> CaseResult:
    from docx import Document
    from openpyxl import load_workbook

    res = CaseResult(case="T2")
    cid = f"regress-t2-{stamp}"
    pico.upload(cid, "roster.csv", make_roster())
    _, wall = pico.chat(
        cid,
        "roster.csv 在附件里。用这个 CSV 做两份东西：1) 按组别汇总人数的 xlsx；2) 一页说明 Word，点名各组人数。不要网页。",
    )
    res.wall_s = wall
    produced = _produced(pico.artifacts(cid))
    res.artifacts = len(produced)
    expect = roster_counts()
    xlsx = _latest_by_suffix(produced, ".xlsx")
    docx = _latest_by_suffix(produced, ".docx")
    ok = True
    if not xlsx:
        ok = False
        res.notes.append("T2: no xlsx")
    else:
        raw = pico.download(xlsx["id"])
        res.artifact_sha[xlsx["title"]] = _sha(raw)
        cells: dict[str, int] = {}
        for ws in load_workbook(io.BytesIO(raw), data_only=False).worksheets:
            for row in ws.iter_rows(values_only=True):
                vals = [v for v in row if v is not None]
                for i, v in enumerate(vals):
                    if isinstance(v, str) and v.strip() in expect and i + 1 < len(vals):
                        nxt = vals[i + 1]
                        try:
                            cells[v.strip()] = int(nxt)
                        except (TypeError, ValueError):
                            pass
        for g, n in expect.items():
            if cells.get(g) != n:
                ok = False
                res.notes.append(f"T2 xlsx: group {g} expected {n} got {cells.get(g)}")
    if not docx:
        ok = False
        res.notes.append("T2: no docx")
    else:
        raw = pico.download(docx["id"])
        res.artifact_sha[docx["title"]] = _sha(raw)
        d = Document(io.BytesIO(raw))
        text = "\n".join(p.text for p in d.paragraphs)
        for t in d.tables:
            for row in t.rows:
                text += "\n" + " ".join(c.text for c in row.cells)
        for g, n in expect.items():
            if g not in text or str(n) not in text:
                ok = False
                res.notes.append(f"T2 docx: group {g}/{n} not named")
    if any(str(a.get("title") or "").lower().endswith(".html") for a in produced):
        res.notes.append("T2: an HTML was produced although 不要网页 (observation)")
    res.tool_calls, _ = _tool_calls(pico, cid)
    res.ok = ok
    return res


def case_t3(pico: Pico, stamp: str) -> CaseResult:
    import re

    res = CaseResult(case="T3-files")
    cid = f"regress-t3-{stamp}"
    _, wall = pico.chat(cid, "做一页离线可开的 HTML 介绍潮汐，再做一个 3 页的 PPT 提纲。")
    res.wall_s = wall
    arts = _produced(pico.artifacts(cid))
    res.artifacts = len(arts)
    html = _latest_by_suffix(arts, ".html")
    pptx = _latest_by_suffix(arts, ".pptx")
    ok = True
    if not html:
        ok = False
        res.notes.append("T3: no html")
    else:
        raw = pico.download(html["id"]).decode("utf-8", errors="replace")
        res.artifact_sha[html["title"]] = _sha(raw.encode("utf-8"))
        if re.search(r"""(src|href)=["']https?://""", raw, re.IGNORECASE):
            ok = False
            res.notes.append("T3: html has http(s) external asset")
    if not pptx:
        ok = False
        res.notes.append("T3: no pptx")
    else:
        raw = pico.download(pptx["id"])
        res.artifact_sha[pptx["title"]] = _sha(raw)
        nslides = _pptx_slide_count(raw)
        blob = _slide_blob(raw)
        if nslides < 3:
            ok = False
            res.notes.append(f"T3: pptx has {nslides} slides (<3)")
        if not any(line.strip() and not line.startswith("slides=") for line in blob.splitlines()):
            ok = False
            res.notes.append("T3: pptx has no visible text")
    res.tool_calls, _ = _tool_calls(pico, cid)
    res.ok = ok
    return res


def case_t4(pico: Pico, stamp: str) -> CaseResult:
    res = CaseResult(case="T4-cancel")
    cid = f"regress-t4-{stamp}"
    pico.upload(cid, "gradebook.xlsx", make_gradebook())
    before = len(_produced(pico.artifacts(cid)))
    cancelled: dict[str, Any] = {}

    def stop_soon() -> None:
        time.sleep(3.0)
        for _ in range(20):
            tasks = pico.tasks(cid)
            if tasks:
                cancelled["resp"] = pico.cancel_active(tasks[0]["id"])
                cancelled["task_id"] = tasks[0]["id"]
                return
            time.sleep(0.5)

    t = threading.Thread(target=stop_soon, daemon=True)
    started = time.perf_counter()
    t.start()
    try:
        pico.chat(
            cid,
            "gradebook.xlsx 在附件里。表头是 姓名/平时/期末/总分。"
            "把总分列 D2:D7 写成期末40%加平时60%的公式，再把每个人的评语写成一段话，保存为 xlsx。",
        )
    except httpx.HTTPError as exc:
        res.notes.append(f"T4: stream ended with {type(exc).__name__} (acceptable on cancel)")
    t.join(timeout=30)
    res.wall_s = time.perf_counter() - started
    time.sleep(2.0)
    after = _produced(pico.artifacts(cid))
    res.artifacts = len(after) - before
    status = "?"
    if cancelled.get("task_id"):
        for run in pico.runs(cancelled["task_id"]):
            status = str(run.get("status") or status)
    res.notes.append(
        f"T4: run status={status} cancel={json.dumps(cancelled.get('resp'), ensure_ascii=False)[:120]}"
    )
    res.tool_calls, _ = _tool_calls(pico, cid)
    res.ok = status in {"cancelled", "cancelling"} and res.artifacts == 0
    if res.artifacts:
        res.notes.append(f"T4: {res.artifacts} artifact(s) landed after Stop")
    return res


def make_week_deck() -> bytes:
    canned = FIXTURES / "week.pptx"
    if canned.is_file():
        return canned.read_bytes()
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    titles = ("封面", "课表", "作业")
    for i, title in enumerate(titles):
        slide = prs.slides.add_slide(blank)
        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.0)
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12), Inches(0.6))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = title
        run.font.size = Pt(28)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        body = slide.shapes.add_textbox(Inches(0.5), Inches(1.4), Inches(12), Inches(2))
        body.text_frame.paragraphs[0].text = f"第{i + 1}页保留标记 KEEP-{title}"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_notice_docx() -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    h = doc.add_heading("通知", level=1)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    p1 = doc.add_paragraph()
    r1 = p1.add_run("第一段春游安排，不得整篇重写。")
    r1.bold = True
    r1.font.size = Pt(12)
    doc.add_paragraph("第二段安全事项，必须留下。")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_chart_book() -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference

    wb = Workbook()
    ws = wb.active
    ws.title = "人数"
    ws.append(["组别", "人数"])
    ws.append(["甲", 12])
    ws.append(["乙", 8])
    ws.append(["丙", 5])
    chart = BarChart()
    chart.title = "各组人数"
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=4), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=4))
    ws.add_chart(chart, "E2")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_formula_book() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "成绩"
    ws.append(["姓名", "平时", "期末", "总分", "备注"])
    ws.append(["甲", 88, 92, "=B2*0.6+C2*0.4", "原备注留着"])
    ws.append(["乙", 75, 80, "=B3*0.6+C3*0.4", "乙备注"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pptx_slide_count(raw: bytes) -> int:
    import zipfile

    with zipfile.ZipFile(io.BytesIO(raw)) as zf:
        return sum(
            1
            for name in zf.namelist()
            if name.startswith("ppt/slides/slide") and name.endswith(".xml")
        )


def _pptx_xml_text(raw: bytes) -> str:
    import re
    import zipfile

    parts: list[str] = []
    with zipfile.ZipFile(io.BytesIO(raw)) as zf:
        names = [
            name
            for name in zf.namelist()
            if name.startswith("ppt/slides/slide") and name.endswith(".xml")
        ]
        for name in sorted(names):
            xml = zf.read(name).decode("utf-8", errors="replace")
            parts.extend(re.findall(r"<a:t>([^<]*)</a:t>", xml))
    return "\n".join(parts)


def _slide_blob(raw: bytes) -> str:
    try:
        from pptx import Presentation
    except ImportError:
        return f"slides={_pptx_slide_count(raw)}\n{_pptx_xml_text(raw)}"
    prs = Presentation(io.BytesIO(raw))
    parts: list[str] = [f"slides={len(prs.slides)}"]
    for s in prs.slides:
        for sh in s.shapes:
            if getattr(sh, "has_text_frame", False):
                parts.append(sh.text_frame.text)
    return "\n".join(parts)


def _docx_blob(raw: bytes) -> str:
    from docx import Document

    d = Document(io.BytesIO(raw))
    return "\n".join(p.text for p in d.paragraphs)


def _docx_has_media(raw: bytes) -> bool:
    import zipfile

    with zipfile.ZipFile(io.BytesIO(raw)) as zf:
        return any(name.startswith("word/media/") for name in zf.namelist())


def case_t5(pico: Pico, stamp: str) -> CaseResult:
    res = CaseResult(case="T5-ppt-template")
    cid = f"regress-t5-{stamp}"
    pico.upload(cid, "周课.pptx", make_week_deck())
    _, wall = pico.chat(
        cid,
        "周课.pptx 在附件里。只改第二页标题「课表」为「本周课表」。"
        "封面页、作业页的文字和蓝条不要动，页数保持 3 页。",
    )
    res.wall_s = wall
    produced = _produced(pico.artifacts(cid))
    res.artifacts = len(produced)
    latest = _latest_by_suffix(produced, ".pptx")
    ok = False
    if not latest:
        res.notes.append("T5: no pptx")
    else:
        raw = pico.download(latest["id"])
        res.artifact_sha[latest["title"]] = _sha(raw)
        nslides = _pptx_slide_count(raw)
        blob = _slide_blob(raw)
        ok = (
            nslides == 3
            and "本周课表" in blob
            and "KEEP-封面" in blob
            and "KEEP-作业" in blob
        )
        if nslides != 3:
            res.notes.append(f"T5: slides={nslides}")
        if "本周课表" not in blob:
            res.notes.append("T5: 本周课表 missing")
        if "KEEP-封面" not in blob or "KEEP-作业" not in blob:
            res.notes.append("T5: unused page mark lost")
    res.tool_calls, _ = _tool_calls(pico, cid)
    res.ok = ok
    return res


def case_t6(pico: Pico, stamp: str) -> CaseResult:
    res = CaseResult(case="T6-word-rounds")
    cid = f"regress-t6-{stamp}"
    pico.upload(cid, "通知.docx", make_notice_docx())
    _, wall = pico.chat(
        cid, "通知.docx 在附件里。只把第一段的「春游」改成「秋游」，标题和其余段不要重排。"
    )
    res.wall_s += wall
    _, wall2 = pico.chat(cid, "在第二段末尾加上「家长签字」。第一段的秋游和标题通知都要还在。")
    res.wall_s += wall2
    produced = _produced(pico.artifacts(cid))
    res.artifacts = len(produced)
    latest = _latest_by_suffix(produced, ".docx")
    ok = False
    if not latest:
        res.notes.append("T6: no docx")
    else:
        raw = pico.download(latest["id"])
        res.artifact_sha[latest["title"]] = _sha(raw)
        blob = _docx_blob(raw)
        ok = "秋游" in blob and "家长签字" in blob and "通知" in blob and "安全" in blob
        if not ok:
            res.notes.append(f"T6: blob={blob[:180]!r}")
    res.tool_calls, _ = _tool_calls(pico, cid)
    res.ok = ok
    return res


def case_t7(pico: Pico, stamp: str) -> CaseResult:
    res = CaseResult(case="T7-chart-to-word")
    cid = f"regress-t7-{stamp}"
    pico.upload(cid, "人数.xlsx", make_chart_book())
    _, wall = pico.chat(
        cid,
        "人数.xlsx 在附件里。做一份 Word，把表里的柱状图嵌进文档（真图片，不要只抄甲12乙8）。",
    )
    res.wall_s = wall
    produced = _produced(pico.artifacts(cid))
    res.artifacts = len(produced)
    latest = _latest_by_suffix(produced, ".docx")
    ok = False
    if not latest:
        res.notes.append("T7: no docx")
    else:
        raw = pico.download(latest["id"])
        res.artifact_sha[latest["title"]] = _sha(raw)
        blob = _docx_blob(raw)
        ok = _docx_has_media(raw)
        if not ok:
            res.notes.append(f"T7: no word/media (text={blob[:120]!r})")
    res.tool_calls, _ = _tool_calls(pico, cid)
    res.ok = ok
    return res


def case_t8(pico: Pico, stamp: str) -> CaseResult:
    res = CaseResult(case="T8-legacy-doc")
    cid = f"regress-t8-{stamp}"
    pico.upload(cid, "教师计划.doc", make_notice_docx())
    _, wall = pico.chat(
        cid,
        "教师计划.doc 是旧版文件名。转成能改的 Word 后，只把「春游」改成「秋游」。"
        "转失败就老实说转不开，不要假装已经交给模型。",
    )
    res.wall_s = wall
    produced = _produced(pico.artifacts(cid))
    res.artifacts = len(produced)
    latest = _latest_by_suffix(produced, ".docx") or _latest_by_suffix(produced, ".doc")
    ok = False
    if not latest:
        res.notes.append("T8: no converted/edited word")
    else:
        raw = pico.download(latest["id"])
        res.artifact_sha[latest["title"]] = _sha(raw)
        try:
            blob = _docx_blob(raw)
        except Exception as exc:  # noqa: BLE001 — not OOXML = fail closed
            res.notes.append(f"T8: not openable docx ({type(exc).__name__})")
            blob = ""
        ok = "秋游" in blob and "安全" in blob
        if not ok and blob:
            res.notes.append(f"T8: blob={blob[:160]!r}")
    res.tool_calls, _ = _tool_calls(pico, cid)
    res.ok = ok
    return res


def case_t9(pico: Pico, stamp: str) -> CaseResult:
    from openpyxl import load_workbook

    res = CaseResult(case="T9-same-name")
    cid = f"regress-t9-{stamp}"
    original = make_gradebook()
    pico.upload(cid, "成绩单.xlsx", original)
    sha0 = _sha(original)
    _, wall = pico.chat(
        cid,
        "成绩单.xlsx 在附件里。把表头第一行改成「期末成绩」，保存时文件名仍是 成绩单.xlsx（覆盖老师盘）。",
    )
    res.wall_s = wall
    produced = _produced(pico.artifacts(cid))
    res.artifacts = len(produced)
    latest = _latest_by_suffix(produced, ".xlsx")
    ok = False
    if not latest:
        res.notes.append("T9: no xlsx")
    else:
        raw = pico.download(latest["id"])
        sha1 = _sha(raw)
        res.artifact_sha[str(latest.get("title"))] = sha1
        titled = str(latest.get("title") or "")
        wb = load_workbook(io.BytesIO(raw))
        blob = json.dumps(
            [[str(c.value) for c in row] for row in wb.active.iter_rows(max_row=2)],
            ensure_ascii=False,
        )
        ok = sha1 != sha0 and "成绩单.xlsx" in titled and "期末成绩" in blob
        if sha1 == sha0:
            res.notes.append("T9: sha unchanged")
        if "成绩单.xlsx" not in titled:
            res.notes.append(f"T9: title={titled!r}")
        if "期末成绩" not in blob:
            res.notes.append("T9: 期末成绩 missing")
    res.tool_calls, _ = _tool_calls(pico, cid)
    res.ok = ok
    return res


def case_t10(pico: Pico, stamp: str) -> CaseResult:
    from openpyxl import load_workbook

    res = CaseResult(case="T10-cell-keep")
    cid = f"regress-t10-{stamp}"
    pico.upload(cid, "定位改.xlsx", make_formula_book())
    _, wall = pico.chat(
        cid,
        "定位改.xlsx 在附件里。只把 B2 改成 99。D 列公式和 E 列备注一个字都不要动。",
    )
    res.wall_s = wall
    produced = _produced(pico.artifacts(cid))
    res.artifacts = len(produced)
    latest = _latest_by_suffix(produced, ".xlsx")
    ok = False
    if not latest:
        res.notes.append("T10: no xlsx")
    else:
        raw = pico.download(latest["id"])
        res.artifact_sha[latest["title"]] = _sha(raw)
        ws = load_workbook(io.BytesIO(raw)).active
        b2, d2, e2 = ws["B2"].value, ws["D2"].value, ws["E2"].value
        ok = (
            b2 == 99
            and isinstance(d2, str)
            and d2.startswith("=")
            and "原备注留着" in str(e2 or "")
        )
        if not ok:
            res.notes.append(f"T10: B2={b2!r} D2={d2!r} E2={e2!r}")
    res.tool_calls, _ = _tool_calls(pico, cid)
    res.ok = ok
    return res


CASES = {
    "t1": case_t1,
    "t2": case_t2,
    "t3": case_t3,
    "t4": case_t4,
    "t5": case_t5,
    "t6": case_t6,
    "t7": case_t7,
    "t8": case_t8,
    "t9": case_t9,
    "t10": case_t10,
}


def main() -> int:
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    ap.add_argument("--base", default=os.environ.get("PICO_BASE") or "http://127.0.0.1:18765")
    ap.add_argument("--key", default=os.environ.get("PICO_OPENAI_PROXY_KEY") or "")
    ap.add_argument(
        "--membership",
        default=os.environ.get("PICO_REGRESS_MEMBERSHIP") or "regress-school:regress-member",
    )
    ap.add_argument("--model", default=os.environ.get("PICO_REGRESS_MODEL") or "pico-fast")
    ap.add_argument("--cases", default="t1,t2,t3,t4,t5,t6,t7,t8,t9,t10")
    ap.add_argument("--timeout", type=float, default=DEFAULT_TIMEOUT_S)
    ap.add_argument("--json", default="", help="write JSON report here")
    args = ap.parse_args()
    if not args.key:
        print("PICO_OPENAI_PROXY_KEY / --key required", file=sys.stderr)
        return 2
    pico = Pico(args.base, args.key, args.membership, args.model, args.timeout)
    stamp = f"{int(time.time())}-{uuid.uuid4().hex[:6]}"
    results: list[CaseResult] = []
    for name in [c.strip().lower() for c in args.cases.split(",") if c.strip()]:
        fn = CASES.get(name)
        if not fn:
            print(f"unknown case {name}", file=sys.stderr)
            return 2
        try:
            results.append(fn(pico, stamp))
        except Exception as exc:  # noqa: BLE001 — report, do not hide
            results.append(
                CaseResult(
                    case=name.upper(), ok=False, notes=[f"exception: {type(exc).__name__}: {exc}"]
                )
            )
    report = {
        "base": args.base,
        "model": args.model,
        "membership": args.membership,
        "stamp": stamp,
        "results": [r.__dict__ for r in results],
        "pass": all(r.ok for r in results),
    }
    if args.json:
        Path(args.json).write_text(
            json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
        )
    print(json.dumps(report, ensure_ascii=False, indent=2))
    print()
    print("| case | ok | tool calls | wall s | artifacts | notes |")
    print("|---|---|---:|---:|---:|---|")
    for r in results:
        print(
            f"| {r.case} | {'✅' if r.ok else '❌'} | {r.tool_calls} | {r.wall_s:.1f} | {r.artifacts} | {'; '.join(r.notes)[:200]} |"
        )
    return 0 if report["pass"] else 1


if __name__ == "__main__":
    sys.exit(main())
