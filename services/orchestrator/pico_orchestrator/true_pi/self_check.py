"""Pre-landing self-check: the answer against what is really in outputs/ (pico#1334 ②).

The runner only lists / serves outputs once the box is stopped (409 while it
runs), and once stopped Pi cannot do another round. So the check runs while the
box is alive, as a read-only script over Pi RPC ``bash`` with
``excludeFromContext`` (never in the model's context or session). What it
reports is the box's word: it can only decide whether the model gets one more
round, or fail its own run. Landing and the API delivery gate still read the
real files through the runner.

Facts only, no how-to: the model gets 「系统核对：…」 lines and nothing else.
"""

from __future__ import annotations

import base64
import json
import os
import re
from collections.abc import Mapping
from typing import Any

from pico_orchestrator.delivery_policy import (
    looks_like_clarification,
    looks_like_delivery_claim,
    looks_like_redelivery_claim,
)
from pico_orchestrator.document_generators import MIN_DOCX_BODY_CHARS, MIN_PPTX_SLIDES

CHECK_TIMEOUT_S = 10.0

# Runs inside the box (python3 + python-docx / python-pptx / openpyxl / pypdf).
# argv[1] = base64 JSON {path: sha256} of outputs/ before this turn; argv[2] = root (tests).
_BOX_SCRIPT = r"""
import base64, hashlib, json, os, sys, zipfile
before = json.loads(base64.b64decode(sys.argv[1]).decode())
skip = {".pyc", ".log", ".tmp", ".lock"}
root = sys.argv[2] if len(sys.argv) > 2 else "/workspace"
files, names = [], []
for dp, _ds, fs in os.walk(os.path.join(root, "outputs"), followlinks=False):
    for f in sorted(fs):
        full = os.path.join(dp, f)
        if os.path.islink(full) or not os.path.isfile(full):
            continue
        rel = os.path.relpath(full, root)
        ext = os.path.splitext(f)[1].lower()
        if f.startswith(".") or ext in skip:
            continue
        names.append(f)
        with open(full, "rb") as fh:
            sha = hashlib.sha256(fh.read()).hexdigest()
        if before.get(rel) == sha:
            continue
        row = {"name": f, "ext": ext, "error": None}
        try:
            if ext in (".docx", ".pptx", ".xlsx"):
                with zipfile.ZipFile(full) as z:
                    z.testzip()
            if ext == ".docx":
                import docx
                d = docx.Document(full)
                parts = [p.text for p in d.paragraphs]
                for t in d.tables:
                    for r in t.rows:
                        parts.extend(c.text for c in r.cells)
                row["chars"] = len("".join("".join(parts).split()))
            elif ext == ".pptx":
                import pptx
                p = pptx.Presentation(full)
                row["slides"] = len(p.slides)
                row["titled"] = sum(
                    1 for s in p.slides
                    if any(getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip() for sh in s.shapes)
                )
            elif ext == ".xlsx":
                import openpyxl
                wb = openpyxl.load_workbook(full, read_only=True)
                row["cells"] = sum(
                    1 for ws in wb.worksheets for r in ws.iter_rows(values_only=True) for v in r if v not in (None, "")
                )
            elif ext == ".pdf":
                import pypdf
                row["pages"] = len(pypdf.PdfReader(full).pages)
        except Exception as exc:
            row["error"] = type(exc).__name__
        files.append(row)
# Files already in the box outside outputs/ (attachments, materials): naming one is not a claim.
others = set()
for dp, ds, fs in os.walk(root, followlinks=False):
    ds[:] = [d for d in ds if not d.startswith(".") and os.path.join(dp, d) != os.path.join(root, "outputs")]
    others.update(f for f in fs if not f.startswith("."))
    if len(others) > 5000:
        break
print(json.dumps({"files": files, "names": names, "others": sorted(others)}, ensure_ascii=False))
"""


def box_check_command(before: Mapping[str, Any] | None) -> str:
    """Shell line for Pi RPC ``bash``: script and snapshot both base64, no quoting games."""
    snap = {str(k): str((v or {}).get("sha256") or "") for k, v in (before or {}).items()}
    script = base64.b64encode(_BOX_SCRIPT.encode()).decode()
    arg = base64.b64encode(json.dumps(snap).encode()).decode()
    return f"echo {script} | base64 -d | python3 - {arg}"


def parse_box_report(output: str) -> dict[str, Any] | None:
    for line in reversed((output or "").strip().splitlines()):
        line = line.strip()
        if line.startswith("{"):
            try:
                data = json.loads(line)
            except ValueError:
                return None
            if isinstance(data, dict) and isinstance(data.get("files"), list):
                return data
            return None
    return None


_EXT = r"\.(?:docx|pptx|xlsx|pdf|html|png|jpg|csv|zip|md|txt)"
# Only names the answer marks off (outputs/ path, `code`, 「」, 《》, quotes); prose
# around a bare name would be read as part of it.
_NAMED_FILE = re.compile(
    r"outputs/([^\s`'\"<>|]+?" + _EXT + r")(?![\w.])"
    r"|[`「《“\"]([^`」》”\"\n/]+?" + _EXT + r")[`」》”\"]",
    re.IGNORECASE,
)
_CN_NUM = {"一": 1, "两": 2, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7, "八": 8, "九": 9, "十": 10}
_NUM = r"(\d{1,3}|[一两二三四五六七八九十])"
_SLIDES = re.compile(_NUM + r"\s*(?:张\s*(?:幻灯片|PPT|课件)?|页\s*(?:幻灯片|PPT|课件))", re.IGNORECASE)
_PDF_PAGES = re.compile(_NUM + r"\s*页\s*(?:的\s*)?PDF", re.IGNORECASE)
_COPIES = re.compile(_NUM + r"\s*份")


def _num(raw: str) -> int:
    return int(raw) if raw.isdigit() else _CN_NUM.get(raw, 0)


def self_check_facts(answer: str, report: Mapping[str, Any], *, had_files_before: bool) -> list[str]:
    """Chinese fact lines when the answer and outputs/ disagree. Empty = fine."""
    text = (answer or "").strip()
    if not text or looks_like_clarification(text):
        return []
    files = [f for f in report.get("files") or [] if isinstance(f, Mapping)]
    names = {str(n) for n in report.get("names") or []}
    others = {str(n) for n in report.get("others") or []}
    facts: list[str] = []
    # Same reading as the API delivery gate (#1167): restating earlier files is fine,
    # saying they were made (again) this turn with nothing new is not.
    if not files and (
        (looks_like_delivery_claim(text) and not had_files_before) or looks_like_redelivery_claim(text)
    ):
        facts.append("回答说交了文件，但这一轮 outputs 里没有新文件，也没有改过的文件。")
    # An outputs/ path must be in outputs; a quoted name may also be a file the teacher gave (attachments/).
    named: dict[str, bool] = {}
    for m in _NAMED_FILE.finditer(text):
        name = os.path.basename(m.group(1) or m.group(2))
        named[name] = named.get(name, False) or bool(m.group(1))
    for name, in_outputs in named.items():
        if name not in names and (in_outputs or name not in others):
            facts.append(f"回答提到的 {name} 在 outputs 里没有。")
    for f in files:
        name = str(f.get("name") or "")
        if f.get("error"):
            facts.append(f"{name} 打不开（{f.get('error')}）。")
            continue
        if f.get("ext") == ".docx" and int(f.get("chars") or 0) < MIN_DOCX_BODY_CHARS:
            facts.append(f"{name} 打开后正文只有 {int(f.get('chars') or 0)} 个字。")
        if f.get("ext") == ".pptx" and int(f.get("titled") or 0) < MIN_PPTX_SLIDES:
            facts.append(f"{name} 打开后没有有字的页。")
        if f.get("ext") == ".xlsx" and int(f.get("cells") or 0) == 0:
            facts.append(f"{name} 打开后所有表都是空的。")
    decks = [f for f in files if f.get("ext") == ".pptx" and not f.get("error")]
    if len(decks) == 1:
        said = {_num(m.group(1)) for m in _SLIDES.finditer(text)} - {0}
        have = int(decks[0].get("slides") or 0)
        if len(said) == 1 and have not in said:
            facts.append(f"{decks[0].get('name')} 有 {have} 张幻灯片，回答说 {said.pop()} 张。")
    pdfs = [f for f in files if f.get("ext") == ".pdf" and not f.get("error")]
    if len(pdfs) == 1:
        said = {_num(m.group(1)) for m in _PDF_PAGES.finditer(text)} - {0}
        have = int(pdfs[0].get("pages") or 0)
        if len(said) == 1 and have not in said:
            facts.append(f"{pdfs[0].get('name')} 有 {have} 页，回答说 {said.pop()} 页。")
    # 「N 份」: several separate files but fewer than N. One combined file is a choice, not a miss.
    said = {_num(m.group(1)) for m in _COPIES.finditer(text)} - {0}
    if len(said) == 1 and 2 <= len(files) < next(iter(said)):
        facts.append(f"这一轮 outputs 里有 {len(files)} 个新文件，回答说 {next(iter(said))} 份。")
    return facts


def self_check_prompt(facts: list[str]) -> str:
    """What the model gets for its one extra round: the facts, nothing else."""
    return "系统核对：\n" + "\n".join(f"- {f}" for f in facts)


def self_check_teacher_note(facts: list[str]) -> str:
    return f"交付前自查：{facts[0].rstrip('。')}，补做一轮"


def self_check_failure(facts: list[str]) -> str:
    """Short enough that user_errors shows it as is (it hides anything over 180)."""
    more = f"（另有 {len(facts) - 1} 处）" if len(facts) > 1 else ""
    return f"交付前核对没通过：{facts[0].rstrip('。')[:110]}{more}。"
