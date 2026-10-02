"""LX3: fourteen turns on three office files for parents; turn-1 rules and numbers must hold.

Truth comes from lx3_data (same seed as the attachments). Workbooks are recalculated
with LibreOffice before reading, since the rules demand formulas. Final files:
  成绩分析.xlsx       统计 = stage-2 truth · every statistic a formula · 分数段 · 期中对比
  致家长的一封信.docx  stage-2 numbers, none of the replaced ones · 页眉 · 落款 · 回执 ·
                      寒假时间安排 table · 返校 2月21日 · 家长会 1月21日 · 物理 progressed most
  期末家长会.pptx      stage-2 numbers · 41 人 · 数学分数段 / 期中 tables · cover · 推荐书目 ·
                      notes on every slide · ≤80 汉字 body per slide
  both                no student name anywhere (notes and headers included)
Per turn (/w/out/turns/Rnn, newest copy of each file so far): workbook delivered with an
Rn row in 修改记录 and formula statistics; letter and slides carry no names and that
stage's numbers (stage 0 before R4, 1 before R11, 2 after), none replaced earlier.
"""

import hashlib
import itertools
import os
import re
import shutil
import subprocess
import sys

sys.path.insert(0, "/w/check")
import lx3_data as D
from _common import OUT, Score

XLSX, DOCX, PPTX = "成绩分析.xlsx", "致家长的一封信.docx", "期末家长会.pptx"
TURNS = 14
METRICS = ["参考人数", "平均分", "最高分", "最低分", "及格率", "优秀率"]
BAND_LABELS = [b[0] for b in D.BANDS]
QUALIFIER = re.compile("原定|原来|原计划|调整|改为|改到|改成|提前|推迟|变更|更改")
CJK = re.compile(r"[一-鿿]")


def stage(n: int) -> int:
    return 0 if n < 4 else 1 if n < 11 else 2


def norm(text: str) -> str:
    text = re.sub(r"\s+", "", text or "").replace("(", "（").replace(")", "）")
    return text.replace("％", "%").replace("～", "-").replace("~", "-").replace("—", "-").replace("–", "-").replace("－", "-")


# ---------- workbooks ----------


def recalc(paths: list[str]) -> dict[str, str]:
    """Recalculate every workbook in one LibreOffice run; path → recalculated copy."""
    src, out, mapping = os.path.abspath("recalc-in"), os.path.abspath("recalc-out"), {}
    os.makedirs(src, exist_ok=True)
    os.makedirs(out, exist_ok=True)
    todo = []
    for path in paths:
        with open(path, "rb") as fh:
            key = hashlib.sha1(fh.read()).hexdigest()[:16] + ".xlsx"
        if not os.path.exists(os.path.join(src, key)):
            shutil.copy(path, os.path.join(src, key))
            todo.append(os.path.join(src, key))
        mapping[path] = os.path.join(out, key)
    if todo and shutil.which("soffice"):
        subprocess.run(
            ["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", out, *todo],
            capture_output=True, timeout=240, check=False, env={**os.environ, "HOME": os.path.abspath(".")},
        )
    return {p: (q if os.path.exists(q) else p) for p, q in mapping.items()}


def sheet(wb, key: str):
    exact = [ws for ws in wb.worksheets if ws.title.strip() == key]
    return exact[0] if exact else next((ws for ws in wb.worksheets if key in ws.title), None)


def grid(ws, rows: list[str], cols: list[str]) -> dict[tuple[str, str], str] | None:
    """{(row label, col label): cell coordinate} for a labelled table anywhere on the sheet."""
    if ws is None:
        return None
    cells = list(ws.iter_rows())
    for i, row in enumerate(cells):
        texts = [norm(str(c.value)) if c.value is not None else "" for c in row]
        where = {}
        for want in cols:
            hit = next((j for j, t in enumerate(texts) if t == want), None)
            if hit is None:
                hit = next((j for j, t in enumerate(texts) if want in t), None)
            if hit is None:
                break
            where[want] = hit
        else:
            out = {}
            for below in cells[i + 1:]:
                if not below:
                    continue
                label = next((norm(str(c.value)) for c in below if c.value not in (None, "")), "")
                for r in rows:
                    if label == r or label.startswith(r):
                        for want, j in where.items():
                            if j < len(below):
                                out.setdefault((r, want), below[j].coordinate)
            if all((r, c) in out for r in rows for c in cols):
                return out
    return None


def number(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    text = norm(str(v)).rstrip("%")
    try:
        return float(text) / (100 if str(v).strip().endswith("%") else 1)
    except ValueError:
        return None


def rate_ok(v, truth) -> bool:
    x = number(v)
    if x is None:
        return False
    x = x * 100 if x <= 1.0001 else x
    return abs(x - float(truth)) <= 0.051


def avg_ok(v, truth) -> bool:
    x = number(v)
    return x is not None and abs(x - float(truth)) <= 0.051


class Book:
    def __init__(self, path: str, values_path: str):
        from openpyxl import load_workbook

        self.f = load_workbook(path, data_only=False)
        self.v = load_workbook(values_path, data_only=True)

    def table(self, key, rows, cols):
        ws_f, ws_v = sheet(self.f, key), sheet(self.v, key)
        where = grid(ws_f, rows, cols)
        if where is None or ws_v is None:
            return None
        return {k: (ws_v[c].value, ws_f[c].value) for k, c in where.items()}

    def stats(self):
        return self.table("统计", METRICS, D.SUBJECTS)

    def log_rounds(self) -> set[int]:
        ws = sheet(self.v, "修改记录")
        found = set()
        for row in ws.iter_rows(values_only=True) if ws is not None else []:
            for c in row:
                m = re.fullmatch(r"R(\d+)", norm(str(c or "")).upper())
                if m:
                    found.add(int(m.group(1)))
        return found


def is_formula(f) -> bool:
    return isinstance(f, str) and f.startswith("=")


def stats_problems(t, st) -> list[str]:
    if t is None:
        return ["统计 table (指标 × 五科) not found"]
    bad = []
    for s in D.SUBJECTS:
        for m in METRICS:
            v, _f = t[(m, s)]
            want = st[s][m]
            ok = avg_ok(v, want) if m == "平均分" else rate_ok(v, want) if m.endswith("率") else number(v) == want
            if not ok:
                bad.append(f"{s}{m} {v!r}≠{want}")
    return bad


# ---------- letter and slides ----------


class Letter:
    def __init__(self, path: str):
        from docx import Document

        doc = Document(path)
        self.paras = [p.text for p in doc.paragraphs]
        self.tables = [[[c.text for c in r.cells] for r in t.rows] for t in doc.tables]
        self.header = ""
        for sec in doc.sections:
            for part in (sec.header, sec.first_page_header, sec.even_page_header, sec.footer):
                try:
                    self.header += "\n".join(p.text for p in part.paragraphs)
                    self.header += "\n".join(c.text for t in part.tables for r in t.rows for c in r.cells)
                except (AttributeError, ValueError):
                    pass
        cells = "\n".join(c for t in self.tables for r in t for c in r)
        self.text = "\n".join(self.paras) + "\n" + cells + "\n" + self.header


class Deck:
    def __init__(self, path: str):
        from pptx import Presentation

        self.slides = []
        for slide in Presentation(path).slides:
            texts, tables, title = [], [], ""
            for shape in _shapes(slide.shapes):
                if shape.has_text_frame and shape.text_frame.text.strip():
                    texts.append((shape, shape.text_frame.text))
                if getattr(shape, "has_table", False) and shape.has_table:
                    tables.append([[c.text for c in r.cells] for r in shape.table.rows])
            if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
                title = slide.shapes.title.text_frame.text
            elif texts:
                title = min(texts, key=lambda t: (t[0].top or 0))[1]
            notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide and slide.notes_slide.notes_text_frame else ""
            body = "\n".join(t for _s, t in texts) + "\n" + "\n".join(c for t in tables for r in t for c in r)
            self.slides.append({"title": title, "body": body, "tables": tables, "notes": notes})
        self.text = "\n".join(s["body"] + "\n" + s["notes"] for s in self.slides)


def _shapes(shapes):
    for shape in shapes:
        if shape.shape_type == 6:  # group
            yield from _shapes(shape.shapes)
        else:
            yield shape


def numbers_problems(text: str, n_stage: int) -> list[str]:
    q = D.quoted(n_stage)
    flat = norm(text)
    missing = [f"{k}={v}" for k, v in q.items() if v not in flat]
    stale = sorted(
        {old for k in range(n_stage) for key, old in D.quoted(k).items() if old != q[key] and old not in q.values() and old in flat}
    )
    out = []
    if missing:
        out.append("missing " + ",".join(missing[:4]))
    if stale:
        out.append("stale " + ",".join(stale[:4]))
    return out


def names_in(text: str) -> list[str]:
    return [n for n in D.names() if n in text]


def unqualified(text: str, old: str) -> bool:
    flat = norm(text)
    for m in re.finditer(re.escape(old), flat):
        window = flat[max(0, m.start() - 14): m.end() + 14]
        if not QUALIFIER.search(window):
            return True
    return False


def band_counts(tables, labels) -> dict[str, int]:
    """Counts next to (or under) each band label in any table."""
    got = {}
    for t in tables:
        rows = [[norm(c) for c in r] for r in t]
        for i, r in enumerate(rows):
            for j, c in enumerate(r):
                lab = next((b for b in labels if c.startswith(b) or c.replace("分", "") == b), None)
                if lab is None or lab in got:
                    continue
                for cand in ([r[j + 1]] if j + 1 < len(r) else []) + ([rows[i + 1][j]] if i + 1 < len(rows) and j < len(rows[i + 1]) else []):
                    m = re.fullmatch(r"(\d+)人?", cand)
                    if m:
                        got[lab] = int(m.group(1))
                        break
    return got


def cjk(text: str) -> int:
    return len(CJK.findall(text or ""))


# ---------- main ----------


def files_at(n: int, carry: dict[str, str]) -> dict[str, str]:
    tdir = os.path.join(OUT, "turns", f"R{n:02d}")
    now = dict(carry)
    for name in (XLSX, DOCX, PPTX):
        path = os.path.join(tdir, name)
        if os.path.exists(path):
            now[name] = path
    return now


def per_turn(sc: Score) -> None:
    hard, nums = [], []
    carry: dict[str, str] = {}
    snaps = []
    for n in range(1, TURNS + 1):
        delivered_x = os.path.exists(os.path.join(OUT, "turns", f"R{n:02d}", XLSX))
        carry = files_at(n, carry)
        snaps.append((n, delivered_x, dict(carry)))
    books = recalc(sorted({s[XLSX] for _n, _d, s in snaps if XLSX in s}))
    for n, delivered_x, cur in snaps:
        st = stage(n)
        if not delivered_x:
            hard.append(f"R{n} {XLSX} not delivered")
        if XLSX in cur:
            try:
                book = Book(cur[XLSX], books[cur[XLSX]])
                t = book.stats()
                if t is None or not all(is_formula(f) for _v, f in t.values()):
                    hard.append(f"R{n} 统计 not all formulas")
                if n not in book.log_rounds():
                    hard.append(f"R{n} 修改记录 has no R{n}")
                if t is not None and any(not avg_ok(t[("平均分", s)][0], D.stats(st)[s]["平均分"]) for s in D.SUBJECTS):
                    nums.append(f"R{n} xlsx 平均分 not stage {st}")
            except Exception as exc:  # noqa: BLE001 — a broken file is a finding, not a crash
                hard.append(f"R{n} {XLSX} unreadable {type(exc).__name__}")
        for name, since, cls in ((DOCX, 2, Letter), (PPTX, 3, Deck)):
            if n < since:
                continue
            if name not in cur:
                hard.append(f"R{n} no {name} yet")
                continue
            try:
                text = cls(cur[name]).text
            except Exception as exc:  # noqa: BLE001
                hard.append(f"R{n} {name} unreadable {type(exc).__name__}")
                continue
            if names_in(text):
                hard.append(f"R{n} {name} names {names_in(text)[:2]}")
            nums += [f"R{n} {name} {p}" for p in numbers_problems(text, st)]
    sc.check(not hard, f"per-turn turn-1 rules: {'; '.join(hard[:8])}")
    sc.check(not nums, f"per-turn numbers follow the corrections: {'; '.join(nums[:8])}")


def main():
    sc = Score()
    paths = {name: os.path.join(OUT, name) for name in (XLSX, DOCX, PPTX)}
    have = {k: os.path.exists(p) for k, p in paths.items()}
    st2 = D.stats(2)

    book = None
    if have[XLSX]:
        try:
            book = Book(paths[XLSX], recalc([paths[XLSX]])[paths[XLSX]])
        except Exception as exc:  # noqa: BLE001
            sc.notes.append(f"{XLSX} unreadable {type(exc).__name__}")
    t = book.stats() if book else None
    bad = stats_problems(t, st2)
    sc.check(not bad, f"统计 vs truth after both corrections: {bad[:5]}")
    bands = book.table("分数段", BAND_LABELS, D.SUBJECTS) if book else None
    mid = book.table("期中对比", D.SUBJECTS, ["期中平均分", "期末平均分", "变化"]) if book else None
    formulas = t is not None and all(is_formula(f) for _v, f in t.values())
    formulas &= bands is not None and all(is_formula(f) for _v, f in bands.values())
    formulas &= mid is not None and all(is_formula(mid[(s, "期中平均分")][1]) for s in D.SUBJECTS)
    sc.check(formulas, "统计 / 分数段 / 期中对比 statistics are not all formulas (turn-1 rule 3)")
    want_bands = D.bands(2)
    bad_bands = [] if bands else ["分数段 table not found"]
    for (lab, s), (v, _f) in (bands or {}).items():
        if number(v) != want_bands[s][lab]:
            bad_bands.append(f"{s}{lab} {v!r}≠{want_bands[s][lab]}")
    sc.check(not bad_bands, f"R6 分数段: {bad_bands[:5]}")
    mids = D.mid_avg()
    bad_mid = [] if mid else ["期中对比 table not found"]
    for s in D.SUBJECTS if mid else []:
        a, b, c = (mid[(s, k)][0] for k in ("期中平均分", "期末平均分", "变化"))
        if not (avg_ok(a, mids[s]) and avg_ok(b, st2[s]["平均分"]) and avg_ok(c, st2[s]["平均分"] - mids[s])):
            bad_mid.append(f"{s} {a!r}/{b!r}/{c!r} want {mids[s]}/{st2[s]['平均分']}")
    sc.check(not bad_mid, f"R8 期中对比 (only students still here, columns reordered): {bad_mid[:4]}")

    letter = deck = None
    try:
        letter = Letter(paths[DOCX]) if have[DOCX] else None
    except Exception as exc:  # noqa: BLE001
        sc.notes.append(f"{DOCX} unreadable {type(exc).__name__}")
    try:
        deck = Deck(paths[PPTX]) if have[PPTX] else None
    except Exception as exc:  # noqa: BLE001
        sc.notes.append(f"{PPTX} unreadable {type(exc).__name__}")
    lt = letter.text if letter else ""
    dt = deck.text if deck else ""

    p = numbers_problems(lt, 2) if letter else ["no letter"]
    sc.check(not p, f"letter numbers after both corrections: {p}")
    p = numbers_problems(dt, 2) if deck else ["no deck"]
    flat = norm(dt)
    if deck and ("41人" not in flat or "42人" in flat):
        p.append("class size not 41人")
    if deck:
        counts = band_counts([t for s in deck.slides for t in s["tables"]], BAND_LABELS)
        if counts != want_bands["数学"]:
            p.append(f"数学分数段 table {counts} want {want_bands['数学']}")
        miss = [f"{s}{v}" for s, v in mids.items() if str(v) not in flat]
        if miss:
            p.append(f"期中 averages missing {miss[:3]}")
    sc.check(not p, f"deck numbers / tables: {p}")

    lflat = norm(lt)
    p = []
    if letter:
        if not ("滨江市第三中学" in norm(letter.header) and "八年级（2）班" in norm(letter.header)):
            p.append(f"页眉 {letter.header[:40]!r}")
        paras = [norm(x) for x in letter.paras if x.strip()]
        if not any(a == norm("八年级（2）班 班主任 林晓") and b == "2027年1月15日" for a, b in itertools.pairwise(paras)):
            p.append("落款 two lines")
        if "回执" not in lflat or not re.search("签字|签名", lflat):
            p.append("回执 with signature")
        table = next((t for t in letter.tables if "日期" in norm("".join(t[0])) and "事项" in norm("".join(t[0]))), None)
        tflat = norm("".join(c for r in table for c in r)) if table else ""
        need = [d for d in ("1月23日", "2月21日", "2月22日", "1月21日") if d not in tflat]
        if table is None or need:
            p.append(f"寒假时间安排 table missing {need or 'table'}")
    else:
        p.append("no letter")
    sc.check(not p, f"letter structure (页眉 / 落款 / 回执 / 时间表): {p}")
    p = []
    if "2月21日" not in lflat or unqualified(lt, "2月20日"):
        p.append("返校 not 2月21日 (notice's later 补充)")
    if "1月21日" not in lflat or unqualified(lt, "1月22日"):
        p.append("家长会 not moved to 1月21日")
    sentences = re.split(r"[。！\n]", lflat)
    if not any("物理" in s and re.search("进步最|提高最|提升最|最大", s) for s in sentences):
        p.append("no sentence naming 物理 as most improved")
    sc.check(not p, f"letter content: {p}")

    p = []
    if deck and deck.slides:
        cover = norm(deck.slides[0]["body"])
        if not ("滨江市第三中学" in cover and "期末家长会" in cover and "1月21日" in cover):
            p.append(f"cover {cover[:40]!r}")
        if unqualified(dt, "1月22日"):
            p.append("deck still says 1月22日")
        miss = [b for b in D.BOOKS if b not in flat]
        if miss:
            p.append(f"推荐书目 missing {miss}")
        bare = [i + 1 for i, s in enumerate(deck.slides) if len(norm(s["notes"])) < 50]
        if bare:
            p.append(f"slides without notes {bare[:5]}")
    else:
        p.append("no deck")
    sc.check(not p, f"deck structure (cover / 书目 / notes): {p}")
    over = [
        f"{i + 1}:{cjk(s['body']) - cjk(s['title'])}"
        for i, s in enumerate(deck.slides if deck else [])
        if cjk(s["body"]) - cjk(s["title"]) > 80
    ]
    sc.check(deck is not None and not over, f"slides over 80 汉字 of body (家委会 rule): {over[:5]}")
    leaks = [f"letter:{n}" for n in names_in(lt)] + [f"deck:{n}" for n in names_in(dt)]
    sc.check(bool(letter and deck) and not leaks, f"student names in parent files (turn-1 rule 2): {leaks[:5]}")
    rounds = book.log_rounds() if book else set()
    sc.check(set(range(1, TURNS + 1)) <= rounds, f"修改记录 rounds {sorted(rounds)} want R1..R{TURNS}", essential=False)
    per_turn(sc)
    sc.emit()


if __name__ == "__main__":
    main()
