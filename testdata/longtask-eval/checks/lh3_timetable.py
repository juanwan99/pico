"""LH3: six-turn timetable — the final solver honours every rule still in force.

Points (all essential): basic rules + same seed same result · the turn-4
reversal (a plan only solvable with 体育 in the first period) · shared
teachers never double-booked · impossible plans raise ValueError in time,
naming the overloaded teacher · own tests green and 课表.xlsx matches the plan.
"""

import json
import os
import shutil
import sys

sys.path.insert(0, "/w/check")
from _common import IN, Score, find, run, tail

PROBE = r'''
import json, math, sys
group = sys.argv[1]
from timetable import make_timetable
MAIN = ("语文", "数学", "英语")
out = {"ok": True, "why": []}

def bad(msg):
    out["ok"] = False
    out["why"].append(msg[:220])

def check(tt, plan, days, periods, teachers=None, tag=""):
    teachers = teachers or {}
    if set(tt) != set(plan):
        bad(f"{tag} classes {sorted(tt)}")
        return
    booked = {}
    for c, subj in plan.items():
        g = [[(x or None) for x in row] for row in tt[c]]
        if len(g) != days or any(len(r) != periods for r in g):
            bad(f"{tag} {c} grid shape")
            return
        cnt, am = {}, {}
        for d, row in enumerate(g):
            day = {}
            for p, s in enumerate(row):
                if s is None:
                    continue
                cnt[s] = cnt.get(s, 0) + 1
                day[s] = day.get(s, 0) + 1
                if p < 4:
                    am[s] = am.get(s, 0) + 1
                if s == "体育" and (p == periods - 1 or p in (3, 4)):
                    bad(f"{tag} {c} 体育 at day{d} p{p}")
                t = teachers.get(c, {}).get(s)
                if t:
                    if (t, d, p) in booked:
                        bad(f"{tag} {t} double-booked day{d} p{p} ({booked[(t, d, p)]}, {c})")
                    booked[(t, d, p)] = c
            for s, n in day.items():
                if n > 2:
                    bad(f"{tag} {c} {s} x{n} on day{d}")
                elif n == 2 and row[row.index(s) + 1: row.index(s) + 2] != [s]:
                    bad(f"{tag} {c} {s} twice on day{d} but not back to back")
        if cnt != {s: h for s, h in subj.items() if h}:
            bad(f"{tag} {c} hours {cnt} != plan")
        for s in MAIN:
            if s in subj and am.get(s, 0) < math.ceil(subj[s] / 2):
                bad(f"{tag} {c} {s} morning {am.get(s, 0)} < {math.ceil(subj[s] / 2)}")

P1 = {"语文": 6, "数学": 6, "英语": 5, "体育": 3, "物理": 3, "音乐": 2, "美术": 2}
try:
    if group == "basic":
        plan = {"七1": dict(P1), "七2": dict(P1)}
        a = make_timetable(plan, seed=5)
        check(a, plan, 5, 6, tag="basic")
        if make_timetable(plan, seed=5) != a:
            bad("same seed, different timetable")
        try:
            make_timetable({"x": {"语文": 31}})
            bad("31 lessons in 30 slots accepted")
        except ValueError:
            pass
    elif group == "reversal":
        plan = {"x": {"体育": 7, "语文": 4, "数学": 4}}
        check(make_timetable(plan, days=5, periods=3), plan, 5, 3, tag="3-period day")
    elif group == "teachers":
        P = {"语文": 6, "数学": 6, "英语": 5, "体育": 3, "物理": 3, "生物": 2, "音乐": 1, "美术": 1}
        plan = {c: dict(P) for c in ("七1", "七2", "七3")}
        teachers = {
            c: {"数学": "张老师", "英语": "李老师", "语文": c + "语文", "物理": "赵老师" if c != "七3" else "钱老师"}
            for c in plan
        }
        check(make_timetable(plan, teachers=teachers, seed=3), plan, 5, 6, teachers, tag="teachers")
    elif group == "infeasible":
        plan = {c: {"语文": 6, "数学": 6} for c in "abc"}
        try:
            make_timetable(plan, teachers={c: {"语文": "王老师", "数学": "王老师"} for c in "abc"})
            bad("王老师 36 lessons in 30 slots accepted")
        except ValueError as exc:
            if "王老师" not in str(exc):
                bad(f"overload error does not name the teacher: {exc}")
        try:
            make_timetable({"a": {"体育": 11, "语文": 5}})
            bad("11 体育 lessons (max 10 legal slots) accepted")
        except ValueError:
            pass
except Exception as exc:
    bad(f"{type(exc).__name__}: {exc}")
print(json.dumps(out, ensure_ascii=False))
'''


def probe(work, group):
    proc = run([sys.executable, "probe.py", group], cwd=work, timeout=150)
    try:
        got = json.loads(proc.stdout.strip().splitlines()[-1])
    except (ValueError, IndexError):
        return False, f"{group}: crashed/timeout rc={proc.returncode} {tail(proc.stderr)}"
    return got["ok"], f"{group}: " + " | ".join(got["why"][:3])


def sheet_ok(path):
    from openpyxl import load_workbook

    with open(os.path.join(IN, "school.json"), encoding="utf-8") as fh:
        school = json.load(fh)
    wb = load_workbook(path, data_only=True)
    why = []
    for cls, subj in school["plan"].items():
        ws = next((w for w in wb.worksheets if cls in w.title), None)
        if ws is None:
            why.append(f"no sheet for {cls}")
            continue
        cells = [str(v).strip() for row in ws.iter_rows(values_only=True) for v in row if v not in (None, "")]
        for day in ("星期一", "星期五"):
            if day not in cells:
                why.append(f"{cls} has no {day} header")
        cnt = {}
        for v in cells:
            s = v.split("/")[0].split("（")[0].split("(")[0].strip()
            if s in subj:
                cnt[s] = cnt.get(s, 0) + 1
        if cnt != subj:
            why.append(f"{cls} counts {cnt}")
        if not any("/" in v and "老师" in v for v in cells):
            why.append(f"{cls} cells lack 科目/老师")
    return not why, "; ".join(why[:3])


def main():
    sc = Score()
    src = find("timetable.py")
    if not src:
        sc.fail("no timetable.py delivered")
        sc.emit()
    work = os.path.abspath("work")
    os.makedirs(work, exist_ok=True)
    for name in os.listdir(IN):
        shutil.copy(os.path.join(IN, name), work)
    for name in ("timetable.py", "test_timetable.py"):
        path = find(name)
        if path:
            shutil.copy(path, work)
    with open(os.path.join(work, "probe.py"), "w", encoding="utf-8") as fh:
        fh.write(PROBE)
    for group in ("basic", "reversal", "teachers", "infeasible"):
        ok, note = probe(work, group)
        sc.check(ok, note)
    why = []
    if not os.path.exists(os.path.join(work, "test_timetable.py")):
        why.append("no test_timetable.py")
    else:
        own = run([sys.executable, "-m", "unittest", "test_timetable"], cwd=work, timeout=240)
        if own.returncode != 0:
            why.append(f"own tests fail {tail(own.stderr)}")
    book = find("课表.xlsx")
    if not book:
        why.append("no 课表.xlsx")
    else:
        ok, note = sheet_ok(book)
        if not ok:
            why.append(note)
    sc.check(not why, "deliver: " + " | ".join(why))
    sc.emit()


main()
