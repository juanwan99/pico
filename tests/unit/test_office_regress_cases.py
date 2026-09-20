"""#1046: office-regress ships 10 cases; fixtures open as real OOXML."""

from __future__ import annotations

import importlib.util
import io
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
_SPEC = importlib.util.spec_from_file_location(
    "office_regress_script", ROOT / "scripts" / "office-regress.py"
)
assert _SPEC and _SPEC.loader
orx = importlib.util.module_from_spec(_SPEC)
sys.modules["office_regress_script"] = orx
_SPEC.loader.exec_module(orx)


def test_ten_named_cases() -> None:
    assert list(orx.CASES) == ["t1", "t2", "t3", "t4", "t5", "t6", "t7", "t8", "t9", "t10"]


def test_week_deck_has_three_marked_slides() -> None:
    raw = orx.make_week_deck()
    blob = orx._slide_blob(raw)
    assert orx._pptx_slide_count(raw) == 3
    assert "slides=3" in blob
    assert "KEEP-封面" in blob and "KEEP-课表" in blob and "KEEP-作业" in blob
    xml_only = orx._pptx_xml_text(raw)
    assert "KEEP-封面" in xml_only


def test_notice_docx_has_spring_and_safety() -> None:
    blob = orx._docx_blob(orx.make_notice_docx())
    assert "春游" in blob and "安全" in blob and "通知" in blob


def test_formula_book_keeps_formula_and_note() -> None:
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(orx.make_formula_book())).active
    assert str(ws["D2"].value).startswith("=")
    assert ws["E2"].value == "原备注留着"


def test_chart_book_is_xlsx() -> None:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(orx.make_chart_book()))
    assert wb.active["A2"].value == "甲"
    assert wb.active._charts


def test_docx_media_helper_false_on_plain() -> None:
    assert orx._docx_has_media(orx.make_notice_docx()) is False


def test_xlsx_group_counts_reads_a_zu_layout() -> None:
    from openpyxl import Workbook

    wb = Workbook()
    summary = wb.active
    summary.title = "组别人数汇总"
    summary.append(["组别", "人数"])
    summary.append(["A组", 5])
    summary.append(["B组", 3])
    summary.append(["C组", 2])
    detail = wb.create_sheet("花名册明细")
    detail.append(["学号", "姓名", "组别"])
    detail.append([2401, "张一", "A"])
    buf = io.BytesIO()
    wb.save(buf)
    assert orx._xlsx_group_counts(buf.getvalue()) == {"A": 5, "B": 3, "C": 2}
