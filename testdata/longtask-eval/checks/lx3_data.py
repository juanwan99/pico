"""LX3 truth: one class's 期末 / 期中 marks and the two rounds of corrections (#1151).

Seeded, so the eval (attachments) and the checker (truth) build the same data.
Stage 0 is the 期末 sheet as uploaded in R1; stage 1 applies 教务处's R4 更正单
(a 数学 mark typed wrong, a 物理 缺考 sat the make-up, one student transferred
out); stage 2 applies R11's second 更正 (an 英语 mark typed wrong). The letter
and slides quote each subject's 平均分 and 及格率; the seed is picked so the
corrections move 数学 / 物理 (R4) and 英语 (R11), no 平均分 is a whole number,
and no replaced number equals a final one.
"""

from __future__ import annotations

import io
import random
import re
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

SEED = 1190  # first seed from 1151 where sound() holds
SUBJECTS = ["语文", "数学", "英语", "物理", "历史"]
MID_ORDER = ["数学", "语文", "英语", "物理", "历史"]  # 期中 sheet's columns, on purpose
N = 42
PASS, GOOD = 60, 85
BANDS = [("90-100", 90, 100), ("80-89", 80, 89), ("70-79", 70, 79), ("60-69", 60, 69), ("0-59", 0, 59)]

SURNAMES = "王李张刘陈杨黄赵吴周徐孙马朱胡郭何罗郑梁谢宋唐许韩冯邓曹彭曾萧田董袁潘蒋蔡余杜叶程苏魏吕丁任沈姚卢姜崔钟谭陆汪范金石廖贾夏韦付方白邹孟熊秦邱江尹薛闫段雷侯龙史陶黎贺顾毛郝龚邵万钱严覃武戴莫孔向汤"
GIVEN = "子梓浩宇轩涵欣怡晨思雨佳嘉俊博文泽睿诗琪一心若可天乐语铭奕然景舒逸辰沐昊宸瑞霖钰彤妍菲琳瑶楠航哲"

TEACHER = "林晓"
CLASS = "八年级（2）班"
SCHOOL = "滨江市第三中学"
BOOKS = ["朝花夕拾", "骆驼祥子", "红星照耀中国", "昆虫记", "海底两万里"]


def _r1(x: float) -> Decimal:
    return Decimal(repr(x)).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP)


def roster(seed: int = SEED) -> list[tuple[str, str]]:
    rng = random.Random(seed)
    names: list[str] = []
    while len(names) < N:
        name = rng.choice(SURNAMES) + rng.choice(GIVEN) + rng.choice(GIVEN)
        if name not in names and name[1] != name[2] and TEACHER not in name:
            names.append(name)
    return [(f"202508{i + 1:02d}", n) for i, n in enumerate(names)]


# Who the corrections touch (index into roster), fixed so the prompts can name them.
MATH_FIX = (5, 58, 63)  # 数学 58 typed for 63
PHYS_MAKEUP = (17, 71)  # 物理 缺考 → make-up 71
TRANSFER = 30  # transferred out before the 期末 results went home
ENG_FIX = (9, 92, 82)  # R11: 英语 92 typed for 82
MID_ABSENT = 12  # off sick for the whole 期中


def _marks(seed: int) -> dict[str, dict[str, int | str]]:
    rng = random.Random(seed * 7 + 1)
    out: dict[str, dict[str, int | str]] = {}
    for i, (sid, _name) in enumerate(roster(seed)):
        ability = rng.gauss(0, 9)
        out[sid] = {s: max(18, min(100, round(rng.gauss(75 + ability, 9)))) for s in SUBJECTS}
    sids = list(out)
    out[sids[MATH_FIX[0]]]["数学"] = MATH_FIX[1]
    out[sids[PHYS_MAKEUP[0]]]["物理"] = "缺考"
    out[sids[ENG_FIX[0]]]["英语"] = ENG_FIX[1]
    return out


def marks(stage: int, seed: int = SEED) -> dict[str, dict[str, int | str]]:
    """{学号: {科目: 分 or "缺考"}} of the class at a stage (0 uploaded, 1 after R4, 2 after R11)."""
    m = {sid: dict(v) for sid, v in _marks(seed).items()}
    sids = [sid for sid, _ in roster(seed)]
    if stage >= 1:
        m[sids[MATH_FIX[0]]]["数学"] = MATH_FIX[2]
        m[sids[PHYS_MAKEUP[0]]]["物理"] = PHYS_MAKEUP[1]
        del m[sids[TRANSFER]]
    if stage >= 2:
        m[sids[ENG_FIX[0]]]["英语"] = ENG_FIX[2]
    return m


def mid_marks(seed: int = SEED) -> dict[str, dict[str, int | None]]:
    """期中: everyone incl. the later transfer; MID_ABSENT has no marks."""
    rng = random.Random(seed * 11 + 3)
    fin = _marks(seed)
    out: dict[str, dict[str, int | None]] = {}
    for i, sid in enumerate(fin):
        if i == MID_ABSENT:
            out[sid] = {s: None for s in SUBJECTS}
            continue
        out[sid] = {}
        for s in SUBJECTS:
            base = fin[sid][s] if isinstance(fin[sid][s], int) else 70
            out[sid][s] = max(15, min(100, round(base + rng.gauss(-2 if s in ("数学", "物理") else 1, 6))))
    return out


def stats(stage: int, seed: int = SEED) -> dict[str, dict[str, Decimal | int]]:
    """Per subject: 参考人数, 平均分, 最高分, 最低分, 及格率, 优秀率 (rates in %, 1 decimal)."""
    out = {}
    m = marks(stage, seed)
    for s in SUBJECTS:
        v = [r[s] for r in m.values() if isinstance(r[s], int)]
        out[s] = {
            "参考人数": len(v),
            "平均分": _r1(sum(v) / len(v)),
            "最高分": max(v),
            "最低分": min(v),
            "及格率": _r1(sum(1 for x in v if x >= PASS) * 100 / len(v)),
            "优秀率": _r1(sum(1 for x in v if x >= GOOD) * 100 / len(v)),
        }
    return out


def raw_stats(stage: int, seed: int = SEED) -> dict[str, dict[str, float]]:
    m = marks(stage, seed)
    out = {}
    for s in SUBJECTS:
        v = [r[s] for r in m.values() if isinstance(r[s], int)]
        out[s] = {
            "平均分": sum(v) / len(v),
            "及格率": sum(1 for x in v if x >= PASS) * 100 / len(v),
            "优秀率": sum(1 for x in v if x >= GOOD) * 100 / len(v),
        }
    return out


def bands(stage: int, seed: int = SEED) -> dict[str, dict[str, int]]:
    m = marks(stage, seed)
    return {
        s: {lab: sum(1 for r in m.values() if isinstance(r[s], int) and lo <= r[s] <= hi) for lab, lo, hi in BANDS}
        for s in SUBJECTS
    }


def mid_avg(seed: int = SEED) -> dict[str, Decimal]:
    """期中 均分 over the students still in class (R8 rule), 缺考 / blank not counted."""
    now = marks(1, seed)
    mid = mid_marks(seed)
    out = {}
    for s in SUBJECTS:
        v = [mid[sid][s] for sid in now if mid[sid][s] is not None]
        out[s] = _r1(sum(v) / len(v))
    return out


def quoted(stage: int, seed: int = SEED) -> dict[str, str]:
    """Strings the letter and slides must carry at a stage: '数学平均分' → '78.4', '数学及格率' → '92.5%'."""
    st = stats(stage, seed)
    out = {}
    for s in SUBJECTS:
        out[f"{s}平均分"] = str(st[s]["平均分"])
        out[f"{s}及格率"] = f"{st[s]['及格率']}%"
    return out


def edgy(stage: int, seed: int = SEED) -> bool:
    """A quoted value whose rounding a float slip could flip (sum/41, sum/42 never hit .x5 exactly)."""
    for row in raw_stats(stage, seed).values():
        for k in ("平均分", "及格率"):
            if abs((row[k] * 10) % 1 - 0.5) < 0.001:
                return True
    return False


def changed(seed: int = SEED) -> dict[str, str]:
    """Quoted values the corrections replaced: {old string: key}, none equal to any final value."""
    q = [quoted(k, seed) for k in range(3)]
    return {q[k][key]: key for k in (0, 1) for key in q[k] if q[k][key] != q[2][key]}


def sound(seed: int) -> bool:
    """Corrections show in the quoted numbers, and an old number never reads as a new one."""
    q = [quoted(k, seed) for k in range(3)]
    if any(edgy(k, seed) for k in range(3)):
        return False
    if any(q[k][f"{s}平均分"].endswith(".0") for k in range(3) for s in SUBJECTS):
        return False  # "75.0" vs "75"
    if set(changed(seed)) & set(q[2].values()):
        return False
    if q[0]["数学平均分"] == q[1]["数学平均分"] or q[0]["物理平均分"] == q[1]["物理平均分"]:
        return False
    if q[1]["英语平均分"] == q[2]["英语平均分"]:
        return False
    if {str(v) for v in mid_avg(seed).values()} & {v for k in range(3) for v in q[k].values()}:
        return False
    sids = [sid for sid, _ in roster(seed)]
    return marks(0, seed)[sids[ENG_FIX[0]]]["英语"] >= GOOD


def final_book() -> bytes:
    """R1 attachment 期末成绩.xlsx (stage 0)."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "期末"
    ws.append(["学号", "姓名", *SUBJECTS])
    m = marks(0)
    for sid, name in roster():
        ws.append([sid, name, *[m[sid][s] for s in SUBJECTS]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def mid_book() -> bytes:
    """R8 attachment 期中成绩.xlsx: other column order, the transfer still on it, one blank row."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "期中"
    ws.append(["学号", "姓名", *MID_ORDER])
    mid = mid_marks()
    for sid, name in roster():
        ws.append([sid, name, *[mid[sid][s] for s in MID_ORDER]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _docx_bytes(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def notice_docx() -> bytes:
    """R2 attachment: the school's 寒假 notice (fixtures/lx3/寒假放假通知.md) as Word."""
    from docx import Document

    src = Path(__file__).resolve().parent.parent / "fixtures" / "lx3" / "寒假放假通知.md"
    doc = Document()
    for line in src.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line == "---":
            continue
        if line.startswith("#"):
            doc.add_heading(line.lstrip("#").strip(), level=1 if line.startswith("# ") else 2)
            continue
        para = doc.add_paragraph()
        for i, part in enumerate(re.split(r"\*\*", line)):
            para.add_run(part).bold = i % 2 == 1
    return _docx_bytes(doc)


def fix_docx(n: int) -> bytes:
    """R4 / R11 attachments: 教务处's 成绩更正单 (一) / (二)."""
    from docx import Document

    rows = []
    if n == 1:
        sid, name = who(MATH_FIX[0])
        rows.append((sid, name, "数学", str(MATH_FIX[1]), str(MATH_FIX[2]), "登分错误，以试卷原始分为准"))
        sid, name = who(PHYS_MAKEUP[0])
        rows.append((sid, name, "物理", "缺考", str(PHYS_MAKEUP[1]), "1 月 12 日病愈补考，按规定计入期末成绩"))
        sid, name = who(TRANSFER)
        rows.append((sid, name, "全部", "—", "—", "已于 1 月 8 日转学，学籍已转出，本次成绩不再计入班级统计"))
        title, date = "成绩更正单（一）", "2027 年 1 月 14 日"
    else:
        sid, name = who(ENG_FIX[0])
        rows.append((sid, name, "英语", str(ENG_FIX[1]), str(ENG_FIX[2]), "复核试卷发现第 47 题误加 10 分"))
        title, date = "成绩更正单（二）", "2027 年 1 月 17 日"
    doc = Document()
    doc.add_heading(f"{SCHOOL} {title}", level=1)
    doc.add_paragraph(f"{CLASS}：经复核，你班期末成绩有以下更正，请据此修改成绩及相关统计。")
    table = doc.add_table(rows=1, cols=6)
    for cell, text in zip(table.rows[0].cells, ("学号", "姓名", "科目", "原登记", "更正为", "原因")):
        cell.text = text
    for row in rows:
        for cell, text in zip(table.add_row().cells, row):
            cell.text = text
    doc.add_paragraph(f"{SCHOOL} 教务处")
    doc.add_paragraph(date)
    return _docx_bytes(doc)


def names() -> list[str]:
    return [n for _, n in roster()]


def who(idx: int) -> tuple[str, str]:
    return roster()[idx]
