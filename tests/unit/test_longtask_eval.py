"""scripts/longtask-eval.py — case pack + CLI, no live network."""

from __future__ import annotations

import importlib.util
import json
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
SCRIPT = ROOT / "scripts" / "longtask-eval.py"
CASES = ROOT / "testdata" / "longtask-eval"

_spec = importlib.util.spec_from_file_location("longtask_eval", SCRIPT)
assert _spec is not None and _spec.loader is not None
lte = importlib.util.module_from_spec(_spec)
sys.modules["longtask_eval"] = lte
_spec.loader.exec_module(lte)


def test_help_exits_zero() -> None:
    proc = subprocess.run(
        [sys.executable, str(SCRIPT), "--help"],
        cwd=ROOT,
        capture_output=True,
        text=True,
        check=False,
    )
    assert proc.returncode == 0, proc.stderr
    assert "long-task" in proc.stdout.lower() or "Long-task" in proc.stdout
    assert "--base" in proc.stdout


def test_eight_cases_cover_required_shapes() -> None:
    cases = lte.load_cases("office")
    ids = [c["id"] for c in cases]
    assert ids == ["LT1", "LT2", "LT3", "LT4", "LT5", "LT6", "LT7", "LT8"]
    titles = " ".join(c["title"] for c in cases)
    assert "Word" in titles
    assert "Excel" in titles
    assert "PPT" in titles
    assert "续改" in titles
    assert "网页" in titles or "研究" in titles
    assert any(int(c.get("timeout_s") or 0) >= 1800 for c in cases)
    assert any(len(c.get("turns") or []) >= 3 for c in cases)
    assert any(
        "web" in json.dumps(c, ensure_ascii=False).lower() or "上网" in json.dumps(c, ensure_ascii=False)
        for c in cases
    )
    for c in cases:
        assert c.get("score_points"), c["id"]
        assert 1 <= len(c["score_points"]) <= 5
        for p in c["score_points"]:
            assert p.get("text")


def test_fixture_files_exist() -> None:
    for name in (
        "memo-a.md",
        "memo-b.md",
        "scores.csv",
        "unit-plan.md",
        "survey.md",
        "calendar.md",
    ):
        path = CASES / "fixtures" / name
        assert path.is_file() and path.stat().st_size > 20, name


def test_triple_workbook_has_three_sheets() -> None:
    import io

    from openpyxl import load_workbook

    raw = lte._triple_workbook()
    wb = load_workbook(io.BytesIO(raw))
    assert set(wb.sheetnames) >= {"成绩", "出勤", "问卷"}


def test_classify_fail_buckets() -> None:
    assert lte._classify_fail("failed", "run owner was lost during API restart") == "deploy_killed"
    assert lte._classify_fail("failed", "HTTP 524") == "upstream_model"
    assert lte._classify_fail("failed", "Tool not allowlisted: x") == "tool_error"
    assert lte._classify_fail("failed", "durable_max exceeded") == "wall_clock"
    assert lte._classify_fail("succeeded", "") == ""


def test_every_case_can_pass_when_all_checks_hold() -> None:
    """LT3/LT4/LT5 top out at 2 points; a fixed >=3 bar made them unpassable."""
    for case in lte.load_cases("office"):
        expect = case.get("expect") or {}
        assert 1 <= lte.pass_bar(expect) <= lte.max_points(expect), case["id"]
    by_id = {c["id"]: c.get("expect") or {} for c in lte.load_cases()}
    assert lte.max_points(by_id["LT4"]) == 2 and lte.pass_bar(by_id["LT4"]) == 2
    assert lte.pass_bar(by_id["LT6"]) == 3


def test_fake_green_flagged_explicitly() -> None:
    """Baseline #1091: 5 runs said succeeded with 0 files. Now a column, never a pass."""
    res = lte.CaseResult(case="LT9", title="t", ok=True, auto_score=3, max_score=5, artifacts=0)
    lte.flag_fake_green(res, "succeeded")
    assert res.fake_green is True and res.ok is False
    with_files = lte.CaseResult(case="LT9", title="t", ok=True, auto_score=3, artifacts=2)
    lte.flag_fake_green(with_files, "succeeded")
    assert with_files.fake_green is False and with_files.ok is True
    table = lte.render_markdown([res, with_files])
    assert "假绿" in table and "| 是 |" in table and "假绿：1" in table


def test_coding_suite_cases_and_checkers() -> None:
    """#1090 v3.1: coding sits beside office. Every LC case names a hidden checker."""
    cases = lte.load_cases("coding")
    assert [c["id"] for c in cases] == ["LC1", "LC2", "LC3", "LC4", "LC5"]
    assert [c["id"] for c in lte.load_cases()][:8] == [f"LT{i}" for i in range(1, 9)]
    assert len(lte.load_cases("office")) + len(cases) == 13
    assert any(len(c.get("turns") or []) >= 3 for c in cases)
    for c in cases:
        check = c["expect"]["check"]
        assert (lte.CHECKS_DIR / check["script"]).is_file(), c["id"]
        assert lte.max_points(c["expect"]) == len(c["expect"]["files"] and [1]) + check["max"]
        for turn in c["turns"]:
            for att in turn.get("attachments") or []:
                assert lte._attachment_bytes(att), (c["id"], att["name"])
        assert 1 <= len(c["score_points"]) <= 5


def test_hard_suite_cases_and_checkers() -> None:
    """LT/LC saturated at 13/13 on two models; LH cases exist to tell them apart."""
    cases = lte.load_cases("hard")
    assert [c["id"] for c in cases] == ["LH1", "LH2", "LH3"]
    assert [c["id"] for c in lte.load_cases()][-3:] == ["LH1", "LH2", "LH3"]
    assert any(len(c["turns"]) >= 6 for c in cases)
    for c in cases:
        check = c["expect"]["check"]
        assert (lte.CHECKS_DIR / check["script"]).is_file(), c["id"]
        assert lte.max_points(c["expect"]) == 1 + check["max"]
        assert int(c["timeout_s"]) >= 1800
        for turn in c["turns"]:
            for att in turn.get("attachments") or []:
                assert lte._attachment_bytes(att), (c["id"], att["name"])


def test_grade_book_is_seeded_and_carries_traps() -> None:
    import io

    from openpyxl import load_workbook

    def cells(raw: bytes) -> list[list[object]]:
        wb = load_workbook(io.BytesIO(raw))
        return [[c.value for c in row] for ws in wb.worksheets for row in ws.iter_rows()]

    raw = lte._grade_book()
    assert cells(raw) == cells(lte._grade_book())
    wb = load_workbook(io.BytesIO(raw))
    mid, fin = wb["期中"], wb["期末"]
    assert mid.max_row == 271 and fin.max_row == 270
    names = [r[0].value for r in mid.iter_rows(min_row=2)]
    assert len({n.strip() for n in names}) == 6 and len(set(names)) > 6
    marks = [c.value for row in fin.iter_rows(min_row=2) for c in row[3:]]
    assert "作废" in marks and "缺考" in marks and None in marks
    ids_mid = [r[1].value for r in mid.iter_rows(min_row=2)]
    ids_fin = [r[1].value for r in fin.iter_rows(min_row=2)]
    assert ids_fin != sorted(ids_fin) and len(set(ids_mid) - set(ids_fin)) == 1


def test_fixture_projects_ship_their_test_data() -> None:
    """.gitignore has ``data/``: tests/data CSVs were never committed (LH1 2026-10-01)."""
    import subprocess

    tracked = subprocess.run(
        ["git", "ls-files", "testdata/longtask-eval/fixtures/code"],
        cwd=ROOT,
        capture_output=True,
        text=True,
        check=True,
    ).stdout.split()
    for name in (
        "roster/tests/data/class.csv",
        "gradebook/tests/data/sample.csv",
        "gradebook/tests/data/attendance.csv",
    ):
        assert f"testdata/longtask-eval/fixtures/code/{name}" in tracked, name


def test_new_fixture_data_is_not_gitignored() -> None:
    """A new fixture's tests/data must not need ``git add -f`` again."""
    import subprocess

    probe = "testdata/longtask-eval/fixtures/code/new-project/tests/data/x.csv"
    rc = subprocess.run(["git", "check-ignore", "-q", probe], cwd=ROOT, check=False).returncode
    assert rc == 1, "fixture tests/data is still ignored"
    assert subprocess.run(["git", "check-ignore", "-q", "data/pico.db"], cwd=ROOT, check=False).returncode == 0


def test_zip_dir_skips_pycache(tmp_path: Path) -> None:
    import io
    import zipfile

    raw = lte._zip_dir("fixtures/code/gradebook")
    names = zipfile.ZipFile(io.BytesIO(raw)).namelist()
    assert "gradebook/gradebook/models.py" in names and "gradebook/tests/test_loader.py" in names
    assert not any("__pycache__" in n for n in names)


def test_zip_dir_is_deterministic_project() -> None:
    import io
    import zipfile

    a = lte._zip_dir("fixtures/code/roster")
    assert a == lte._zip_dir("fixtures/code/roster")
    names = zipfile.ZipFile(io.BytesIO(a)).namelist()
    assert "roster/tests/test_roster.py" in names and "roster/roster/stats.py" in names


def test_delivered_files_skip_uploads_and_summary() -> None:
    arts = [
        {"id": "1", "title": "bank.json", "kind": "file", "created_at": "1"},
        {"id": "2", "title": "roster.zip", "kind": "edu_office", "created_at": "1"},
        {"id": "3", "title": "回复摘要", "kind": "doc", "created_at": "1"},
        {"id": "4", "title": "quiz.py", "kind": "py", "created_at": "1"},
        {"id": "5", "title": "quiz.py", "kind": "py", "created_at": "2"},
    ]
    latest = lte._latest_by_title(arts)
    assert list(latest) == ["quiz.py"] and latest["quiz.py"]["id"] == "5"


def test_unzip_skips_entries_escaping_dest(tmp_path: Path) -> None:
    import io
    import zipfile

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("ok/a.py", "x = 1")
        zf.writestr("../evil.py", "boom")
    lte._unzip_into(buf.getvalue(), tmp_path / "out")
    assert (tmp_path / "out" / "ok" / "a.py").is_file()
    assert not (tmp_path / "evil.py").exists()


def test_unzip_reads_utf8_names_without_the_utf8_bit(tmp_path: Path) -> None:
    """Info-ZIP `zip -r` in the box stores UTF-8 names with flag 0 (LX2 run 1)."""
    import io
    import zipfile

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("p/XXXXXX.csv", "学号")  # same byte length as 春游 in UTF-8
    raw = buf.getvalue().replace(b"XXXXXX", "春游".encode())
    assert zipfile.ZipFile(io.BytesIO(raw)).infolist()[0].flag_bits & 0x800 == 0
    lte._unzip_into(raw, tmp_path / "out")
    assert (tmp_path / "out" / "p" / "春游.csv").read_text(encoding="utf-8") == "学号"


def test_coding_case_fails_when_files_missing() -> None:
    case = {"id": "LCX", "expect": {"files": ["a.py", "b.py"]}}
    res = lte.CaseResult(case="LCX")
    lte.score_code(case, {"a.py": b"x"}, res, "img", "runc")
    assert res.ok is False and res.auto_score == 0 and "missing files: b.py" in res.notes
    done = lte.CaseResult(case="LCX")
    lte.score_code(case, {"a.py": b"x", "b.py": b"y"}, done, "img", "runc")
    assert done.ok is True and done.auto_score == 1


def test_follow_waits_out_an_api_restart(monkeypatch) -> None:
    """#1133: the stream drops on deploy; the run is followed until it ends."""
    import httpx

    monkeypatch.setattr(lte.time, "sleep", lambda _s: None)
    replies = iter(
        [
            httpx.ConnectError("api down"),
            [{"latest_run": {"status": "running"}}],
            [{"latest_run": {"status": "succeeded"}}],
        ]
    )

    def tasks(_cid: str):
        r = next(replies)
        if isinstance(r, Exception):
            raise r
        return r

    pico = lte.Pico.__new__(lte.Pico)
    pico.tasks = tasks
    assert pico.follow("c", timeout_s=60) >= 0
    assert next(replies, None) is None


def _lx1_answer(tmp_path: Path, wrong_early: int = 0) -> Path:
    from openpyxl import Workbook

    sys.path.insert(0, str(lte.CHECKS_DIR))
    import lx_corpus

    truth = lx_corpus.students()
    early = sorted((t for t in truth if t["file"] < "21_"), key=lambda t: t["file"])
    flip = {t["id"] for t in early[:wrong_early]}
    wb = Workbook()
    ws = wb.active
    ws.title = "评语"
    ws.append(["学号", "姓名", "进步方向", "家校配合", "期末评语"])
    counts: dict[str, int] = {}
    for t in truth:
        prog = t["progress"]
        if t["id"] in flip:
            prog = next(p for p in lx_corpus.PROGRESS if p != prog)
        text = f"{t['name']}，这学期你在{prog}方面的变化老师都看在眼里，每一点努力都没有白费，继续保持这份认真和坚持，下学期一定会更好，老师相信你。"
        ws.append([t["id"], t["name"], prog, t["home"], text])
        for v in (prog, t["home"]):
            counts[v] = counts.get(v, 0) + 1
    st = wb.create_sheet("统计")
    for label in [*lx_corpus.PROGRESS, *lx_corpus.HOME]:
        st.append([label, counts.get(label, 0)])
    out = tmp_path / "期末评语.xlsx"
    wb.save(out)
    return out


def _run_lx1_checker(monkeypatch, capsys, path: Path) -> dict:
    sys.path.insert(0, str(lte.CHECKS_DIR))
    import lx1_records

    monkeypatch.setattr(lx1_records, "find", lambda name: str(path))
    try:
        lx1_records.main()
    except SystemExit:
        pass
    return json.loads(capsys.readouterr().out.strip().splitlines()[-1])


def test_xlong_suite_overflows_one_window_and_stays_out_of_all() -> None:
    """#1139: LX cases must outgrow 256k tokens so compaction fires; never in --suite all."""
    cases = lte.load_cases("xlong")
    assert [c["id"] for c in cases] == ["LX1", "LX2", "LX4"]
    assert not [c for c in lte.load_cases() if c["id"].startswith("LX")]
    sys.path.insert(0, str(lte.CHECKS_DIR))
    import lx_corpus

    rows = lx_corpus.students()
    texts = [lx_corpus.record(r) for r in rows]
    assert sum(len(t) for t in texts) > 500_000
    assert max(len(t.encode()) for t in texts) < 50_000  # one Pi read per file
    for t in texts:
        for label in [*lx_corpus.PROGRESS, *lx_corpus.HOME]:
            assert label not in t  # read, don't grep
    assert lte._attachment_bytes(cases[0]["turns"][0]["attachments"][0]) == lx_corpus.corpus_zip()
    assert lte.max_points(cases[0]["expect"]) == 6


def test_lx1_checker_full_marks_on_truth(monkeypatch, capsys, tmp_path: Path) -> None:
    verdict = _run_lx1_checker(monkeypatch, capsys, _lx1_answer(tmp_path))
    assert verdict["pass"] is True, verdict
    assert verdict["points"] == 5


def test_lx1_checker_catches_early_amnesia(monkeypatch, capsys, tmp_path: Path) -> None:
    """6 of the first 20 wrong: 92.5% overall would pass, the early bar must not."""
    verdict = _run_lx1_checker(monkeypatch, capsys, _lx1_answer(tmp_path, wrong_early=6))
    assert verdict["pass"] is False
    assert any("first20 14/20" in n for n in verdict["notes"]), verdict


def test_compaction_events_counted() -> None:
    class Fake:
        def tasks(self, cid):
            return [{"latest_run": {"id": "r1", "status": "succeeded"}}]

        def run(self, rid):
            return {"status": "succeeded"}

        def events(self, rid):
            return [{"type": "compaction.begin"}, {"type": "compaction.end"}, {"type": "compaction.failed"}]

    res = lte.CaseResult(case="LX1")
    lte._run_events(Fake(), "c", res)
    assert (res.compactions, res.compaction_failed) == (1, 1)
    assert "| 1（败 1） |" in lte.render_markdown([res])


def test_compaction_turn_recorded() -> None:
    class Fake:
        def tasks(self, cid):
            return [{"id": "t7", "latest_run": {"id": "r7"}}, {"id": "t13", "latest_run": {"id": "r13"}}]

        def run(self, rid):
            return {"status": "succeeded"}

        def events(self, rid):
            return [{"type": "compaction.end"}]

    res = lte.CaseResult(case="LX4")
    lte._run_events(Fake(), "c", res, {"t7": 7, "t13": 13})
    assert res.compaction_turns == [7, 13]
    assert "| 2 @R7/R13 |" in lte.render_markdown([res])


def test_lx4_is_lx2_with_pasted_chat() -> None:
    """#1152: same turns, truth and checker as LX2; long pastes push compaction early."""
    cases = {c["id"]: c for c in lte.load_cases("xlong")}
    lx2, lx4 = cases["LX2"], cases["LX4"]
    assert lx4["expect"] == lx2["expect"] and len(lx4["turns"]) == 15
    for a, b in zip(lx2["turns"], lx4["turns"]):
        assert a["prompt"] == b["prompt"] and a.get("attachments") == b.get("attachments")
    prompts = [lte.turn_prompt(t) for t in lx4["turns"]]
    assert prompts[0] == lx2["turns"][0]["prompt"]
    pasted = [n for n, (t, p) in enumerate(zip(lx4["turns"], prompts), 1) if len(p) > len(t["prompt"]) + 20000]
    assert pasted[:6] == [2, 3, 4, 5, 6, 7] and len(pasted) >= 10
    assert max(len(p) for p in prompts) < 100_000  # pico_chat_max_prompt_chars
    assert prompts == [lte.turn_prompt(t) for t in lx4["turns"]]  # seeded
    chat = "".join(p[len(t["prompt"]):] for t, p in zip(lx4["turns"], prompts))
    for word in ("班费", "转学", "退费", "¥", "高若溪", "许静怡"):
        assert word not in chat  # filler never moves an LX2 answer
    assert "api_v1" in chat  # the edit-v1 bait


def test_per_turn_files_land_under_turns(monkeypatch, tmp_path: Path) -> None:
    seen = {}

    def fake_checker(case, out_dir, res, image, runtime):
        seen["files"] = sorted(str(p.relative_to(out_dir)) for p in out_dir.rglob("*") if p.is_file())
        return 1, True

    monkeypatch.setattr(lte, "run_checker", fake_checker)
    case = {"expect": {"files": ["a.zip"], "check": {"script": "x.py", "max": 1, "per_turn": True}}}
    res = lte.CaseResult(case="X")
    lte.score_code(case, {"a.zip": b"v2"}, res, "img", "rt", [{"a.zip": b"v1"}, {}, {"a.zip": b"v2"}])
    assert seen["files"] == ["a.zip", "turns/R01/a.zip", "turns/R03/a.zip"]
    assert res.ok and res.auto_score == 2


LX2_SOLUTION = {
    "classfund/ledger.py": '''
    def total_by_category(self, start=None, end=None):
        out = {}
        for e in self.entries:
            if (start is None or e.day >= start) and (end is None or e.day <= end):
                out[e.category] = out.get(e.category, 0) + e.amount_fen
        return out
''',
    "classfund/trip.py": '''
def split_evenly(total_fen, sids):
    sids = sorted(sids)
    base, rest = divmod(total_fen, len(sids))
    return {s: base + (1 if i < rest else 0) for i, s in enumerate(sids)}
''',
    "classfund/refund.py": '''
def refund_fen(paid_fen, weeks_left, total_weeks=20):
    return paid_fen * weeks_left // total_weeks
''',
    "classfund/bank.py": '''
from classfund.money import parse_yuan


def parse_bank_amount(text):
    s = text.strip().replace(",", "").replace("\\uffe5", "").lstrip("+")
    if s.startswith("(") and s.endswith(")"):
        return -parse_yuan(s[1:-1])
    return parse_yuan(s)
''',
    "classfund/report.py": '''
from classfund.money import fmt_yuan


def month_report_html(ledger, roster, month):
    rows = [e for e in ledger.entries if e.day.isoformat().startswith(month)]
    income = sum(e.amount_fen for e in rows if e.amount_fen > 0)
    spent = -sum(e.amount_fen for e in rows if e.amount_fen < 0)
    bal = sum(e.amount_fen for e in ledger.entries if e.day.isoformat()[:7] <= month)
    paid = "".join(f"<li>{e.sid} {fmt_yuan(e.amount_fen)}</li>" for e in rows if e.sid in roster)
    return f"<html><body>{fmt_yuan(income)} {fmt_yuan(spent)} {fmt_yuan(bal)}<ul>{paid}</ul></body></html>"
''',
    "classfund/api_v2.py": '''
from classfund.money import fmt_yuan


def list_entries(ledger, sid):
    return [
        {"date": e.day.isoformat(), "amount_fen": e.amount_fen, "amount": fmt_yuan(e.amount_fen), "category": e.category}
        for e in ledger.entries_for(sid)
    ]


def get_balance(ledger, start=None, end=None):
    rows = [e for e in ledger.entries if (start is None or e.day >= start) and (end is None or e.day <= end)]
    income = sum(e.amount_fen for e in rows if e.amount_fen > 0)
    spent = -sum(e.amount_fen for e in rows if e.amount_fen < 0)
    bal = ledger.balance()
    return {"balance_fen": bal, "balance": fmt_yuan(bal), "income_fen": income, "income": fmt_yuan(income),
            "spent_fen": spent, "spent": fmt_yuan(spent)}


def get_statement(ledger, sid):
    paid = sum(e.amount_fen for e in ledger.entries_for(sid) if e.category == "班费")
    return {"sid": sid, "paid_fen": paid, "paid": fmt_yuan(paid), "entries": list_entries(ledger, sid)}
''',
    "classfund/budget.py": '''
def budget_status(ledger, budgets, start=None, end=None):
    out = []
    for cat, budget in budgets.items():
        used = -sum(
            e.amount_fen for e in ledger.entries
            if e.category == cat and e.amount_fen < 0
            and (start is None or e.day >= start) and (end is None or e.day <= end)
        )
        out.append({"category": cat, "budget_fen": budget, "used_fen": used,
                    "used_pct": used * 100 // budget, "over": used > budget})
    return out


def budget_reminder(ledger, budgets, start=None, end=None):
    return "\\n".join(
        f"{r['category']} {r['used_pct']}%" + (" 【超预算】" if r["over"] else "")
        for r in budget_status(ledger, budgets, start, end)
    )
''',
    "classfund/settle.py": '''
from classfund.trip import split_evenly


def final_refunds(balance_fen, sids, keep_fen):
    return split_evenly(max(0, balance_fen - keep_fen), sids)
''',
    "classfund/uniform.py": '''
from classfund.trip import split_evenly


def uniform_shares(unit_fen, sids, sponsor_fen=0):
    return split_evenly(max(0, unit_fen * len(sids) - sponsor_fen), sids)
''',
    "classfund/fees.py": '''
def next_term_fees(base_fen, sids, half_sids=()):
    return {s: base_fen // 2 if s in half_sids else base_fen for s in sorted(sids)}
''',
    "classfund/summary.py": '''
from classfund.money import fmt_yuan


def term_summary_html(ledger, start, end):
    rows = [e for e in ledger.entries if start <= e.day <= end]
    cats = {}
    for e in rows:
        cats[e.category] = cats.get(e.category, 0) + abs(e.amount_fen)
    income = sum(e.amount_fen for e in rows if e.amount_fen > 0)
    spent = -sum(e.amount_fen for e in rows if e.amount_fen < 0)
    bal = sum(e.amount_fen for e in ledger.entries if e.day <= end)
    table = "".join(f"<tr><td>{c}</td><td>{fmt_yuan(v)}</td></tr>" for c, v in cats.items())
    return f"<html><body>{fmt_yuan(income)} {fmt_yuan(spent)} {fmt_yuan(bal)}<table>{table}</table></body></html>"
''',
}

# R11: receipts still to book (the three already in the fixture ledger make up the rest).
LX2_Q4 = [("2026-10-12", -3250, "奖品"), ("2026-10-21", -2760, "卫生用品"), ("2026-10-30", -4500, "奖品"),
          ("2026-11-03", -5970, "图书角"), ("2026-11-08", -6840, "班级活动"), ("2026-11-15", -3600, "卫生用品"),
          ("2026-11-22", -23880, "图书角"), ("2026-11-28", -4130, "奖品"), ("2026-12-05", -8520, "班级活动"),
          ("2026-12-12", -2390, "卫生用品"), ("2026-12-18", -9600, "班级活动"), ("2026-12-22", -8880, "奖品"),
          ("2026-12-26", -1200, "图书角")]


def _lx2_project(
    tmp_path: Path, *, edit_v1: bool = False, leak: bool = False, forget_r3: bool = False, slip_turn: int = 0
) -> Path:
    """Correct LX2 delivery built on the fixture; flags break one turn-1 rule each.

    Every turn's delivery lands under out/turns/R01… as a copy whose CHANGELOG tops
    at that turn; ``slip_turn`` edits api_v1.py in that one turn only (restored later).
    """
    import shutil

    sys.path.insert(0, str(lte.CHECKS_DIR))
    import lx2_classfund as chk

    root = tmp_path / "out" / "classfund"
    shutil.copytree(CASES / "fixtures" / "code" / "classfund", root)
    for rel, code in LX2_SOLUTION.items():
        path = root / rel
        if path.exists():  # ledger.py: the method goes into class Ledger
            text = path.read_text(encoding="utf-8").replace("\n\ndef load_csv", code + "\n\ndef load_csv", 1)
            path.write_text(text, encoding="utf-8")
        else:
            path.write_text(code.lstrip(), encoding="utf-8")
    if edit_v1:
        v1 = root / "classfund" / "api_v1.py"
        v1.write_text(v1.read_text(encoding="utf-8").replace("{e.day.year}/{e.day.month}/{e.day.day}", "{e.day.isoformat()}"), encoding="utf-8")
    names = dict(line.split(",") for line in (root / "data" / "roster.csv").read_text(encoding="utf-8").split()[1:])
    now = chk.ALL if forget_r3 else chk.CURRENT
    settle = chk.split(123456 - 20000, now)
    uniform = chk.split(len(set(now) - chk.OWN_SHIRT) * 4680 - 30000, sorted(set(now) - chk.OWN_SHIRT))
    data = root / "data"
    _lx2_sheet(data / "春游分摊.csv", chk.TRIP, names if leak else None)
    _lx2_sheet(data / "退费.csv", chk.TRANSFERRED)
    _lx2_sheet(data / "结余返还.csv", settle)
    _lx2_sheet(data / "班服分摊.csv", uniform)
    _lx2_sheet(data / "下学期收费.csv", {s: 15000 if s in chk.HALF else 30000 for s in [*now, chk.NEW_SID]})
    (data / "budgets.json").write_text(json.dumps(chk.BUDGETS, ensure_ascii=False), encoding="utf-8")
    with (data / "roster.csv").open("a", encoding="utf-8") as fh:
        fh.write(f"{chk.NEW_SID},秦朗\n")
    with (data / "entries.csv").open("a", encoding="utf-8") as fh:
        fh.writelines(f"{day},,{fen},{cat},\n" for day, fen, cat in LX2_Q4)
    own = root / "tests" / "test_ledger.py"
    own.write_text(own.read_text(encoding="utf-8").replace("48)", f"{48 + len(LX2_Q4)})"), encoding="utf-8")
    q4 = [("2026-10-10", 8600, "图书角"), ("2026-10-18", 15000, "班级活动"), ("2026-10-24", 1980, "卫生用品")]
    q4 += [(day, -fen, cat) for day, fen, cat in LX2_Q4]
    (data / "10-12月支出明细.csv").write_text(
        "日期,类别,说明,金额\n" + "".join(f"{d},{c},买东西,¥{f // 100}.{f % 100:02d}\n" for d, f, c in sorted(q4))
        + "合计,,,¥1111.00\n",
        encoding="utf-8",
    )
    log = root / "CHANGELOG.md"
    heads = "".join(f"## R{n} 第 {n} 轮\n\n" for n in range(15, 0, -1))
    log.write_text(log.read_text(encoding="utf-8").replace("## R0", heads + "## R0"), encoding="utf-8")
    final = log.read_text(encoding="utf-8")
    for n in range(1, 16):
        turn = tmp_path / "out" / "turns" / f"R{n:02d}" / "classfund"
        shutil.copytree(root, turn)
        (turn / "CHANGELOG.md").write_text(final[final.index(f"## R{n} "):], encoding="utf-8")
        if n == slip_turn:
            v1 = turn / "classfund" / "api_v1.py"
            v1.write_text(v1.read_text(encoding="utf-8") + "\n", encoding="utf-8")
    zin = tmp_path / "in"
    zin.mkdir()
    (zin / "classfund.zip").write_bytes(lte._zip_dir("fixtures/code/classfund"))
    return tmp_path


def _lx2_sheet(path: Path, rows: dict[str, int], names: dict[str, str] | None = None) -> None:
    lines = ["学号,金额"]
    for sid, fen in sorted(rows.items()):
        lines.append(f"{sid},¥{fen // 100}.{fen % 100:02d}" + (f",{names[sid]}" if names else ""))
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _run_lx2_checker(monkeypatch, capsys, w: Path) -> dict:
    sys.path.insert(0, str(lte.CHECKS_DIR))
    import _common
    import lx2_classfund as chk

    out, real_find = str(w / "out"), _common.find
    monkeypatch.setattr(chk, "find", lambda name, root=out: real_find(name, root))
    monkeypatch.setattr(chk, "IN", str(w / "in"))
    monkeypatch.setattr(chk, "OUT", out)
    monkeypatch.chdir(w)
    try:
        chk.main()
    except SystemExit:
        pass
    return json.loads(capsys.readouterr().out.strip().splitlines()[-1])


def test_lx2_is_fifteen_turns_on_one_project() -> None:
    case = next(c for c in lte.load_cases("xlong") if c["id"] == "LX2")
    assert len(case["turns"]) == 15
    assert case["turns"][0]["attachments"][0]["zip_dir"] == "fixtures/code/classfund"
    for turn in case["turns"]:
        for att in turn.get("attachments") or []:
            assert lte._attachment_bytes(att)
    assert lte.max_points(case["expect"]) == 24


def test_lx2_checker_full_marks_on_reference(monkeypatch, capsys, tmp_path: Path) -> None:
    verdict = _run_lx2_checker(monkeypatch, capsys, _lx2_project(tmp_path))
    assert verdict == {"points": 23, "pass": True, "notes": []}, verdict


def test_lx2_checker_catches_broken_turn1_rules(monkeypatch, capsys, tmp_path: Path) -> None:
    """Each flag forgets one thing a compaction could drop; each must fail the case."""
    for kw, needle in (
        ({"edit_v1": True}, "api_v1.py was edited"),
        ({"leak": True}, "student names"),
        ({"forget_r3": True}, "R8 结余返还.csv ok=False"),
        ({"slip_turn": 9}, "R9 api_v1 edited"),  # broken in one turn, restored by the last
    ):
        verdict = _run_lx2_checker(monkeypatch, capsys, _lx2_project(tmp_path / next(iter(kw)), **kw))
        assert verdict["pass"] is False, (kw, verdict)
        assert any(needle in n for n in verdict["notes"]), (kw, verdict)
