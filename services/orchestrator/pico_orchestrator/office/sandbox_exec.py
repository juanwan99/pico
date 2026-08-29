"""Isolated python-pptx subprocess entry. Copied into a temp dir.

Not a second Office OS. User source runs with an allowlisted import
hook: pptx / pptx_helpers plus a pathlib stub. Presentation.save always
lands on the ledger OUTPUT_PATH so naked GPT `prs.save("/tmp/…")` still
books. Host open/eval/os stay denied.
"""

from __future__ import annotations

import json
import sys
import types
from typing import Any

import pptx as _pptx_pkg
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.presentation import Presentation as PresentationClass
from pptx.util import Emu, Inches, Pt
from pptx_helpers import (
    ImagePathMap,
    add_content_slide,
    add_table,
    add_title_slide,
    install_blank_slide_compat,
)

_pptx_pkg.Inches = Inches
_pptx_pkg.Pt = Pt
_pptx_pkg.Emu = Emu
_pptx_pkg.RGBColor = RGBColor
install_blank_slide_compat()


class _SandboxPath:
    """pathlib.Path stand-in: mkdir is ignored; host IO is closed."""

    def __init__(self, *parts: Any) -> None:
        bits: list[str] = []
        for part in parts:
            text = str(part)
            if bits and text.startswith("/"):
                bits = [text]
            else:
                bits.append(text)
        self._parts = tuple(bits)

    def mkdir(self, *args: Any, **kwargs: Any) -> None:
        return None

    def __truediv__(self, other: Any) -> _SandboxPath:
        return _SandboxPath(*self._parts, other)

    def __rtruediv__(self, other: Any) -> _SandboxPath:
        return _SandboxPath(other, *self._parts)

    def __str__(self) -> str:
        if not self._parts:
            return ""
        joined = self._parts[0]
        for part in self._parts[1:]:
            joined = joined.rstrip("/") + "/" + str(part).lstrip("/")
        return joined

    def __repr__(self) -> str:
        return f"Path({str(self)!r})"

    def __fspath__(self) -> str:
        return str(self)

    def __eq__(self, other: object) -> bool:
        return str(self) == str(other)

    @property
    def parent(self) -> _SandboxPath:
        if len(self._parts) <= 1:
            return _SandboxPath("/")
        return _SandboxPath(*self._parts[:-1])

    @property
    def name(self) -> str:
        if not self._parts:
            return ""
        return str(self._parts[-1]).rstrip("/").rsplit("/", 1)[-1]

    def resolve(self) -> _SandboxPath:
        return self

    def absolute(self) -> _SandboxPath:
        return self

    def exists(self) -> bool:
        return False

    def is_file(self) -> bool:
        return False

    def is_dir(self) -> bool:
        return False

    def as_posix(self) -> str:
        return str(self)

    def with_name(self, name: Any) -> _SandboxPath:
        if not self._parts:
            return _SandboxPath(name)
        return _SandboxPath(*self._parts[:-1], name)

    def _deny_host_io(self, *args: Any, **kwargs: Any) -> None:
        raise PermissionError("sandbox pathlib cannot read or write host files")

    open = _deny_host_io
    read_text = _deny_host_io
    read_bytes = _deny_host_io
    write_text = _deny_host_io
    write_bytes = _deny_host_io
    unlink = _deny_host_io
    rmdir = _deny_host_io
    rename = _deny_host_io
    replace = _deny_host_io
    touch = _deny_host_io
    chmod = _deny_host_io
    glob = _deny_host_io
    rglob = _deny_host_io
    iterdir = _deny_host_io

    @classmethod
    def cwd(cls) -> _SandboxPath:
        return cls(".")

    @classmethod
    def home(cls) -> _SandboxPath:
        return cls("/sandbox")


def _pathlib_stub() -> types.ModuleType:
    mod = types.ModuleType("pathlib")
    mod.Path = _SandboxPath  # type: ignore[attr-defined]
    mod.PurePath = _SandboxPath  # type: ignore[attr-defined]
    mod.PurePosixPath = _SandboxPath  # type: ignore[attr-defined]
    return mod


_PATHLIB = _pathlib_stub()


def _sandbox_import(
    name: Any,
    globals: Any = None,
    locals: Any = None,
    fromlist: Any = (),
    level: int = 0,
) -> Any:
    if level != 0:
        raise ImportError("relative import denied")
    text = str(name or "")
    root = text.split(".")[0]
    if root == "pathlib":
        if text != "pathlib":
            raise ImportError("denied import " + text)
        return _PATHLIB
    if root not in ("pptx", "pptx_helpers"):
        raise ImportError("denied import " + text)
    mod = __import__(name, globals, locals, fromlist, level)
    if root == "pptx":
        import pptx as pkg

        pkg.Inches = Inches
        pkg.Pt = Pt
        pkg.Emu = Emu
        pkg.RGBColor = RGBColor
    return mod


_SAFE_BUILTINS: dict[str, Any] = {
    "abs": abs,
    "all": all,
    "any": any,
    "bool": bool,
    "bytes": bytes,
    "bytearray": bytearray,
    "callable": callable,
    "chr": chr,
    "classmethod": classmethod,
    "dict": dict,
    "divmod": divmod,
    "enumerate": enumerate,
    "filter": filter,
    "float": float,
    "format": format,
    "frozenset": frozenset,
    "getattr": getattr,
    "hasattr": hasattr,
    "hash": hash,
    "hex": hex,
    "int": int,
    "isinstance": isinstance,
    "issubclass": issubclass,
    "iter": iter,
    "len": len,
    "list": list,
    "map": map,
    "max": max,
    "min": min,
    "next": next,
    "object": object,
    "oct": oct,
    "ord": ord,
    "pow": pow,
    "print": print,
    "property": property,
    "range": range,
    "repr": repr,
    "reversed": reversed,
    "round": round,
    "set": set,
    "setattr": setattr,
    "slice": slice,
    "sorted": sorted,
    "staticmethod": staticmethod,
    "str": str,
    "sum": sum,
    "super": super,
    "tuple": tuple,
    "type": type,
    "zip": zip,
    "Exception": Exception,
    "BaseException": BaseException,
    "TypeError": TypeError,
    "ValueError": ValueError,
    "KeyError": KeyError,
    "IndexError": IndexError,
    "AttributeError": AttributeError,
    "RuntimeError": RuntimeError,
    "OSError": OSError,
    "PermissionError": PermissionError,
    "StopIteration": StopIteration,
    "ArithmeticError": ArithmeticError,
    "ZeroDivisionError": ZeroDivisionError,
    "OverflowError": OverflowError,
    "LookupError": LookupError,
    "NameError": NameError,
    "UnboundLocalError": UnboundLocalError,
    "AssertionError": AssertionError,
    "NotImplementedError": NotImplementedError,
    "StopAsyncIteration": StopAsyncIteration,
    "__build_class__": __build_class__,
    "__name__": "builtins",
    "__import__": _sandbox_import,
}


def _ledger_save(orig_save: Any, output_path: str) -> Any:
    def _save(self: Any, path: Any = None, *args: Any, **kwargs: Any) -> Any:
        return orig_save(self, output_path)

    return _save


def run_user(cfg: dict[str, Any]) -> None:
    output_path = str(cfg["output"])
    orig_save = PresentationClass.save
    PresentationClass.save = _ledger_save(orig_save, output_path)  # type: ignore[method-assign]

    def save_deck(prs: Any) -> None:
        if not hasattr(prs, "save"):
            raise SystemExit("save_deck 需要 Presentation")
        prs.save(output_path)

    exec(  # noqa: S102 — isolated allowlisted snippet
        cfg["source"],
        {
            "__builtins__": _SAFE_BUILTINS,
            "Presentation": Presentation,
            "Inches": Inches,
            "Emu": Emu,
            "Pt": Pt,
            "RGBColor": RGBColor,
            "Path": _SandboxPath,
            "add_title_slide": add_title_slide,
            "add_content_slide": add_content_slide,
            "add_table": add_table,
            "save_deck": save_deck,
            "IMAGE_PATHS": ImagePathMap(cfg.get("images") or {}),
            "OUTPUT_PATH": output_path,
            "__name__": "__main__",
        },
    )


def main() -> None:
    if len(sys.argv) > 1:
        from pathlib import Path as _RealPath

        raw = _RealPath(sys.argv[1]).read_text(encoding="utf-8")
    else:
        raw = sys.stdin.buffer.read().decode("utf-8")
    cfg: dict[str, Any] = json.loads(raw or "{}")
    run_user(cfg)


if __name__ == "__main__":
    main()
