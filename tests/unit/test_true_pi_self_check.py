"""pico#1334 ②: before landing, the answer is checked against outputs/ in the live box."""

from __future__ import annotations

import base64
import json
import subprocess
import sys
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import pytest

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "services" / "orchestrator"))

from pico_orchestrator.run_types import RunCaps
from pico_orchestrator.true_pi.client import FakeTransport
from pico_orchestrator.true_pi.runtime import run_true_pi_agent
from pico_orchestrator.true_pi.self_check import (
    _BOX_SCRIPT,
    box_check_command,
    parse_box_report,
    self_check_facts,
    self_check_failure,
    self_check_prompt,
)


def _report(files: list[dict[str, Any]] | None = None, names: list[str] | None = None) -> dict[str, Any]:
    files = files or []
    return {"files": files, "names": names if names is not None else [f["name"] for f in files]}


def test_named_file_missing_is_a_fact() -> None:
    # Ruler s13 R16 turn 2: renamed in the answer, never written.
    text = "好的，已按要求修改并重新生成 Word 文件。文件已保存到 `/workspace/outputs/高三（2）班班级公约.docx`。"
    facts = self_check_facts(text, _report(names=["班级公约.docx"]), had_files_before=True)
    assert facts == ["回答提到的 高三（2）班班级公约.docx 在 outputs 里没有。"]


def test_naming_the_teachers_upload_is_not_a_miss() -> None:
    # Ruler s4b R10 / R11: the answer named the uploaded input; the old check sent the model
    # to copy it into outputs/ (R10) or failed the run (R11).
    report = {**_report([{"name": "期中考试通知.docx", "ext": ".docx", "chars": 300}]), "others": ["r11_notice.pdf"]}
    text = "已根据 `r11_notice.pdf` 转成 Word：`/workspace/outputs/期中考试通知.docx`。"
    assert self_check_facts(text, report, had_files_before=False) == []
    # An outputs/ path still has to be in outputs.
    assert self_check_facts("见 outputs/r11_notice.pdf", report, had_files_before=False) == [
        "回答提到的 r11_notice.pdf 在 outputs 里没有。"
    ]


def test_claim_without_new_file_same_as_delivery_gate() -> None:
    claim = "通知已生成，请下载。"
    assert self_check_facts(claim, _report(), had_files_before=False)
    # Restating an earlier turn's file is not a miss (#1167) …
    assert self_check_facts(claim, _report(names=["a.docx"]), had_files_before=True) == []
    # … saying it was made again this turn is.
    assert self_check_facts("已重新生成课件。", _report(names=["a.pptx"]), had_files_before=True)


def test_good_delivery_and_clarification_have_no_facts() -> None:
    good = _report([{"name": "班级公约.docx", "ext": ".docx", "chars": 300}])
    assert self_check_facts("已完成！文件已保存到 `/workspace/outputs/班级公约.docx`。", good, had_files_before=False) == []
    ask = "开始之前想先确认两点：要几页？用哪个模板？"
    assert self_check_facts(ask, _report(), had_files_before=False) == []


def test_unopenable_and_empty_files() -> None:
    facts = self_check_facts(
        "已生成。",
        _report(
            [
                {"name": "a.docx", "ext": ".docx", "error": "BadZipFile"},
                {"name": "b.docx", "ext": ".docx", "chars": 3},
                {"name": "c.pptx", "ext": ".pptx", "slides": 2, "titled": 0},
                {"name": "d.xlsx", "ext": ".xlsx", "cells": 0},
            ]
        ),
        had_files_before=False,
    )
    assert facts == [
        "a.docx 打不开（BadZipFile）。",
        "b.docx 打开后正文只有 3 个字。",
        "c.pptx 打开后没有有字的页。",
        "d.xlsx 打开后所有表都是空的。",
    ]


def test_counts_against_the_files() -> None:
    deck = _report([{"name": "光合作用.pptx", "ext": ".pptx", "slides": 6, "titled": 6}])
    assert self_check_facts("已生成 10 张幻灯片的课件。", deck, had_files_before=False) == [
        "光合作用.pptx 有 6 张幻灯片，回答说 10 张。"
    ]
    assert self_check_facts("已生成六页 PPT。", deck, had_files_before=False) == []
    pdf = _report([{"name": "卷.pdf", "ext": ".pdf", "pages": 3}])
    assert self_check_facts("已导出 4 页 PDF。", pdf, had_files_before=False) == ["卷.pdf 有 3 页，回答说 4 页。"]
    three = _report([{"name": f"{i}.docx", "ext": ".docx", "chars": 300} for i in range(3)])
    assert self_check_facts("5 份家长通知已生成。", three, had_files_before=False) == [
        "这一轮 outputs 里有 3 个新文件，回答说 5 份。"
    ]
    # Five notices in one document is a choice, not a miss.
    one = _report([{"name": "通知.docx", "ext": ".docx", "chars": 900}])
    assert self_check_facts("5 份家长通知已生成，在一个文档里。", one, had_files_before=False) == []


def test_prompt_is_facts_only_and_failure_fits_user_errors() -> None:
    facts = ["回答提到的 a.docx 在 outputs 里没有。", "b.pptx 有 6 张幻灯片，回答说 10 张。"]
    assert self_check_prompt(facts) == "系统核对：\n- 回答提到的 a.docx 在 outputs 里没有。\n- b.pptx 有 6 张幻灯片，回答说 10 张。"
    msg = self_check_failure(facts)
    assert msg.startswith("交付前核对没通过：回答提到的 a.docx") and "另有 1 处" in msg
    assert len(self_check_failure(["很长" * 200])) <= 180


def test_box_script_reads_outputs(tmp_path: Path) -> None:
    import docx
    import pptx

    out = tmp_path / "outputs"
    out.mkdir()
    d = docx.Document()
    d.add_paragraph("各位家长：本周五下午三点在报告厅召开家长会，请准时参加。")
    d.save(out / "通知.docx")
    p = pptx.Presentation()
    for _ in range(3):
        s = p.slides.add_slide(p.slide_layouts[1])
        s.shapes.title.text = "页"
    p.save(out / "课件.pptx")
    (out / "坏.xlsx").write_bytes(b"not a zip")
    (out / "旧.md").write_text("old")
    (tmp_path / "attachments").mkdir()
    (tmp_path / "attachments" / "r11_notice.pdf").write_bytes(b"%PDF")
    (tmp_path / ".pico").mkdir()
    (tmp_path / ".pico" / "session.jsonl").write_text("{}")
    import hashlib

    before = {"outputs/旧.md": {"sha256": hashlib.sha256(b"old").hexdigest()}}
    snap = base64.b64encode(json.dumps({k: v["sha256"] for k, v in before.items()}).encode()).decode()
    run = subprocess.run(
        [sys.executable, "-", snap, str(tmp_path)], input=_BOX_SCRIPT, capture_output=True, text=True, check=True
    )
    report = parse_box_report(run.stdout)
    assert report is not None
    rows = {f["name"]: f for f in report["files"]}
    assert set(rows) == {"通知.docx", "课件.pptx", "坏.xlsx"}
    assert rows["通知.docx"]["chars"] > 20
    assert rows["课件.pptx"]["slides"] == 3 and rows["课件.pptx"]["titled"] == 3
    assert rows["坏.xlsx"]["error"]
    assert "旧.md" in report["names"]
    assert report["others"] == ["r11_notice.pdf"]
    assert box_check_command(before).startswith("echo ")


def _turn(answer: str) -> list[dict[str, Any]]:
    msg = {"role": "assistant", "content": [{"type": "text", "text": answer}]}
    return [
        {"type": "agent_start"},
        {"type": "turn_start"},
        {"type": "message_end", "message": msg},
        {"type": "turn_end", "toolResults": [], "message": msg},
        {"type": "agent_end", "willRetry": False, "messages": []},
    ]


class _Principal:
    school_id = "school-a"
    membership_id = "member-a"
    scopes = ("ai:run",)


async def _not_cancelled() -> bool:
    return False


async def _run(monkeypatch, transport: FakeTransport) -> tuple[Any, list[tuple[str, dict[str, Any]]]]:
    from pico_orchestrator.true_pi import runner

    async def fake_list(_key: Any) -> dict[str, Any]:
        # Turn 1 of the conversation already wrote 班级公约.docx.
        return {"outputs/班级公约.docx": {"sha256": "v1"}}

    async def fake_land(**_kw: Any) -> list[tuple[str, dict[str, Any]]]:
        return [("workspace_output", {"artifact_id": "a1", "title": "高三（2）班班级公约.docx"})]

    monkeypatch.setattr(runner, "list_outputs", fake_list)
    monkeypatch.setattr(runner, "land_outputs", fake_land)
    transport.runner = SimpleNamespace(key="ws-key")
    events: list[tuple[str, dict[str, Any]]] = []

    async def emit(k: str, p: dict[str, Any]) -> None:
        events.append((k, p))

    result = await run_true_pi_agent(
        prompt="把标题改成高三（2）班班级公约",
        principal=_Principal(),
        emit=emit,
        is_cancelled=_not_cancelled,
        caps=RunCaps(min_artifacts=0, max_seconds=20),
        transport=transport,
    )
    return result, events


_MISSING = json.dumps({"files": [], "names": ["班级公约.docx"]})
_FIXED = json.dumps(
    {"files": [{"name": "高三（2）班班级公约.docx", "ext": ".docx", "chars": 200}], "names": ["高三（2）班班级公约.docx"]}
)
_CLAIM = "已修改，文件已保存到 `/workspace/outputs/高三（2）班班级公约.docx`。"


@pytest.mark.asyncio
async def test_one_extra_round_on_facts_then_succeeds(monkeypatch) -> None:
    transport = FakeTransport(
        scripted=_turn(_CLAIM),
        scripted_after_prompt=[_turn(_CLAIM)],
        bash_outputs=[_MISSING, _FIXED],
        assistant_text="",
    )
    result, events = await _run(monkeypatch, transport)
    assert result.status == "succeeded"
    prompts = [c["message"] for c in transport.sent if c.get("type") == "prompt"]
    assert prompts[-1] == "系统核对：\n- 回答提到的 高三（2）班班级公约.docx 在 outputs 里没有。"
    bash = [c for c in transport.sent if c.get("type") == "bash"]
    assert len(bash) == 2 and all(c["excludeFromContext"] is True for c in bash)
    resume = [p for k, p in events if k == "run.resume"]
    assert len(resume) == 1 and resume[0]["reason"] == "self_check"
    assert resume[0]["text"].startswith("交付前自查：")


@pytest.mark.asyncio
async def test_still_wrong_after_the_round_fails_on_the_fact(monkeypatch) -> None:
    transport = FakeTransport(
        scripted=_turn(_CLAIM),
        scripted_after_prompt=[_turn(_CLAIM)],
        bash_outputs=[_MISSING, _MISSING],
        assistant_text="",
    )
    result, _events = await _run(monkeypatch, transport)
    assert result.status == "failed"
    assert result.code == "delivery.self_check"
    assert "高三（2）班班级公约.docx 在 outputs 里没有" in (result.error or "")
    assert sum(1 for c in transport.sent if c.get("type") == "prompt") == 2


@pytest.mark.asyncio
async def test_check_that_cannot_run_never_blocks(monkeypatch) -> None:
    transport = FakeTransport(scripted=_turn(_CLAIM), bash_outputs=["python3: not found"], assistant_text="")
    result, events = await _run(monkeypatch, transport)
    assert result.status == "succeeded"
    assert not [k for k, _ in events if k == "run.resume"]
    assert sum(1 for c in transport.sent if c.get("type") == "prompt") == 1
